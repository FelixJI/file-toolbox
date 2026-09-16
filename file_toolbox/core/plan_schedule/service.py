"""计划排布核心:项点清单(名称+起止日期)-> 按月分块的排布工作簿。

纯 openpyxl 实现,跨平台、不依赖 Office/WPS COM。openpyxl 在方法内延迟导入,
避免 CLI 入口与 GUI 启动链为 `--help`/首屏付出导入成本(与 excel_merge 同策略)。

输出布局参照交付计划模板:A 列项点名称,每月一个块 —— MONTH 合并行("2026年9月")
+ DATE 行(1..N)+ 项点行 + 并行数行。项点行在活动日期写"项点内第几天"
(跨月连续编号)并按项点循环填色;周末整列浅灰底、DATE 表头红字;并行数行给出
每天并行的项点数。输出文件已存在时自动加序号,绝不覆盖。
"""

from __future__ import annotations

import calendar
from collections.abc import Callable
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import TYPE_CHECKING, Any

from file_toolbox.common.history import JsonHistoryStore
from file_toolbox.common.loggable import LoggableMixin
from file_toolbox.core.plan_schedule.constants import (
    CELL_NAME,
    DATE_FORMATS_FULL,
    DATE_FORMATS_NO_YEAR,
    DAY_COLUMN_WIDTH,
    END_HEADERS,
    ITEM_FILLS,
    MAX_HEADER_SCAN_ROWS,
    NAME_COLUMN_WIDTH,
    NAME_HEADERS,
    NAME_MODE_DAY_WIDTH_MAX,
    SHEET_NAME,
    START_HEADERS,
    SUPPORTED_SUFFIXES,
    WEEKEND_FILL,
    WEEKEND_HEADER_FONT,
)
from file_toolbox.core.plan_schedule.types import (
    InvalidRow,
    MonthPlan,
    PlanItem,
    ScheduleOptions,
    ScheduleResult,
)

if TYPE_CHECKING:
    from openpyxl import Workbook
    from openpyxl.worksheet.worksheet import Worksheet

ProgressCallback = Callable[[int, int, str], None]

# Excel 1900 日期系统的序列日期起点(openpyxl 读到的裸数字按此换算;
# 1900-01-01 之前的序号无意义)
_EXCEL_SERIAL_EPOCH = date(1899, 12, 30)


class PlanScheduleService(LoggableMixin):
    """计划排布服务:parse 解析清单、plan 排布计算、generate 生成工作簿。"""

    def __init__(self, history_store: JsonHistoryStore | None = None) -> None:
        """初始化服务。

        Args:
            history_store: 历史存储;传入则 generate 成功写出输出后记录一条
                plan_schedule 历史。None 表示不记录(默认)。
        """
        self._history_store = history_store

    # ==================== 输入解析 ====================

    def parse(
        self, input_path: Path, options: ScheduleOptions | None = None
    ) -> tuple[list[PlanItem], list[InvalidRow]]:
        """读取项点清单:定位表头,逐行解析 项点名称/起始日期/终止日期。

        Args:
            input_path: 输入 Excel(.xlsx/.xlsm,取第一个工作表)。
            options: 选项(default_year 供缺年份的日期串使用);None 用默认。

        Returns:
            (项点列表, 无效行列表)。行级问题进 invalid,不中断其余行。

        Raises:
            ValueError: 文件级问题(格式不支持/不存在/无法读取/找不到表头)。
        """
        if options is None:
            options = ScheduleOptions()
        rows = self._read_rows(input_path)
        located = self._locate_header(rows)
        if located is None:
            raise ValueError(
                f"前 {MAX_HEADER_SCAN_ROWS} 行内未找到表头:"
                f"需同时包含 项点名称({'/'.join(NAME_HEADERS)})、"
                f"起始日期({'/'.join(START_HEADERS)})、终止日期({'/'.join(END_HEADERS)}) 列"
            )
        header_row, name_col, start_col, end_col = located

        items: list[PlanItem] = []
        invalid: list[InvalidRow] = []
        for offset, row in enumerate(rows[header_row + 1 :]):
            row_no = header_row + 2 + offset  # 1-based Excel 行号
            name_v = self._cell(row, name_col)
            start_v = self._cell(row, start_col)
            end_v = self._cell(row, end_col)
            if self._is_blank(name_v) and self._is_blank(start_v) and self._is_blank(end_v):
                continue  # 整行空:跳过(清单常见尾部空行)
            name = self._text(name_v)
            if not name:
                invalid.append(InvalidRow(row_no, "缺少项点名称"))
                continue
            start = self._parse_date(start_v, options.default_year)
            if start is None:
                invalid.append(InvalidRow(row_no, f"起始日期无法识别:{self._display(start_v)}"))
                continue
            end = self._parse_date(end_v, options.default_year)
            if end is None:
                invalid.append(InvalidRow(row_no, f"终止日期无法识别:{self._display(end_v)}"))
                continue
            if end < start:
                invalid.append(InvalidRow(row_no, f"终止日期 {end} 早于起始日期 {start}"))
                continue
            items.append(PlanItem(name=name, start=start, end=end, row=row_no))
        if invalid:
            self.logger.warning("计划排布输入有 %d 行无效: %s", len(invalid), invalid[:3])
        return items, invalid

    # ==================== 排布计算(纯逻辑) ====================

    def plan(self, items: list[PlanItem]) -> list[MonthPlan]:
        """把项点按月分块:每月给出活跃项点、逐日"项点内第几天"与并行数。

        月份范围覆盖最早开始到最晚结束;项点按(开始日期, 名称)排序,
        编号跨月连续(9/30 的次日在 10 月块里接续编号)。
        """
        if not items:
            return []
        ordered = sorted(items, key=lambda it: (it.start, it.name))
        first = min(it.start for it in ordered)
        last = max(it.end for it in ordered)
        months: list[MonthPlan] = []
        year, month = first.year, first.month
        while (year, month) <= (last.year, last.month):
            days = calendar.monthrange(year, month)[1]
            month_start = date(year, month, 1)
            month_end = date(year, month, days)
            weekends = frozenset(
                d for d in range(1, days + 1) if date(year, month, d).weekday() >= 5
            )
            active: list[tuple[PlanItem, list[tuple[int, int]]]] = []
            parallel: dict[int, int] = {}
            for item in ordered:
                if item.start > month_end or item.end < month_start:
                    continue
                cells = [
                    (d, (date(year, month, d) - item.start).days + 1)
                    for d in range(1, days + 1)
                    if item.start <= date(year, month, d) <= item.end
                ]
                active.append((item, cells))
                for d, _ in cells:
                    parallel[d] = parallel.get(d, 0) + 1
            months.append(
                MonthPlan(
                    year=year,
                    month=month,
                    days=days,
                    weekends=weekends,
                    items=active,
                    parallel=parallel,
                )
            )
            year, month = (year + 1, 1) if month == 12 else (year, month + 1)
        return months

    @staticmethod
    def peak_parallel(months: list[MonthPlan]) -> tuple[int, date] | None:
        """全部月份中的最大并行数及首个达到日(供预览/摘要展示)。"""
        best: tuple[int, date] | None = None
        for mp in months:
            for d, count in mp.parallel.items():
                current = date(mp.year, mp.month, d)
                if best is None or count > best[0]:
                    best = (count, current)
        return best

    # ==================== 模板 ====================

    def write_template(self, path: Path) -> Path:
        """写出输入清单模板(表头 + 两行虚构示例),已存在时自动加序号。"""
        from openpyxl import Workbook

        target = path if path.suffix.lower() == ".xlsx" else path.with_suffix(".xlsx")
        wb = Workbook()
        ws = wb.active
        assert ws is not None  # 新建工作簿必有默认表
        ws.title = "项点清单"
        ws.append(["项点名称", "起始日期", "终止日期"])
        ws.append(["示例项点A", date(2026, 9, 17), date(2026, 9, 21)])
        ws.append(["示例项点B", date(2026, 9, 21), date(2026, 9, 27)])
        output = self._save_workbook(wb, target)
        self.logger.info("已写出计划排布输入模板: %s", output)
        return output

    # ==================== 生成 ====================

    def generate(
        self,
        input_path: Path,
        output: Path,
        options: ScheduleOptions | None = None,
        progress_callback: ProgressCallback | None = None,
    ) -> ScheduleResult:
        """解析清单并生成按月排布工作簿。

        Args:
            input_path: 项点清单(.xlsx/.xlsm)。
            output: 输出文件(后缀强制 .xlsx;已存在时自动加序号,绝不覆盖)。
            options: 生成选项;None 用默认。
            progress_callback: (current, total, message) 进度回调(按月份粒度)。

        Returns:
            ScheduleResult:success = 输出已写出(允许部分行无效)。
        """
        if options is None:
            options = ScheduleOptions()
        try:
            items, invalid = self.parse(input_path, options)
        except ValueError as e:
            return ScheduleResult(error_message=str(e))
        if not items:
            reason = "没有可排布的项点" + ("(全部行无效)" if invalid else "(清单为空)")
            return ScheduleResult(invalid=invalid, error_message=reason)

        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        assert ws is not None  # 新建工作簿必有默认表
        ws.title = SHEET_NAME
        months = self.plan(items)
        self._render(ws, months, progress_callback, options.cell_mode)

        output_path = self._normalize_output(output)
        try:
            output_path = self._save_workbook(wb, output_path)
        except Exception as e:
            self.logger.error("排布输出写出失败: %s (%s)", output_path, e)
            return ScheduleResult(items=items, invalid=invalid, error_message=f"输出失败: {e}")
        self.logger.info(
            "计划排布完成: %d 个项点 / %d 个月 -> %s", len(items), len(months), output_path
        )
        result = ScheduleResult(output=output_path, items=items, invalid=invalid)
        try:
            self._record_history(result, len(months))
        except Exception as e:
            self.logger.exception("排布已写出,但历史保存失败: %s", output_path)
            result.warning_message = f"输出已生成,但历史未保存: {e}"
        return result

    # ==================== 内部实现:读取与解析 ====================

    def _read_rows(self, path: Path) -> list[tuple[Any, ...]]:
        """读取输入第一个工作表的全部行(仅值)。文件级问题转 ValueError。"""
        if path.suffix.lower() not in SUPPORTED_SUFFIXES:
            raise ValueError(
                f"不支持的格式 {path.suffix or '(无后缀)'},仅支持 {'/'.join(SUPPORTED_SUFFIXES)}"
            )
        if not path.is_file():
            raise ValueError(f"输入文件不存在:{path}")
        try:
            from openpyxl import load_workbook
        except ImportError as e:  # pragma: no cover -- openpyxl 是 base 依赖,测试环境必装
            raise ImportError("计划排布需要 openpyxl 依赖,请重新安装 file-toolbox") from e
        try:
            wb = load_workbook(path, data_only=True, read_only=True)
        except Exception as e:
            raise ValueError(f"无法读取:{e}") from e
        try:
            if not wb.worksheets:
                raise ValueError("输入工作簿没有工作表")
            return [tuple(row) for row in wb.worksheets[0].iter_rows(values_only=True)]
        finally:
            wb.close()

    def _locate_header(self, rows: list[tuple[Any, ...]]) -> tuple[int, int, int, int] | None:
        """在前 MAX_HEADER_SCAN_ROWS 行内定位 (表头行, 名称列, 起始列, 终止列)。"""
        for idx, row in enumerate(rows[:MAX_HEADER_SCAN_ROWS]):
            norm = [self._header_text(v) for v in row]
            name_col = self._find_col(norm, NAME_HEADERS)
            start_col = self._find_col(norm, START_HEADERS)
            end_col = self._find_col(norm, END_HEADERS)
            if name_col is not None and start_col is not None and end_col is not None:
                return idx, name_col, start_col, end_col
        return None

    @staticmethod
    def _find_col(norm_row: list[str], aliases: tuple[str, ...]) -> int | None:
        """返回首个命中别名的列下标;未命中返回 None。"""
        for idx, text in enumerate(norm_row):
            if text in aliases:
                return idx
        return None

    @staticmethod
    def _cell(row: tuple[Any, ...], col: int) -> Any:
        """安全取第 col 列(行元组可能短于表头宽)。"""
        return row[col] if col < len(row) else None

    @staticmethod
    def _is_blank(value: Any) -> bool:
        """单元格是否为空(None 或纯空白字符串)。"""
        return value is None or (isinstance(value, str) and not value.strip())

    @staticmethod
    def _text(value: Any) -> str:
        """单元格转展示文本:整数值浮点去掉 .0,其余 str() 后去空白。"""
        if value is None:
            return ""
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value).strip()

    @classmethod
    def _display(cls, value: Any) -> str:
        """无效行错误信息里的单元格展示(空值给出占位)。"""
        return cls._text(value) or "(空)"

    @staticmethod
    def _header_text(value: Any) -> str:
        """表头规范化:转字符串并去除所有空白(含全角空格)。"""
        return "" if value is None else "".join(str(value).split())

    def _parse_date(self, value: Any, default_year: int | None) -> date | None:
        """解析日期单元格:datetime/date、Excel 序列数或常见格式字符串。

        缺年份的字符串按 default_year(未指定取当前年份)补全;
        无法识别返回 None,由调用方生成行级错误。
        """
        if value is None or isinstance(value, bool):
            return None
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        if isinstance(value, (int, float)):
            serial = int(value)
            return _EXCEL_SERIAL_EPOCH + timedelta(days=serial) if serial > 0 else None
        if isinstance(value, str):
            text = value.strip()
            if not text:
                return None
            for fmt in DATE_FORMATS_FULL:
                try:
                    return datetime.strptime(text, fmt).date()
                except ValueError:
                    continue
            year = default_year if default_year is not None else date.today().year
            for fmt, sep in DATE_FORMATS_NO_YEAR:
                try:
                    return datetime.strptime(f"{year}{sep}{text}", f"%Y{sep}{fmt}").date()
                except ValueError:
                    continue  # 含 2-29 等非法日期,继续尝试下一格式
            return None
        return None

    # ==================== 内部实现:输出 ====================

    def _render(
        self,
        ws: Worksheet,
        months: list[MonthPlan],
        progress: ProgressCallback | None,
        cell_mode: str,
    ) -> None:
        """把排布月块渲染进工作表:样式先行、再写值、最后合并 MONTH 行。

        cell_mode 决定活动日期格内容:index=项点内第几天,name=项点名称
        (name 模式下日期列按最长名称自适应加宽,不超过 NAME_MODE_DAY_WIDTH_MAX)。
        """
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter

        base_font = Font(name="宋体", size=11)
        bold_font = Font(name="宋体", size=11, bold=True)
        weekend_font = Font(name="宋体", size=11, color=WEEKEND_HEADER_FONT)
        thin = Side(style="thin")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        center = Alignment(horizontal="center", vertical="center")
        weekend_fill = PatternFill("solid", fgColor=WEEKEND_FILL)
        item_fills = [PatternFill("solid", fgColor=color) for color in ITEM_FILLS]

        total = len(months)
        row = 1
        for idx, mp in enumerate(months, start=1):
            if progress is not None:
                progress(idx, total, f"生成 {mp.year}年{mp.month}月")
            last_col = 1 + mp.days  # A=名称,B..=1..N
            block_rows = 2 + len(mp.items) + 1  # MONTH + DATE + 项点行 + 并行数
            # 先建格并统一字体/对齐/边框(合并前完成,避免写 MergedCell)
            for r in range(row, row + block_rows):
                for c in range(1, last_col + 1):
                    cell = ws.cell(row=r, column=c)
                    cell.font = base_font
                    cell.alignment = center
                    cell.border = border

            # MONTH 行
            month_cell = ws.cell(row=row, column=1, value="MONTH")
            month_cell.font = bold_font
            ws.cell(row=row, column=2, value=f"{mp.year}年{mp.month}月")

            # DATE 行:周末表头灰底红字
            date_row = row + 1
            ws.cell(row=date_row, column=1, value="DATE")
            for d in range(1, mp.days + 1):
                cell = ws.cell(row=date_row, column=1 + d, value=d)
                if d in mp.weekends:
                    cell.fill = weekend_fill
                    cell.font = weekend_font

            # 项点行:活动日写"项点内第几天"或项点名称(按 cell_mode)并循环填色;
            # 周末空格灰底
            name_mode = cell_mode == CELL_NAME
            for i, (item, cells) in enumerate(mp.items):
                r = date_row + 1 + i
                ws.cell(row=r, column=1, value=item.name).data_type = "s"
                fill = item_fills[i % len(item_fills)]
                day_index = dict(cells)
                for d in range(1, mp.days + 1):
                    cell = ws.cell(row=r, column=1 + d)
                    if d in day_index:
                        cell.value = item.name if name_mode else day_index[d]
                        if name_mode:
                            cell.data_type = "s"
                        cell.fill = fill
                    elif d in mp.weekends:
                        cell.fill = weekend_fill

            # 并行数行:每天并行项点数(>=1 才写)
            parallel_row = date_row + 1 + len(mp.items)
            label = ws.cell(row=parallel_row, column=1, value="并行数")
            label.font = bold_font
            for d in range(1, mp.days + 1):
                cell = ws.cell(row=parallel_row, column=1 + d)
                count = mp.parallel.get(d, 0)
                if count:
                    cell.value = count
                elif d in mp.weekends:
                    cell.fill = weekend_fill

            # 值与样式全部就位后再合并 MONTH 行
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=last_col)
            row = parallel_row + 1

        ws.column_dimensions["A"].width = NAME_COLUMN_WIDTH
        day_width = DAY_COLUMN_WIDTH
        if cell_mode == CELL_NAME:
            # 名称模式:加宽日期列放下最长项点名(CJK 每字约 2 单位),封顶防 31 列过宽
            max_len = max((len(item.name) for mp in months for item, _ in mp.items), default=0)
            day_width = min(2.0 * max_len + 1.0, NAME_MODE_DAY_WIDTH_MAX)
        for c in range(2, 33):  # B..AF 覆盖最长 31 天
            ws.column_dimensions[get_column_letter(c)].width = day_width

    def _save_workbook(self, wb: Workbook, output: Path) -> Path:
        """在实际写入边界排他创建;同名竞争只重选序号,不覆盖其他文件。"""
        output.parent.mkdir(parents=True, exist_ok=True)
        while True:
            candidate = self._resolve_output_path(output)
            try:
                stream = candidate.open("xb")
            except FileExistsError:
                continue
            try:
                with stream:
                    wb.save(stream)
            except Exception:
                # 仅移除本次排他创建但未完成的输出,保留其他写入者的文件。
                candidate.unlink(missing_ok=True)
                raise
            return candidate

    def _normalize_output(self, output: Path) -> Path:
        """输出统一为 .xlsx。"""
        if output.suffix.lower() != ".xlsx":
            return output.with_suffix(".xlsx")
        return output

    def _resolve_output_path(self, output: Path) -> Path:
        """输出已存在时自动加序号,绝不覆盖已有文件(与 pdf/excel_merge 一致)。"""
        if not output.exists(follow_symlinks=False):
            return output
        counter = 1
        while True:
            candidate = output.with_name(f"{output.stem}_{counter}{output.suffix}")
            if not candidate.exists(follow_symlinks=False):
                return candidate
            counter += 1

    def _record_history(self, result: ScheduleResult, month_count: int) -> None:
        """记录 plan_schedule 历史(若注入了 history_store 且生成成功)。"""
        if self._history_store is None or not result.success:
            return
        assert result.output is not None  # success 蕴含 output 已写出
        self._history_store.add_record(
            "plan_schedule",
            {
                "output": str(result.output),
                "item_count": len(result.items),
                "invalid_count": len(result.invalid),
                "month_count": month_count,
                "success": True,
            },
        )
