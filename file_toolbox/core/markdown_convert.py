"""Markdown 批量转换核心:Word(内置 Pandoc 子进程)与 Excel(openpyxl)。

纯跨平台实现,不依赖 Office/WPS COM:

- Word(docx):调用确定位置的内置 Pandoc(pypandoc 包内 ``files/pandoc.exe``,
  打包形态为发行目录 ``pypandoc/files/pandoc.exe``),markdown 从 stdin 进入,
  docx 从 stdout 二进制写出(``--output=-``);``--sandbox`` 隔离文件访问,
  ``--fail-if-warnings`` 保证缺图片等告警显式失败,绝不静默丢内容。
  不回退 PATH 上的 pandoc,避免运行环境漂移。
- Excel(xlsx):markdown-it-py(commonmark + table)解析 token 流提取表格与
  正文,openpyxl 写工作簿。单元格一律文本(data_type="s" + "@"):防 ``=`` 开头
  被当公式执行、保前导零;超 Excel 单元格上限显式失败,不静默截断。
  代码块里的伪表格是 fence token,不会被误认成表格。

输出复用 write_numbered_output:暂存-原子提交、冲突自动加序号、永不覆盖。
取消在文件间与提交前检查,已完成文件保留。单文件失败记录后继续批量。
"""

from __future__ import annotations

import subprocess
import sys
from collections.abc import Callable
from dataclasses import dataclass, field
from importlib.util import find_spec as _find_spec
from pathlib import Path
from typing import IO, TYPE_CHECKING

from file_toolbox.common.history import JsonHistoryStore
from file_toolbox.common.loggable import LoggableMixin
from file_toolbox.common.operation_errors import preserve_history_result
from file_toolbox.common.runtime import is_packaged_runtime
from file_toolbox.core.output_file import write_numbered_output

if TYPE_CHECKING:
    from markdown_it.token import Token
    from openpyxl.workbook import Workbook
    from openpyxl.worksheet.worksheet import Worksheet

SUPPORTED_SUFFIXES = (".md", ".markdown")
TARGET_DOCX = "docx"
TARGET_XLSX = "xlsx"
SUPPORTED_TARGETS = (TARGET_DOCX, TARGET_XLSX)
EXCEL_MODE_TABLES = "tables"
EXCEL_MODE_DOCUMENT = "document"
SUPPORTED_EXCEL_MODES = (EXCEL_MODE_TABLES, EXCEL_MODE_DOCUMENT)
HISTORY_TOOL = "markdown_convert"
PANDOC_TIMEOUT_SECONDS = 120
# gfm 提供管道表/删除线/任务列表,+tex_math_dollars 保留数学,-raw_html 禁裸 HTML。
PANDOC_FROM_FORMAT = "gfm+tex_math_dollars-raw_html"
EXCEL_CELL_MAX_CHARS = 32767
EXCEL_MAX_ROWS = 1_048_576
EXCEL_MAX_COLUMNS = 16_384
_BODY_SHEET_NAME = "正文"
_BODY_HEADER = ("类型", "内容")

ProgressCallback = Callable[[int, int, str], None]
CancelCheck = Callable[[], bool]


class _ConversionCancelled(Exception):
    """提交边界取消:内容已生成但尚未提交,由 convert 统一转为 cancelled 状态。"""


@dataclass
class ConversionItem:
    """单个源文件的转换结果。

    output 已写出即成功;error 记录失败原因;skipped=True 表示按规则跳过
    (如 tables 模式无表格、批量取消),不是失败。
    """

    source: Path
    output: Path | None = None
    error: str = ""
    skipped: bool = False


@dataclass
class ConversionResult:
    """批量转换结果:success = 至少一个输出且未取消。"""

    items: list[ConversionItem] = field(default_factory=list)
    cancelled: bool = False

    @property
    def success(self) -> bool:
        """至少写出一个输出文件且未取消(部分源文件失败不影响)。"""
        return not self.cancelled and any(item.output is not None for item in self.items)


@dataclass
class _Table:
    """一张 Markdown 表:表头行 + 数据行(均已按内联规则拼接为文本)。"""

    header: list[str]
    rows: list[list[str]]


class MarkdownConvertService(LoggableMixin):
    """Markdown 转换服务:Word(Pandoc)/Excel(openpyxl)批量转换。"""

    def __init__(self, history_store: JsonHistoryStore | None = None) -> None:
        """初始化服务。

        Args:
            history_store: 历史存储;传入则批量结束且至少一个输出时记录一条
                markdown_convert 历史。None 表示不记录(默认)。
        """
        self._history_store = history_store

    def convert(
        self,
        files: list[Path],
        output_dir: Path | None,
        target: str = TARGET_DOCX,
        excel_mode: str = EXCEL_MODE_TABLES,
        progress_callback: ProgressCallback | None = None,
        cancel_check: CancelCheck | None = None,
    ) -> ConversionResult:
        """把 files 逐个转换为 docx/xlsx 并写出。

        Args:
            files: 源文件列表(.md/.markdown)。
            output_dir: 输出目录;None 表示各源文件所在目录。
            target: docx 或 xlsx。
            excel_mode: tables(每表一工作表,无表跳过)或 document
                (正文工作表 + 每表一工作表);仅对 xlsx 生效。
            progress_callback: (current, total, message) 进度回调。
            cancel_check: 返回 True 时在文件间/提交前取消,已完成输出保留。

        Returns:
            ConversionResult:单文件错误记录在 item.error 后继续,不中断批量。
        """
        if target not in SUPPORTED_TARGETS:
            raise ValueError(f"无效的目标格式 {target},可选 {'/'.join(SUPPORTED_TARGETS)}")
        if excel_mode not in SUPPORTED_EXCEL_MODES:
            raise ValueError(
                f"无效的 Excel 模式 {excel_mode},可选 {'/'.join(SUPPORTED_EXCEL_MODES)}"
            )
        pandoc = _locate_pandoc() if target == TARGET_DOCX else None
        items: list[ConversionItem] = []
        cancelled = False
        total = len(files)
        for index, path in enumerate(files, start=1):
            if cancel_check is not None and cancel_check():
                cancelled = True
                self.logger.info("Markdown 转换被取消,已处理 %d/%d 个文件", index - 1, total)
                for pending in files[index - 1 :]:
                    items.append(
                        ConversionItem(source=pending, skipped=True, error="已取消,未处理")
                    )
                break
            if progress_callback is not None:
                progress_callback(index, total, f"转换 {path.name}")
            try:
                item = self._convert_one(path, output_dir, target, excel_mode, pandoc, cancel_check)
            except _ConversionCancelled:
                cancelled = True
                self.logger.info("Markdown 转换在提交前被取消: %s", path)
                items.append(ConversionItem(source=path, skipped=True, error="已取消,未写出输出"))
                for pending in files[index:]:
                    items.append(
                        ConversionItem(source=pending, skipped=True, error="已取消,未处理")
                    )
                break
            except MemoryError:
                raise
            except Exception as error:
                self.logger.warning("Markdown 转换失败,记录后继续: %s (%s)", path, error)
                items.append(ConversionItem(source=path, error=str(error) or type(error).__name__))
            else:
                items.append(item)
        result = ConversionResult(items=items, cancelled=cancelled)
        with preserve_history_result(result):
            self._record_history(result, target, excel_mode)
        return result

    # ==================== 内部实现 ====================

    def _convert_one(
        self,
        path: Path,
        output_dir: Path | None,
        target: str,
        excel_mode: str,
        pandoc: Path | None,
        cancel_check: CancelCheck | None,
    ) -> ConversionItem:
        """转换单个文件;异常上抛由 convert 记录,取消哨兵上抛由 convert 转状态。"""
        if path.suffix.lower() not in SUPPORTED_SUFFIXES:
            return ConversionItem(
                source=path,
                error=f"不支持的格式 {path.suffix or '(无后缀)'},仅支持 {'/'.join(SUPPORTED_SUFFIXES)}",
            )
        text = self._read_source(path)
        destination = (
            output_dir if output_dir is not None else path.parent
        ) / f"{path.stem}.{target}"
        if target == TARGET_DOCX:
            assert pandoc is not None  # convert 已在入口解析
            output = self._write_docx(destination, text, pandoc, cancel_check)
        else:
            blocks, tables = _parse_markdown(text)
            if excel_mode == EXCEL_MODE_TABLES and not tables:
                return ConversionItem(
                    source=path, skipped=True, error="没有表格,已跳过(tables 模式仅提取表格)"
                )
            output = self._write_xlsx(destination, blocks, tables, excel_mode, cancel_check)
        return ConversionItem(source=path, output=output)

    def _read_source(self, path: Path) -> str:
        """按 utf-8-sig 读取(BOM 兼容);原文件只读不动。"""
        try:
            return path.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError as error:
            raise ValueError(f"无法按 UTF-8 解码(不支持其他编码): {path.name}") from error
        except OSError as error:
            raise OSError(f"无法读取源文件: {error}") from error

    def _write_docx(
        self, destination: Path, text: str, pandoc: Path, cancel_check: CancelCheck | None
    ) -> Path:
        """Pandoc docx 写出:stdout 二进制直写输出流,提交前检查取消。"""

        def write(stream: IO[bytes]) -> None:
            _run_pandoc(stream, text, pandoc)
            if cancel_check is not None and cancel_check():
                raise _ConversionCancelled()

        return write_numbered_output(destination, write)

    def _write_xlsx(
        self,
        destination: Path,
        blocks: list[tuple[str, str]],
        tables: list[_Table],
        excel_mode: str,
        cancel_check: CancelCheck | None,
    ) -> Path:
        """openpyxl 工作簿写出:内容先行生成,提交前检查取消。"""
        workbook = _build_workbook(blocks, tables, excel_mode)

        def write(stream: IO[bytes]) -> None:
            workbook.save(stream)
            if cancel_check is not None and cancel_check():
                raise _ConversionCancelled()

        return write_numbered_output(destination, write)

    def _record_history(self, result: ConversionResult, target: str, excel_mode: str) -> None:
        """记录 markdown_convert 历史(键与 GUI 历史摘要 _summary_label 对齐)。

        只要实际写出过输出就记录(取消时已完成输出同样保留,历史如实标记
        cancelled);没有任何输出时不制造空记录。
        """
        if self._history_store is None:
            return
        outputs = [item.output for item in result.items if item.output is not None]
        if not outputs:
            return
        data: dict[str, object] = {
            "success": len(outputs),
            "file_count": len(result.items),
            "target": target,
            "outputs": [str(output) for output in outputs],
            "cancelled": result.cancelled,
        }
        if target == TARGET_XLSX:
            data["excel_mode"] = excel_mode
        self._history_store.add_record(HISTORY_TOOL, data)


def _locate_pandoc() -> Path:
    """定位内置 Pandoc:只用确定位置,不回退 PATH。

    源码运行取 pypandoc 包目录(files/pandoc.exe 来自 pypandoc_binary);
    打包发行取发行目录 pypandoc/files/pandoc.exe。
    """
    executable = "pandoc.exe" if sys.platform == "win32" else "pandoc"
    if is_packaged_runtime():
        candidate = Path(sys.executable).parent / "pypandoc" / "files" / executable
    else:
        spec = _find_spec("pypandoc")
        if spec is None or spec.origin is None:
            raise ImportError("Word 转换需要内置 Pandoc(pypandoc_binary),当前环境缺少 pypandoc 包")
        candidate = Path(spec.origin).parent / "files" / executable
    if not candidate.is_file():
        raise ImportError(f"未找到内置 Pandoc: {candidate}(不回退 PATH 上的 pandoc)")
    return candidate


def _run_pandoc(stream: IO[bytes], text: str, pandoc: Path) -> None:
    """运行内置 Pandoc:stdin 进 markdown,stdout docx 直写 stream。"""
    command = [
        str(pandoc),
        f"--from={PANDOC_FROM_FORMAT}",
        "--to=docx",
        "--output=-",
        "--sandbox",
        "--fail-if-warnings",
    ]
    creationflags = subprocess.CREATE_NO_WINDOW if sys.platform == "win32" else 0
    try:
        completed = subprocess.run(
            command,
            input=text.encode("utf-8"),
            stdout=stream,
            stderr=subprocess.PIPE,
            creationflags=creationflags,
            timeout=PANDOC_TIMEOUT_SECONDS,
            check=False,
        )
    except subprocess.TimeoutExpired as error:
        raise TimeoutError(f"Pandoc 超时({PANDOC_TIMEOUT_SECONDS} 秒),已终止") from error
    except OSError as error:
        raise OSError(f"无法启动内置 Pandoc: {error}") from error
    if completed.returncode != 0:
        detail = completed.stderr.decode("utf-8", "replace").strip()
        raise RuntimeError(
            f"Pandoc 转换失败(退出码 {completed.returncode}): {detail or '无诊断输出'}"
        )


def _parse_markdown(text: str) -> tuple[list[tuple[str, str]], list[_Table]]:
    """解析 markdown(commonmark + table)为 (块列表, 表格列表)。

    blocks 按阅读顺序记录 (类型, 内容);表格块的内容是 1 起始序号,
    写正文工作表时替换为工作表引用。遍历 token 流而非按行猜测:
    代码块中的伪表格、转义竖线、空单元格都由解析器语义保证。
    """
    from markdown_it import MarkdownIt

    parser = MarkdownIt("commonmark").enable("table")
    tokens = parser.parse(text)
    blocks: list[tuple[str, str]] = []
    tables: list[_Table] = []
    index = 0
    list_depth = 0
    quote_depth = 0
    while index < len(tokens):
        token = tokens[index]
        kind = token.type
        if kind == "heading_open":
            blocks.append((f"标题{token.tag[1]}", _inline_text(tokens[index + 1])))
            index += 2  # 跳过 inline;heading_close 由兜底分支越过
        elif kind in ("fence", "code_block"):
            blocks.append(("代码", token.content.rstrip("\n")))
            index += 1
        elif kind == "table_open":
            table, consumed = _parse_table(tokens, index)
            tables.append(table)
            blocks.append(("表格", str(len(tables))))
            index += consumed
        elif kind in ("bullet_list_open", "ordered_list_open"):
            list_depth += 1
            index += 1
        elif kind in ("bullet_list_close", "ordered_list_close"):
            list_depth -= 1
            index += 1
        elif kind == "blockquote_open":
            quote_depth += 1
            index += 1
        elif kind == "blockquote_close":
            quote_depth -= 1
            index += 1
        elif kind == "inline":
            if list_depth:
                label = "列表"
            elif quote_depth:
                label = "引用"
            else:
                label = "段落"
            blocks.append((label, _inline_text(token)))
            index += 1
        else:
            index += 1
    return blocks, tables


def _parse_table(tokens: list[Token], start: int) -> tuple[_Table, int]:
    """从 table_open 起解析一张表,返回 (表格, 消耗的 token 数)。"""
    header: list[str] = []
    rows: list[list[str]] = []
    row: list[str] | None = None
    in_header = False
    index = start
    while index < len(tokens):
        token = tokens[index]
        if token.type == "table_close":
            return _Table(header=header, rows=rows), index - start + 1
        if token.type == "thead_open":
            in_header = True
        elif token.type == "thead_close":
            in_header = False
        elif token.type == "tr_open":
            row = []
        elif token.type == "inline":
            if row is not None:
                row.append(_inline_text(token))
        elif token.type == "tr_close":
            if row is not None:
                if in_header:
                    header = row
                else:
                    rows.append(row)
            row = None
        index += 1
    return _Table(header=header, rows=rows), len(tokens) - start


def _inline_text(token: Token) -> str:
    """把 inline token 的 children 拼成纯文本(含行内代码与图片替代文本)。"""
    parts: list[str] = []
    for child in token.children or []:
        if child.type in ("text", "code_inline", "image"):
            parts.append(child.content)
        elif child.type in ("softbreak", "hardbreak"):
            parts.append("\n")
    return "".join(parts)


def _build_workbook(
    blocks: list[tuple[str, str]], tables: list[_Table], excel_mode: str
) -> Workbook:
    """构建输出工作簿:document 模式含正文工作表,表格各自成表。"""
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    if excel_mode == EXCEL_MODE_DOCUMENT:
        body = workbook.create_sheet(_BODY_SHEET_NAME)
        _write_row(body, 1, _BODY_HEADER)
        row_index = 2
        for label, content in blocks:
            if label == "表格":
                content = f"内容见工作表「表{content}」"
            _write_row(body, row_index, (label, content))
            row_index += 1
    for number, table in enumerate(tables, start=1):
        sheet = workbook.create_sheet(f"表{number}")
        row_index = 1
        for row_values in [table.header, *table.rows]:
            _write_row(sheet, row_index, tuple(row_values))
            row_index += 1
    return workbook


def _write_row(sheet: Worksheet, row: int, values: tuple[str, ...]) -> None:
    """写一行纯文本单元格(全字符串,防公式/保前导零,超限显式失败)。

    openpyxl 不校验行列上限(可写出 Excel 打不开的工作簿),这里显式拦截。
    """
    if row > EXCEL_MAX_ROWS:
        raise ValueError(f"行号 {row} 超过 Excel 上限 {EXCEL_MAX_ROWS},拒绝写出")
    for column, value in enumerate(values, start=1):
        if column > EXCEL_MAX_COLUMNS:
            raise ValueError(
                f"列数 {len(values)} 超过 Excel 上限 {EXCEL_MAX_COLUMNS},拒绝写出"
                "(请拆分源 Markdown 表格)"
            )
        if len(value) > EXCEL_CELL_MAX_CHARS:
            raise ValueError(
                f"单元格内容 {len(value)} 字符超过 Excel 上限 {EXCEL_CELL_MAX_CHARS},"
                "拒绝截断写出(请拆分源 Markdown)"
            )
        cell = sheet.cell(row=row, column=column)
        cell.value = value
        cell.data_type = "s"
        cell.number_format = "@"
