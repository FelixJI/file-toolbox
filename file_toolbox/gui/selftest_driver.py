"""显式 ``--selftest`` 成品自测驱动(仅 gui_entry 显式参数分支 lazy import)。

真实路径边界:
- QApplication/FreezeWatchdog/单实例守卫/prepare_gui_runtime/日志/Velopack hook
  均来自正常启动链,本模块不伪造 packaged 标志、不跳过任何 hook;
- 仅更新协调器经既有 ``MainWindow(coordinator=...)`` 测试缝注入 fake(不触发
  真实更新检查/下载/apply),其余业务页面/worker/按钮/信号全部真实;
- 模态框(QMessageBox 静态方法)记录并自动应答——等价 UI 自动化点击确认,
  业务逻辑与真实用户路径一致;
- 批量任务异步:一律以真实 ``finished`` 后 ``_task.busy`` 翻转为终态,等待用
  QEventLoop 轮询(queued 信号自然派发),不做 processEvents 自旋或 sleep 掩盖。

模式:
- pure:纯文件/图片/PDF/Markdown(含包内 Pandoc docx 与两种 xlsx 模式)、
  重命名+历史、txt/md 替换、任务取消、fake 更新互斥——CI 无 Office 可跑;
- full:pure 之外追加真实 Office 代表路径(docx/xlsx→PDF、docx 替换、考勤
  真实 Excel 虚构样本);Office 缺失或代表路径未执行 → EVIDENCE_MISSING
  (退出码 3),不得 skip 变 PASS;
- reopen:第二次启动验证同数据根历史与跨启动撤销(真实 HistoryDialog 入口)。

退出码:0=全部通过;1=存在失败;3=存在 EVIDENCE_MISSING(无失败时)。
"""

from __future__ import annotations

import json
import logging
import os
import sys
import threading
import time
from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path
from typing import TYPE_CHECKING, Any, TypedDict

from PySide6.QtCore import QEventLoop, QMetaObject, QObject, Qt, QThread, QTimer
from PySide6.QtWidgets import QApplication, QMessageBox, QWidget

from file_toolbox.common.history import JsonHistoryStore
from file_toolbox.common.paths import current_data_root
from file_toolbox.core.attendance import AttendancePlan
from file_toolbox.core.batch_pdf.engine_manager import ProbeState
from file_toolbox.core.office_capability import format_status, tool_capability_statuses
from file_toolbox.updater.coordinator import UpdateRequest
from file_toolbox.updater.models import (
    UpdateApplyResult,
    UpdateApplyStatus,
    UpdateCheckResult,
    UpdateCheckStatus,
)

if TYPE_CHECKING:
    from openpyxl.workbook.workbook import Workbook

    from file_toolbox.gui.dialogs.attendance_tab import AttendanceTab
    from file_toolbox.gui.main_window import MainWindow
    from file_toolbox.gui.task_lifecycle import TaskLifecycle

logger = logging.getLogger(__name__)

SELFTEST_MODES = ("pure", "full", "reopen")
EXIT_PASS = 0
EXIT_FAIL = 1
EXIT_EVIDENCE_MISSING = 3
# 各场景等待真实 finished 的超时(秒):纯路径快速;Office/考勤给足冷启动余量
_WAIT_FAST_S = 60
_WAIT_OFFICE_S = 300
_WAIT_ATTENDANCE_S = 600
_KIND_LABELS = {"word": "Word", "excel": "Excel", "ppt": "PowerPoint"}
# 报告工件只承载 JSON 标量/路径证据,不做通用 Any 框架
ArtifactValue = str | list[str] | dict[str, str] | bool | int
# 窗口关闭的协作收尾期限(秒):超时先写失败报告,但保留窗口/worker 与事件
# 循环继续等真实 finished/关闭,绝不以返回/sys.exit 截断 QThread 生命周期;
# 进程级超时由外层 product_selftest 负责(报告 owned PID 并保留现场,不杀)。
_CLOSE_WAIT_S = 30.0
# 取消场景的有界同步等待(秒):等待主线程真实执行按钮取消链后页面取消使
# worker 标志置位;超时即放弃等待让 worker 继续真实转换,由输出断言暴露
# 空取消(不悬挂 QThread)。
_CANCEL_SYNC_TIMEOUT_S = 5.0
# 自测数据根范围标记文件(语义标记 + 精确记录契约,非安全框架)
_SELFTEST_SCOPE_FILE = "selftest-scope.json"


# --------------------------------------------------------------------------- #
#  数据根门禁与自测范围标记(F2):拒绝已有真实数据根,不改动其任何内容
# --------------------------------------------------------------------------- #


class SelftestScope(TypedDict):
    """自测范围标记的实际语义字段(窄化解析,不落宽 Any)。"""

    owner: str
    rename_record_id: int | None


def _selftest_scope_path() -> Path:
    return current_data_root() / _SELFTEST_SCOPE_FILE


def load_selftest_scope() -> SelftestScope | None:
    """读取并语义校验自测范围标记;缺失/损坏/owner 不符/字段类型不符 → None。"""
    path = _selftest_scope_path()
    if not path.is_file():
        return None
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return None
    if not isinstance(payload, dict) or payload.get("owner") != "file-toolbox-selftest":
        return None
    record_id: object = payload.get("rename_record_id")
    if record_id is not None and (not isinstance(record_id, int) or isinstance(record_id, bool)):
        return None
    return SelftestScope(owner="file-toolbox-selftest", rename_record_id=record_id)


def _write_selftest_scope(scope: SelftestScope) -> None:
    path = _selftest_scope_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(scope, ensure_ascii=False, indent=2), encoding="utf-8")


def ensure_selftest_data_root() -> tuple[bool, str]:
    """显式 selftest 的数据根门禁:拒绝已存在的真实数据根。

    - 数据根已存在且无**语义有效**的自测范围标记(缺失/损坏/owner 不符/
      字段类型不符)→ 拒绝(绝不写入/迁移/删除任何内容);
    - 不存在或标记有效(源码隔离根/新解包便携根/同根 reopen)→ 建立或保留标记。
    标记是语义归属与精确记录 ID 契约,不是 hash/token 安全框架。
    """
    data_root = current_data_root()
    if data_root.exists():
        if load_selftest_scope() is not None:
            return True, ""
        return False, (
            f"数据根已存在且不是有效自测根(无/损坏/非自测标记),拒绝执行 --selftest"
            f"(不改动该目录): {data_root}"
        )
    _write_selftest_scope(SelftestScope(owner="file-toolbox-selftest", rename_record_id=None))
    return True, ""


def record_selftest_rename_record(record_id: int) -> None:
    """rename 场景把本轮精确记录 ID 写入范围标记,供同根 reopen 定向撤销。"""
    _write_selftest_scope(SelftestScope(owner="file-toolbox-selftest", rename_record_id=record_id))


# --------------------------------------------------------------------------- #
#  基础设施:事件循环等待、模态框自动应答、fake 更新协调器、场景结果
# --------------------------------------------------------------------------- #


def _wait_until(app: QApplication, predicate: Callable[[], bool], timeout_s: float) -> None:
    """在真实事件循环中等待谓词翻转;超时抛 TimeoutError(fail closed)。

    QEventLoop + 短周期 QTimer 轮询:queued 信号(worker→页面)在本循环内自然
    派发,不用 processEvents 自旋,也不用任意 sleep 掩盖时序。app 形参保留
    调用方语义(QApplication 需已存在),循环本身即派发通道。
    """
    del app  # 未直接使用:QEventLoop 自身派发当前线程事件
    loop = QEventLoop()
    poll = QTimer()
    poll.setInterval(10)
    deadline = QTimer()
    deadline.setSingleShot(True)
    expired: list[bool] = []

    def _poll() -> None:
        if predicate():
            loop.quit()

    def _expire() -> None:
        expired.append(True)
        loop.quit()

    poll.timeout.connect(_poll)
    deadline.timeout.connect(_expire)
    poll.start()
    deadline.start(int(timeout_s * 1000))
    loop.exec()
    poll.stop()
    deadline.stop()
    if expired:
        raise TimeoutError(f"等待条件超时({timeout_s:.0f}s)")


def _settle(app: QApplication, ms: int = 60) -> None:
    """短暂进入事件循环,让 show/布局与排队信号落地。"""
    del app
    loop = QEventLoop()
    QTimer.singleShot(ms, loop.quit)
    loop.exec()


class _MessageBoxScript:
    """记录并自动应答 QMessageBox 静态模态框(退出时恢复原实现)。

    question 的自动应答按调用方按钮集区分:提供 Apply(更新确认框)回 Apply,
    否则回 Yes(重命名执行/考勤覆盖等);information/warning/critical 回 Ok。
    只替代"人点确认"这一步,不改变任何业务分支。
    """

    def __init__(self) -> None:
        self.records: list[dict[str, str]] = []
        self._originals: dict[str, Any] = {}

    def _question_answer(self, *args: Any, **kwargs: Any) -> QMessageBox.StandardButton:
        buttons = args[3] if len(args) > 3 else kwargs.get("buttons")
        if buttons is not None:
            try:
                if int(buttons) & int(QMessageBox.StandardButton.Apply):
                    return QMessageBox.StandardButton.Apply
            except (TypeError, ValueError):
                pass
        return QMessageBox.StandardButton.Yes

    def _make(self, name: str) -> Any:
        def handler(*args: Any, **kwargs: Any) -> QMessageBox.StandardButton:
            # QMessageBox.question(parent, title, text, buttons?, default?) → 文本是第 3 个位置参数
            if len(args) > 2:
                text = str(args[2])
            else:
                text = str(kwargs.get("text", args[-1] if args else ""))
            self.records.append({"kind": name, "text": text})
            if name == "question":
                return self._question_answer(*args, **kwargs)
            return QMessageBox.StandardButton.Ok

        return handler

    def __enter__(self) -> _MessageBoxScript:
        for name in ("question", "information", "warning", "critical"):
            self._originals[name] = getattr(QMessageBox, name)
            setattr(QMessageBox, name, self._make(name))
        return self

    def __exit__(self, *exc: object) -> None:
        for name, original in self._originals.items():
            setattr(QMessageBox, name, original)
        self._originals.clear()


class SelftestUpdateCoordinator:
    """fake 更新协调器:check 结果可编程;download 阻塞并尊重真实取消门。

    ``download_and_apply`` 在 worker 线程循环检查 ``request.check_cancelled()``,
    场景经真实 UI(更新页取消按钮)取消后,UpdateCancelled 由 worker 转为
    CANCELLED 终态;``release`` 供异常收尾放行,避免悬挂 worker 线程。
    """

    def __init__(self) -> None:
        self.check_result: UpdateCheckResult = UpdateCheckResult(UpdateCheckStatus.LATEST)
        self.download_calls = 0
        self.release = threading.Event()

    def check(self) -> UpdateCheckResult:
        return self.check_result

    def download_and_apply(
        self,
        progress: Callable[[int], None] | None = None,
        *,
        request: UpdateRequest | None = None,
        before_apply: Callable[[], None] | None = None,
        expected_version: str | None = None,
    ) -> UpdateApplyResult:
        from file_toolbox.updater.coordinator import UpdateCancelled

        self.download_calls += 1
        if request is None:
            return UpdateApplyResult(UpdateApplyStatus.FAILED, "selftest: worker 未携带取消门请求")
        try:
            while not self.release.wait(0.02):
                request.check_cancelled()  # 取消被接受 → UpdateCancelled → CANCELLED
        except UpdateCancelled:
            return UpdateApplyResult(UpdateApplyStatus.CANCELLED)
        return UpdateApplyResult(UpdateApplyStatus.FAILED, "selftest: 下载门被异常释放")


@dataclass
class ScenarioOutcome:
    """单个场景的结果(pass / fail / evidence_missing)。"""

    name: str
    status: str
    detail: str = ""
    duration_s: float = 0.0
    artifacts: dict[str, ArtifactValue] = field(default_factory=dict)


@dataclass
class _SelftestContext:
    """场景共享的真实对象(窗口/工作目录/fake 协调器/模态框记录)。"""

    app: QApplication
    window: MainWindow
    workdir: Path
    coordinator: SelftestUpdateCoordinator
    boxes: _MessageBoxScript
    outcomes: list[ScenarioOutcome] = field(default_factory=list)


def _pass(
    name: str, detail: str = "", artifacts: dict[str, ArtifactValue] | None = None
) -> ScenarioOutcome:
    return ScenarioOutcome(name, "pass", detail, artifacts=dict(artifacts or {}))


def _fail(name: str, detail: str) -> ScenarioOutcome:
    return ScenarioOutcome(name, "fail", detail)


def _missing(name: str, detail: str) -> ScenarioOutcome:
    return ScenarioOutcome(name, "evidence_missing", detail)


def _switch_to_tab[TabT: QWidget](
    ctx: _SelftestContext, tool_id: str, tab_type: type[TabT]
) -> TabT:
    """按稳定 ID 切换到登记页并返回已构造页面(按具体类型窄化)。"""
    index = next((i for i, spec in enumerate(ctx.window._specs) if spec.tool_id == tool_id), -1)
    if index < 0:
        raise RuntimeError(f"登记中不存在工具: {tool_id}")
    ctx.window._tabs.setCurrentIndex(index)
    spec = ctx.window._specs[index]
    widget = getattr(ctx.window, spec.attr, None)
    if widget is None:
        raise RuntimeError(f"页面未构造: {tool_id}")
    if not isinstance(widget, tab_type):
        raise RuntimeError(f"页面类型不符: {tool_id} → {type(widget).__name__}")
    return widget


def _office_ready(tool_id: str, kinds: tuple[str, ...]) -> tuple[bool, str]:
    """full 模式前置检查:指定 kind 是否预筛可用。

    预筛不可用(缺失或检测失败)按 EVIDENCE_MISSING 处理,不冒充可运行;
    预筛可用也不代表转换必然成功——真实 Dispatch 结果才是运行证据。
    """
    wanted = {_KIND_LABELS.get(kind, kind) for kind in kinds}
    problems = [
        format_status(status)
        for status in tool_capability_statuses(tool_id)
        if status.requirement.split("(", 1)[0] in wanted
        and status.state is not ProbeState.AVAILABLE
    ]
    return (not problems), "; ".join(problems)


def _collect_worker_problems(worker: QObject | None, sink: list[str]) -> None:
    """把真实 worker 的失败/清理警告信号订阅到共享 sink(F6)。

    真实 finished 不等于资源清理成功:结果输出成功但严格清理失败(cleanup_
    warning)或失败信号都必须让场景非零,不靠模态框文案猜测。回调直写调用方
    持有的 sink——不能返回新列表(调用方 extend/+= 会在订阅时复制空列表,
    异步写入落在被弃引用上,问题永远丢失)。
    """

    if worker is None:
        return
    for signal_name, label in (
        ("failed", "failed"),
        ("cleanup_warning", "cleanup_warning"),
        ("warning", "warning"),
    ):
        signal = getattr(worker, signal_name, None)
        if signal is None:
            continue

        def record(message: object, lbl: str = label) -> None:
            sink.append(f"{lbl}: {message}")

        signal.connect(record)


def _observe_task_problems(
    lifecycle: TaskLifecycle,
    sink: list[str],
    extra: Callable[[QThread], None] | None = None,
) -> None:
    """经 TaskLifecycle.track 的 start 前统一注册点订阅问题信号(F6)。

    track 在 worker.start() 之前执行,订阅严格先于任何发射——工作线程可能在
    主线程连接前立即失败/立即清理告警,启动后再连接会丢这些信号;"start 后
    立即连接必先于发射"的时序假设不成立,必须经 start 前注册点订阅。
    extra:同一时刻的附加观察(如 start 前连接进度),与本观察链组合,不另建
    注册框架。
    """

    def observe(worker: QThread) -> None:
        _collect_worker_problems(worker, sink)
        if extra is not None:
            extra(worker)

    lifecycle.on_worker_tracked = observe


# --------------------------------------------------------------------------- #
#  场景:pure(纯文件/图片/PDF/Markdown/重命名/替换/取消/更新互斥)
# --------------------------------------------------------------------------- #


def _scenario_pages(ctx: _SelftestContext) -> ScenarioOutcome:
    """切页:全部登记页懒构造成功,无页面加载错误。"""
    window = ctx.window
    for index in range(window._tabs.count()):
        window._tabs.setCurrentIndex(index)
        _settle(ctx.app, 10)
        if window._tab_error_message is not None:
            return _fail("pages", f"切到第 {index} 页失败: {window._tab_error_message}")
    unfilled = [spec.tool_id for spec in window._specs if getattr(window, spec.attr, None) is None]
    if unfilled:
        return _fail("pages", f"页面未构造: {unfilled}")
    return _pass("pages", artifacts={"tabs": window._tabs.count()})


def _scenario_rename(ctx: _SelftestContext) -> ScenarioOutcome:
    """重命名:预览→执行(真实 worker)→读回文件与历史记录。"""
    from file_toolbox.gui.dialogs.rename_tab import FileRenamerDialog

    work = ctx.workdir / "rename"
    work.mkdir(parents=True, exist_ok=True)
    names = [f"selftest-{i}.txt" for i in range(3)]
    for name in names:
        (work / name).write_text(f"content of {name}\n", encoding="utf-8")
    tab = _switch_to_tab(ctx, "rename", FileRenamerDialog)
    problems: list[str] = []
    _observe_task_problems(tab._task, problems)  # track 时订阅:严格先于 start
    tab.selected_files = [work / name for name in names]
    tab.operations = [
        {"type": "replace_text", "params": {"find": "selftest", "replace": "ftb-renamed"}}
    ]
    tab._do_refresh_preview()
    _wait_until(
        ctx.app, lambda: not tab._task.busy and tab._preview_snapshot is not None, _WAIT_FAST_S
    )
    tab._execute()  # 确认框自动 Yes
    _wait_until(ctx.app, lambda: not tab._task.busy, _WAIT_FAST_S)
    if problems:
        return _fail("rename", "; ".join(problems))
    expected = sorted(name.replace("selftest", "ftb-renamed") for name in names)
    actual = sorted(path.name for path in work.iterdir() if path.is_file())
    if actual != expected:
        return _fail("rename", f"重命名结果不符: 期望 {expected}, 实际 {actual}")
    records = tab._history.get_records("rename", limit=1)
    if not records:
        return _fail("rename", "rename 历史未记录")
    record_id = int(records[0]["id"])
    record_selftest_rename_record(record_id)  # 范围标记精确记录 ID,供同根 reopen
    data = records[0].get("data", {})
    raw_map = data.get("rename_map", {}) if isinstance(data, dict) else {}
    if not isinstance(raw_map, dict) or len(raw_map) != len(names):
        return _fail("rename", f"历史 rename_map 不符: {raw_map}")
    rename_map = {str(key): str(value) for key, value in raw_map.items()}
    return _pass(
        "rename",
        artifacts={
            "files": actual,
            "rename_map": rename_map,
            "record_id": record_id,
            "workdir": str(work),
            "history_tool": "rename",
        },
    )


def _scenario_pdf_pure(ctx: _SelftestContext) -> ScenarioOutcome:
    """纯 PDF 路径:图片→PDF(真实转换)与 PDF→PDF 复制,不启动 Office。"""
    from PIL import Image

    from file_toolbox.gui.dialogs.pdf_tab import PDFGeneratorDialog

    work = ctx.workdir / "pdf-pure"
    work.mkdir(parents=True, exist_ok=True)
    images = []
    for index in range(3):
        path = work / f"pure-image-{index}.png"
        Image.new("RGB", (800, 600), (30 * index + 20, 90, 150)).save(path)
        images.append(path)
    tab = _switch_to_tab(ctx, "pdf", PDFGeneratorDialog)
    problems: list[str] = []
    _observe_task_problems(tab._task, problems)  # track 时订阅:严格先于 start
    tab.selected_files = images
    tab._generate()
    _wait_until(ctx.app, lambda: not tab._task.busy, _WAIT_FAST_S)
    if problems:
        return _fail("pdf_pure", "; ".join(problems))  # 信号证据优先于输出推断
    outputs = [work / f"{path.stem}.pdf" for path in images]
    for source, output in zip(images, outputs, strict=True):
        if not output.is_file():
            return _fail("pdf_pure", f"缺少输出: {output}(源 {source.name})")
    # pdf→pdf(可编辑型=复制)路径
    tab.selected_files = [outputs[0]]
    tab._generate()
    _wait_until(ctx.app, lambda: not tab._task.busy, _WAIT_FAST_S)
    if problems:
        return _fail("pdf_pure", "; ".join(problems))
    copy_output = work / f"{outputs[0].stem}_1.pdf"
    if not copy_output.is_file():
        return _fail("pdf_pure", f"pdf→pdf 复制缺少输出: {copy_output}")
    if problems:
        return _fail("pdf_pure", "; ".join(problems))
    return _pass(
        "pdf_pure",
        artifacts={"pdfs": [str(p) for p in outputs], "pdf_copy": str(copy_output)},
    )


def _scenario_pdf_cancel(ctx: _SelftestContext) -> ScenarioOutcome:
    """任务取消真实路径:真实按钮→页面取消→worker 标志,确定性生效。

    可控同步点(本自测限定,无泛化框架):progress 以 DirectConnection 在
    worker 线程内同步处理——首个进度(每个文件开始前发射)把真实取消按钮的
    click 以 QueuedConnection 投递到主线程执行(不在后台线程触碰 QWidget),
    并有界等待页面取消链使 worker 标志置位后才放行;下一个文件边界的取消
    检查因此确定命中,不依赖线程调度/样本量/延时。若取消链任一环退化(空
    cancel/按钮无效),等待超时后 worker 继续真实转换全部文件,输出断言
    使场景失败——不会悬挂 QThread。

    fail closed:PASS 必须同时满足(1)实际观察到取消标志置位(空取消+
    部分文件转换失败产生的伪部分输出不得误判为取消成功);(2)至少一个输出
    且至少一个输入未处理。
    """
    from PIL import Image

    from file_toolbox.gui.dialogs.pdf_tab import PDFGeneratorDialog
    from file_toolbox.gui.workers.pdf_worker import PdfGenerateWorker

    work = ctx.workdir / "pdf-cancel"
    work.mkdir(parents=True, exist_ok=True)
    images = []
    for index in range(3):
        path = work / f"cancel-image-{index}.png"
        Image.new("RGB", (1600, 1200), (40 * index + 15, 70, 110)).save(path)
        images.append(path)
    tab = _switch_to_tab(ctx, "pdf", PDFGeneratorDialog)
    problems: list[str] = []
    requested: list[int] = []
    cancel_observed = False
    worker_ref: PdfGenerateWorker | None = None

    def on_progress(current: int, total: int, message: str) -> None:  # noqa: ARG001
        """worker 线程内同步执行(有界):投递真实按钮点击并等页面取消链生效。"""
        nonlocal cancel_observed
        if requested:
            return
        requested.append(current)
        target = worker_ref
        if target is None:
            return
        # 真实 UI 路径:click 槽在按钮所属主线程执行(→ _on_cancel → _task.cancel)
        QMetaObject.invokeMethod(tab.ui.btn_cancel, "click", Qt.ConnectionType.QueuedConnection)
        deadline = time.monotonic() + _CANCEL_SYNC_TIMEOUT_S
        while not target._cancel and time.monotonic() < deadline:
            time.sleep(0.005)
        cancel_observed = target._cancel

    def observe_extra(worker: QThread) -> None:
        nonlocal worker_ref
        if isinstance(worker, PdfGenerateWorker):
            worker_ref = worker
            worker.progress.connect(on_progress, Qt.ConnectionType.DirectConnection)

    _observe_task_problems(tab._task, problems, extra=observe_extra)  # track 时订阅:严格先于 start
    tab.selected_files = images
    tab._generate()
    _wait_until(ctx.app, lambda: not tab._task.busy, _WAIT_FAST_S)
    if not requested:
        return _fail("pdf_cancel", "未收到首个进度(start 前订阅失效或 worker 未发射)")
    outputs = sorted(path.name for path in work.glob("cancel-image-*.pdf"))
    if not outputs:
        return _fail("pdf_cancel", "取消场景无任何输出(首个文件应已完成)")
    if len(outputs) >= len(images):
        return _fail(
            "pdf_cancel",
            f"取消未在剩余文件边界生效:产出 {len(outputs)}/{len(images)},按钮→页面→worker"
            f"取消链未生效(等待标志置位: {cancel_observed})",
        )
    if not cancel_observed:
        return _fail(
            "pdf_cancel",
            "取消链未在有界等待内置位 worker 取消标志(按钮→页面→worker 任一环失效);"
            "输出计数达标不构成取消成功",
        )
    if not tab.ui.btn_generate.isEnabled():
        return _fail("pdf_cancel", "真实 finished 后控件未恢复")
    if problems:
        return _fail("pdf_cancel", "; ".join(problems))
    return _pass(
        "pdf_cancel",
        detail=(
            f"取消于文件 {requested[0]} 进度处经真实按钮链同步生效"
            f"(标志置位: {cancel_observed}),产出 {len(outputs)}/{len(images)}"
        ),
        artifacts={"outputs": outputs},
    )


_MARKDOWN_BODY = """# 自测标题

正文段落 selftest-body,用于 docx 读回验证。

| 名称 | 数量 |
| --- | --- |
| 甲 | 1 |
| 乙 | 2 |
"""


def _scenario_markdown(ctx: _SelftestContext) -> ScenarioOutcome:
    """Markdown:docx(包内 Pandoc)与两种 xlsx 模式(纯库),逐一读回。"""
    from file_toolbox.gui.dialogs.markdown_tab import MarkdownConvertTab

    work = ctx.workdir / "markdown"
    work.mkdir(parents=True, exist_ok=True)
    tab = _switch_to_tab(ctx, "markdown_convert", MarkdownConvertTab)
    rounds = (("docx", None), ("xlsx", "tables"), ("xlsx", "document"))
    outputs: dict[str, Path] = {}
    problems: list[str] = []
    _observe_task_problems(tab._task, problems)  # track 时订阅:严格先于 start
    for target, mode in rounds:
        label = target if mode is None else mode
        subdir = work / label
        subdir.mkdir(parents=True, exist_ok=True)
        source = subdir / f"sample-{label}.md"
        source.write_text(_MARKDOWN_BODY, encoding="utf-8")
        tab.ui.cmb_target.setCurrentIndex(0 if target == "docx" else 1)
        if mode is not None:
            tab.ui.cmb_excel_mode.setCurrentIndex(0 if mode == "tables" else 1)
        tab.ui.edit_outdir.setText(str(subdir))
        tab._add_paths([source])
        _wait_until(ctx.app, lambda: len(tab._files) == 1 and not tab._task.busy, _WAIT_FAST_S)
        tab._convert()
        _wait_until(ctx.app, lambda: not tab._task.busy, _WAIT_FAST_S)
        output = subdir / f"sample-{label}.{target}"
        if not output.is_file():
            return _fail("markdown", f"{label} 输出缺失: {output}")
        outputs[label] = output
        tab._clear()
    # 就地读回关键内容(外层脚本再独立复核)
    from openpyxl import load_workbook

    tables_book = load_workbook(outputs["tables"], read_only=True)
    try:
        if "表1" not in tables_book.sheetnames or tables_book["表1"]["A1"].value != "名称":
            return _fail("markdown", f"tables 工作簿结构不符: {tables_book.sheetnames}")
    finally:
        tables_book.close()
    document_book = load_workbook(outputs["document"], read_only=True)
    try:
        if "正文" not in document_book.sheetnames:
            return _fail("markdown", f"document 工作簿缺正文表: {document_book.sheetnames}")
    finally:
        document_book.close()
    import zipfile

    with zipfile.ZipFile(outputs["docx"]) as package:
        xml = package.read("word/document.xml").decode("utf-8")
    if "自测标题" not in xml:
        return _fail("markdown", "docx document.xml 缺少标题文本")
    if problems:
        return _fail("markdown", "; ".join(problems))
    return _pass(
        "markdown",
        artifacts={
            **{f"{label}_output": str(path) for label, path in outputs.items()},
            **{f"{label}_source": str(path.with_suffix(".md")) for label, path in outputs.items()},
        },
    )


def _scenario_replace_pure(ctx: _SelftestContext) -> ScenarioOutcome:
    """纯替换:txt/md 预览→执行(真实 worker)→读回文件内容。"""
    from file_toolbox.gui.dialogs.replace_tab import ContentReplaceDialog

    work = ctx.workdir / "replace-pure"
    work.mkdir(parents=True, exist_ok=True)
    txt = work / "notes.txt"
    txt.write_text("hello selftest hello\n", encoding="utf-8")
    md = work / "notes.md"
    md.write_text("hello selftest hello\n", encoding="utf-8")
    tab = _switch_to_tab(ctx, "replace", ContentReplaceDialog)
    problems: list[str] = []
    _observe_task_problems(tab._task, problems)  # track 时订阅:严格先于 start
    tab.selected_files = [txt, md]
    tab.operations = [{"type": "simple_replace", "params": {"find": "hello", "replace": "nihao"}}]
    tab._do_refresh_preview()
    _wait_until(ctx.app, lambda: not tab._task.busy, _WAIT_FAST_S)
    tab._execute()  # 确认框自动 Yes
    _wait_until(ctx.app, lambda: not tab._task.busy, _WAIT_FAST_S)
    if problems:
        return _fail("replace_pure", "; ".join(problems))
    for path in (txt, md):
        if "nihao selftest nihao" not in path.read_text(encoding="utf-8"):
            return _fail("replace_pure", f"替换后内容不符: {path}")
    return _pass("replace_pure", artifacts={"txt": str(txt), "md": str(md)})


def _scenario_update_mutex(ctx: _SelftestContext) -> ScenarioOutcome:
    """fake 更新互斥:真实检查→横幅→更新页→下载期间业务锁定→取消→解锁。"""
    window = ctx.window
    coordinator = ctx.coordinator
    coordinator.check_result = UpdateCheckResult(
        UpdateCheckStatus.AVAILABLE, version="9.9.9-selftest"
    )
    worker = window._update_worker
    if not worker.isRunning():
        worker.start()
    QMetaObject.invokeMethod(worker, "do_check", Qt.ConnectionType.QueuedConnection)
    _wait_until(
        ctx.app,
        lambda: (
            window._last_check is not None
            and window._last_check.status is UpdateCheckStatus.AVAILABLE
        ),
        _WAIT_FAST_S,
    )
    window._update_banner.click()  # 横幅 → 打开更新页(仅导航,不下载)
    _wait_until(ctx.app, lambda: window._update_tab is not None, _WAIT_FAST_S)
    window._start_download()  # 确认框自动 Apply
    _wait_until(ctx.app, lambda: coordinator.download_calls == 1, _WAIT_FAST_S)
    rename_tab = getattr(window, "_rename_tab", None)
    if rename_tab is None:
        return _fail("update_mutex", "rename 页未构造,无法验证业务锁")
    if rename_tab.isEnabled():
        return _fail("update_mutex", "下载期间业务页未锁定")
    update_tab = window._update_tab
    if update_tab is None or not update_tab.isEnabled():
        return _fail("update_mutex", "下载期间更新页不可用")
    if update_tab.btn_cancel_update.isHidden():
        return _fail("update_mutex", "下载期间取消按钮不可见")
    update_tab.btn_cancel_update.click()  # 真实取消入口
    _wait_until(ctx.app, lambda: window._download_request is None, _WAIT_FAST_S)
    if not rename_tab.isEnabled():
        return _fail("update_mutex", "取消后业务页未解锁")
    if coordinator.download_calls != 1:
        return _fail("update_mutex", "重复下载请求")
    return _pass(
        "update_mutex",
        artifacts={
            "download_calls": coordinator.download_calls,
            "locked_then_unlocked": True,
        },
    )


# --------------------------------------------------------------------------- #
#  场景:full(真实 Office 代表路径)
# --------------------------------------------------------------------------- #


def _markdown_artifact(ctx: _SelftestContext, key: str) -> Path | None:
    """取 markdown 场景已产出的输入文件(供 Office 代表路径复用)。"""
    markdown = next((outcome for outcome in ctx.outcomes if outcome.name == "markdown"), None)
    if markdown is None:
        return None
    value = markdown.artifacts.get(key)
    return Path(value) if isinstance(value, str) else None


def _scenario_pdf_office(ctx: _SelftestContext) -> ScenarioOutcome:
    """Office→PDF:docx(Word)与 xlsx(Excel)真实转换并读回页数。"""
    ready, missing = _office_ready("pdf", ("word", "excel"))
    if not ready:
        return _missing("pdf_office", f"Office 预筛不可用: {missing}")
    docx_source = _markdown_artifact(ctx, "docx_output")
    xlsx_source = _markdown_artifact(ctx, "tables_output")
    if docx_source is None or xlsx_source is None:
        return _fail("pdf_office", "缺少 markdown 输出作为 Office 转换输入(前置场景未通过)")
    work = ctx.workdir / "pdf-office"
    work.mkdir(parents=True, exist_ok=True)
    from file_toolbox.gui.dialogs.pdf_tab import PDFGeneratorDialog

    tab = _switch_to_tab(ctx, "pdf", PDFGeneratorDialog)
    outputs: dict[str, ArtifactValue] = {}
    problems: list[str] = []
    _observe_task_problems(tab._task, problems)  # track 时订阅:严格先于 start
    for label, source in (("docx", docx_source), ("xlsx", xlsx_source)):
        target = work / f"office-{label}-source{source.suffix}"
        target.write_bytes(source.read_bytes())
        tab.selected_files = [target]
        tab._generate()
        _wait_until(ctx.app, lambda: not tab._task.busy, _WAIT_OFFICE_S)
        output = work / f"{target.stem}.pdf"
        if not output.is_file():
            return _fail("pdf_office", f"{label}→PDF 缺少输出: {output}")
        outputs[label] = str(output)
    from pypdf import PdfReader

    for label, path in outputs.items():
        if len(PdfReader(str(path)).pages) < 1:
            return _fail("pdf_office", f"{label}→PDF 页数异常")
    if problems:
        return _fail("pdf_office", "; ".join(problems))
    return _pass("pdf_office", artifacts=outputs)


def _scenario_replace_office(ctx: _SelftestContext) -> ScenarioOutcome:
    """docx 内容替换(真实 Word COM):预览→执行→读回 document.xml。"""
    ready, missing = _office_ready("replace", ("word",))
    if not ready:
        return _missing("replace_office", f"Office 预筛不可用: {missing}")
    docx_source = _markdown_artifact(ctx, "docx_output")
    if docx_source is None:
        return _fail("replace_office", "缺少 markdown 输出作为替换输入(前置场景未通过)")
    work = ctx.workdir / "replace-office"
    work.mkdir(parents=True, exist_ok=True)
    docx = work / "office-replace.docx"
    docx.write_bytes(docx_source.read_bytes())
    from file_toolbox.gui.dialogs.replace_tab import ContentReplaceDialog

    tab = _switch_to_tab(ctx, "replace", ContentReplaceDialog)
    problems: list[str] = []
    _observe_task_problems(tab._task, problems)  # track 时订阅:严格先于 start
    tab.selected_files = [docx]
    tab.operations = [
        {
            "type": "simple_replace",
            "params": {"find": "selftest-body", "replace": "office-replaced"},
        }
    ]
    tab._do_refresh_preview()
    _wait_until(ctx.app, lambda: not tab._task.busy, _WAIT_OFFICE_S)
    tab._execute()  # 确认框自动 Yes(自动备份真实发生)
    _wait_until(ctx.app, lambda: not tab._task.busy, _WAIT_OFFICE_S)
    if problems:
        return _fail("replace_office", "; ".join(problems))
    import zipfile

    with zipfile.ZipFile(docx) as package:
        xml = package.read("word/document.xml").decode("utf-8")
    if "office-replaced" not in xml or "selftest-body" in xml:
        return _fail("replace_office", "docx 替换读回不符")
    return _pass("replace_office", artifacts={"docx": str(docx)})


def _build_attendance_samples(workdir: Path) -> dict[str, Path]:
    """生成最小正确虚构考勤样本(openpyxl,不依赖任何 TEMP 外部路径)。

    前序验证修正过的坑必须保持:
    - 源考勤部门列统一为同一部门(普通模式多部门不支持);
    - 模板明细第 6 行日期表头写字符串 '1'..'30'(Excel 读数字成 1.0 会匹配失败);
    - 名单模式源内姓名唯一(测试丙/丁/戊),普通模式才允许跨组同名。
    """
    from openpyxl import Workbook

    workdir.mkdir(parents=True, exist_ok=True)
    department = "验收部门"

    def write_source(path: Path, people: list[tuple[str, str]]) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sheet1"
        for row, (name, group) in enumerate(people, start=2):
            sheet.cell(row, 1, name)
            sheet.cell(row, 2, group)
            sheet.cell(row, 3, department)
            for day in range(30):
                value = "迟到" if (row == 2 and day == 2) else "正常"
                sheet.cell(row, 7 + day, value)
        workbook.save(path)
        workbook.close()

    write_source(
        workdir / "source.xlsx",
        [("测试甲", "甲组"), ("同名", "甲组"), ("同名", "乙组"), ("测试乙", "乙组")],
    )
    write_source(
        workdir / "source-roster.xlsx",
        [("测试丙", "丙组"), ("测试丁", "丁组"), ("测试戊", "戊组")],
    )

    template = workdir / "template.xlsx"
    workbook = Workbook()
    detail = workbook.active
    detail.title = "出勤明细"
    summary = workbook.create_sheet("考勤汇总表")
    for day in range(1, 31):
        detail.cell(6, 3 + day, str(day))  # 字符串表头(见 docstring 第 2 点)
    for row in range(7, 22):
        detail.cell(row, 3, "模板占位")
        summary.cell(row + 1, 3, "模板占位")
        summary.cell(row + 1, 5, f"=COUNTIF('出勤明细'!D{row}:AG{row},\"√\")")
    workbook.save(template)
    workbook.close()

    roster = workdir / "roster.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    people = [
        ("丙组", department, "测试丙", "900001"),
        ("丁组", department, "测试丁", "900002"),
        ("戊组", department, "测试戊", "900003"),
    ]
    for row, values in enumerate(people, start=1):
        for column, value in enumerate(values, start=1):
            sheet.cell(row, column, value)
    workbook.save(roster)
    workbook.close()
    return {
        "source": workdir / "source.xlsx",
        "source_roster": workdir / "source-roster.xlsx",
        "template": template,
        "roster": roster,
    }


def _attendance_plan(template: Path, *, roster_mode: bool) -> AttendancePlan:
    """构造最小正确方案(普通/名单两用,布局与模板严格对应)。"""
    from dataclasses import replace

    from file_toolbox.core.attendance import (
        CellMapping,
        CellRef,
        GroupSheetConfig,
        RosterConfig,
        RosterLayout,
        SourceLayout,
        TargetLayout,
    )

    groups = (
        (("甲组", "甲明细", "甲汇总"), ("乙组", "乙明细", "乙汇总"))
        if not roster_mode
        else (
            ("丙组", "丙明细", "丙汇总"),
            ("丁组", "丁明细", "丁汇总"),
            ("戊组", "戊明细", "戊汇总"),
        )
    )
    plan = AttendancePlan(
        "自测方案",
        template,
        SourceLayout("Sheet1", CellRef(2, 1), CellRef(2, 3), CellRef(2, 7), CellRef(2, 2)),
        TargetLayout("出勤明细", CellRef(7, 3), CellRef(7, 4), "考勤汇总表", CellRef(8, 3)),
        mappings=(CellMapping("出勤明细", CellRef(3, 1), "{{year}}年{{month}}月"),),
        split_by_group=True,
        group_sheet_configs=tuple(
            GroupSheetConfig(group, detail, summary, f"{group}别名")
            for group, detail, summary in groups
        ),
    )
    if roster_mode:
        return replace(
            plan,
            roster=RosterConfig(
                template.parent / "roster.xlsx",
                RosterLayout("Sheet1", CellRef(1, 1), CellRef(1, 2), CellRef(1, 3), CellRef(1, 4)),
            ),
        )
    return plan


def _round_wait(
    ctx: _SelftestContext, tab: AttendanceTab, timeout_s: float, sink: list[str]
) -> bool:
    """为当前 worker 订阅问题信号到共享 sink 并等待真实 finished。

    返回 worker 是否存在(极端竞态下无 worker 由调用方如实失败)。
    """
    worker = tab._task.worker
    if worker is not None:
        _collect_worker_problems(worker, sink)
    _wait_until(ctx.app, lambda: not tab._task.busy, timeout_s)
    return worker is not None


def _run_attendance_round(
    ctx: _SelftestContext,
    tab: AttendanceTab,
    plan: AttendancePlan,
    source: Path,
    output: Path,
    *,
    roster_mode: bool,
    problems: list[str],
) -> str:
    """一轮考勤:预览→(调整)→重新预览→生成,返回错误描述(空=成功)。

    普通模式做人员调组(与编辑委托同源的 model.setData);名单模式经真实预览
    勾选排除测试戊(工号 900003,与计划层排除等效但走真实 UI 路径);两者都
    经过真实预览门与生成 worker,按真实 finished 等待。
    """
    tab._apply_plan(plan)
    tab.ui.edit_source.setText(str(source))
    tab.ui.edit_output_dir.setText(str(output.parent))
    tab.ui.edit_output_name.setText(output.name)
    tab.ui.spin_year.setValue(2026)
    tab.ui.spin_month.setValue(9)
    tab._preview()
    _round_wait(ctx, tab, _WAIT_ATTENDANCE_S, problems)
    if problems:
        return "; ".join(problems)
    if tab._preview_request is None:
        return "预览未通过"
    if not tab.ui.btn_generate.isEnabled():
        return "预览通过但生成按钮未启用"
    view = tab.ui.table_employee_preview
    model = view.model()
    if roster_mode:
        rows = [
            row for row in range(model.rowCount()) if model.data(model.index(row, 1)) == "900003"
        ]
        if len(rows) != 1:
            return f"名单预览缺少工号 900003 行: {rows}"
        if not model.setData(
            model.index(rows[0], 0), Qt.CheckState.Unchecked, Qt.ItemDataRole.CheckStateRole
        ):
            return "名单排除勾选失败"
    else:
        rows = [
            row
            for row in range(model.rowCount())
            if model.data(model.index(row, 0)) == "测试甲"
            and model.data(model.index(row, 1)) == "甲组"
        ]
        if len(rows) != 1:
            return f"普通预览缺少测试甲/甲组行: {rows}"
        if not model.setData(model.index(rows[0], 2), "乙组"):
            return "人员调组编辑失败"
    if tab.ui.btn_generate.isEnabled():
        return "调整后未要求重新预览"
    tab._apply_preview_adjustments()
    _round_wait(ctx, tab, _WAIT_ATTENDANCE_S, problems)
    if problems:
        return "; ".join(problems)
    if tab._preview_request is None:
        return "调整后预览未通过"
    tab._generate()
    _round_wait(ctx, tab, _WAIT_ATTENDANCE_S, problems)
    if problems:
        return "; ".join(problems)
    if not output.is_file():
        return f"结果缺失: {output}"
    return ""


def _detail_names(book: Workbook, sheet_name: str) -> list[str]:
    """读明细表姓名列(第 3 列、第 7 行起)的已有值。"""
    return [
        str(row[0])
        for row in book[sheet_name].iter_rows(
            min_row=7, max_row=25, min_col=3, max_col=3, values_only=True
        )
        if row[0]
    ]


def _scenario_attendance(ctx: _SelftestContext) -> ScenarioOutcome:
    """考勤(真实 Excel):普通(分组+调组)与名单(排除)两轮,读回工作簿。"""
    ready, missing = _office_ready("attendance", ("excel",))
    if not ready:
        return _missing("attendance", f"Office 预筛不可用: {missing}")
    work = ctx.workdir / "attendance"
    samples = _build_attendance_samples(work)
    from file_toolbox.gui.dialogs.attendance_tab import AttendanceTab as _AttendanceTab

    tab = _switch_to_tab(ctx, "attendance", _AttendanceTab)
    problems: list[str] = []
    _observe_task_problems(tab._task, problems)  # track 时订阅:严格先于 start
    artifacts: dict[str, ArtifactValue] = {key: str(value) for key, value in samples.items()}
    plain_output = work / "plain-output.xlsx"
    error = _run_attendance_round(
        ctx,
        tab,
        _attendance_plan(samples["template"], roster_mode=False),
        samples["source"],
        plain_output,
        roster_mode=False,
        problems=problems,
    )
    if error:
        return _fail("attendance", f"普通模式: {error}")
    artifacts["plain_output"] = str(plain_output)
    roster_output = work / "roster-output.xlsx"
    error = _run_attendance_round(
        ctx,
        tab,
        _attendance_plan(samples["template"], roster_mode=True),
        samples["source_roster"],
        roster_output,
        roster_mode=True,
        problems=problems,
    )
    if error:
        return _fail("attendance", f"名单模式: {error}")
    artifacts["roster_output"] = str(roster_output)
    from openpyxl import load_workbook

    plain = load_workbook(plain_output, read_only=True)
    try:
        for sheet in ("甲明细", "乙明细", "乙汇总"):
            if sheet not in plain.sheetnames:
                return _fail("attendance", f"普通模式输出缺少工作表 {sheet}: {plain.sheetnames}")
        names_jia = _detail_names(plain, "甲明细")
        names_yi = _detail_names(plain, "乙明细")
        if "测试甲" in names_jia:
            return _fail("attendance", "测试甲 调组后仍留在 甲明细")
        if "测试甲" not in names_yi:
            return _fail("attendance", f"测试甲 调组后未进入 乙明细: {names_yi}")
    finally:
        plain.close()
    roster_book = load_workbook(roster_output, read_only=True)
    try:
        if "丙明细" not in roster_book.sheetnames or "丁明细" not in roster_book.sheetnames:
            return _fail("attendance", f"名单模式输出缺工作表: {roster_book.sheetnames}")
        roster_names = [
            name
            for sheet in roster_book.sheetnames
            if sheet.endswith("明细")
            for name in _detail_names(roster_book, sheet)
        ]
        if "测试丙" not in roster_names or "测试丁" not in roster_names:
            return _fail("attendance", f"名单模式人员缺失: {roster_names}")
        if "测试戊" in roster_names:
            return _fail("attendance", "被排除的测试戊仍出现在明细输出")
    finally:
        roster_book.close()
    return _pass("attendance", artifacts=artifacts)


# --------------------------------------------------------------------------- #
#  场景:reopen(第二次启动:同数据根历史 + 跨启动撤销)
# --------------------------------------------------------------------------- #


def _scenario_reopen_history_undo(ctx: _SelftestContext) -> ScenarioOutcome:
    """第二次启动:同数据根历史仍在,按范围标记的精确记录经真实 HistoryDialog 撤销。

    只接受首轮写入范围标记的 rename 记录 ID,且其 rename_map 必须全部落在
    固定虚构目录 selftest-work/rename 内——不从任意未 undone 历史选一条,
    也不碰真实用户历史。
    """
    work = ctx.workdir / "rename"
    scope = load_selftest_scope()
    record_id = scope["rename_record_id"] if scope is not None else None
    if record_id is None:
        return _fail(
            "reopen_history_undo", "自测范围标记缺少首轮 rename 记录 ID(非自测根或首轮未执行)"
        )
    store = JsonHistoryStore()  # 同数据根策略 → 读到第一次启动的历史
    record = store.get_record("rename", record_id)
    if record is None:
        return _fail("reopen_history_undo", f"范围标记指向的记录不存在: id={record_id}")
    if record.get("undone"):
        return _fail("reopen_history_undo", f"目标记录已被撤销: id={record_id}")
    data = record.get("data", {})
    raw_map = data.get("rename_map", {}) if isinstance(data, dict) else {}
    rename_map = (
        {str(key): str(value) for key, value in raw_map.items()}
        if isinstance(raw_map, dict) and raw_map
        else {}
    )
    if not rename_map:
        return _fail("reopen_history_undo", "记录缺少 rename_map")
    for old, new in rename_map.items():
        if Path(old).parent != work or Path(new).parent != work:
            return _fail("reopen_history_undo", f"记录路径不在自测虚构目录内: {old} → {new}")
        if not Path(new).is_file():
            return _fail("reopen_history_undo", f"重命名后的文件缺失: {new}")
    from file_toolbox.gui.dialogs.history_dialog import HistoryDialog

    dialog = HistoryDialog(store, "rename")
    target_row = next(
        (
            row
            for row in range(dialog.list_widget.count())
            if dialog.list_widget.item(row).data(0x0100) == record["id"]
        ),
        -1,
    )
    if target_row < 0:
        return _fail("reopen_history_undo", "历史对话框中找不到对应记录行")
    dialog.list_widget.setCurrentRow(target_row)
    dialog.btn_undo.click()  # 确认框自动 Yes → 真实 FileRenameService.undo_record
    for old, new in rename_map.items():
        if not Path(old).is_file():
            return _fail("reopen_history_undo", f"撤销后原文件未恢复: {old}")
        if Path(new).exists():
            return _fail("reopen_history_undo", f"撤销后新文件仍存在: {new}")
    refreshed = store.get_record("rename", record["id"])
    if refreshed is None or not refreshed.get("undone"):
        return _fail("reopen_history_undo", "撤销后历史记录未标记 undone")
    dialog.deleteLater()
    return _pass(
        "reopen_history_undo",
        artifacts={
            "restored": list(rename_map),
            "record_id": record_id,
            "workdir": str(work),
        },
    )


# --------------------------------------------------------------------------- #
#  组装:模式 → 场景表;执行;报告;退出码
# --------------------------------------------------------------------------- #


def _scenarios_for(
    mode: str,
) -> list[tuple[str, Callable[[_SelftestContext], ScenarioOutcome]]]:
    pure: list[tuple[str, Callable[[_SelftestContext], ScenarioOutcome]]] = [
        ("pages", _scenario_pages),
        ("rename", _scenario_rename),
        ("pdf_pure", _scenario_pdf_pure),
        ("pdf_cancel", _scenario_pdf_cancel),
        ("markdown", _scenario_markdown),
        ("replace_pure", _scenario_replace_pure),
        ("update_mutex", _scenario_update_mutex),
    ]
    if mode == "pure":
        return pure
    if mode == "full":
        return [
            *pure,
            ("pdf_office", _scenario_pdf_office),
            ("replace_office", _scenario_replace_office),
            ("attendance", _scenario_attendance),
        ]
    if mode == "reopen":
        return [("reopen_history_undo", _scenario_reopen_history_undo)]
    raise ValueError(f"未知的 selftest 模式: {mode!r}")


def _run_scenario(
    name: str,
    scenario: Callable[[_SelftestContext], ScenarioOutcome],
    ctx: _SelftestContext,
) -> ScenarioOutcome:
    started = time.monotonic()
    try:
        outcome = scenario(ctx)
    except Exception as error:  # noqa: BLE001 - 场景失败要落报告,不能中断后续场景
        logger.exception("selftest 场景异常: %s", name)
        outcome = ScenarioOutcome(name, "fail", f"{type(error).__name__}: {error}")
    outcome.duration_s = round(time.monotonic() - started, 3)
    logger.info(
        "selftest 场景 %s → %s(%.1fs): %s",
        name,
        outcome.status,
        outcome.duration_s,
        outcome.detail[:300],
    )
    return outcome


def _aggregate(outcomes: list[ScenarioOutcome]) -> int:
    if any(outcome.status == "fail" for outcome in outcomes):
        return EXIT_FAIL
    if any(outcome.status == "evidence_missing" for outcome in outcomes):
        return EXIT_EVIDENCE_MISSING
    return EXIT_PASS


def _default_report_path() -> Path:
    return current_data_root().parent / "selftest-report.json"


def _write_report(
    path: Path,
    mode: str,
    outcomes: list[ScenarioOutcome],
    exit_code: int,
    boxes: _MessageBoxScript | None = None,
) -> None:
    from file_toolbox.common.metadata import runtime_version

    payload: dict[str, object] = {
        "mode": mode,
        "exit_code": exit_code,
        "exe": sys.executable,
        "pid": os.getpid(),
        "version": runtime_version(),
        "platform": sys.platform,
        "data_root": str(current_data_root()),
        "workdir": str(current_data_root().parent / "selftest-work"),
        "history_dir": str(current_data_root().parent / ".file_toolbox" / "history"),
        "scenarios": [
            {
                "name": outcome.name,
                "status": outcome.status,
                "detail": outcome.detail,
                "duration_s": outcome.duration_s,
                "artifacts": outcome.artifacts,
            }
            for outcome in outcomes
        ],
        "message_boxes": boxes.records if boxes is not None else [],
        "evidence_missing": exit_code == EXIT_EVIDENCE_MISSING,
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    logger.info("selftest 报告已写出: %s(exit=%s)", path, exit_code)


def _write_report_best_effort(
    path: Path,
    mode: str,
    outcomes: list[ScenarioOutcome],
    exit_code: int,
    boxes: _MessageBoxScript | None = None,
) -> bool:
    """尽力写报告:OSError 记为失败 outcome 返回 False,不向调用方抛出。

    报告写出失败不能截断收尾流程(F1):中间/最终报告写失败都先落内存失败
    证据(保证退出码非零),由调用方继续等待真实 finished/关闭后再返回。
    """
    try:
        _write_report(path, mode, outcomes, exit_code, boxes)
        return True
    except OSError as error:
        logger.exception("selftest 报告写出失败: %s", path)
        outcomes.append(ScenarioOutcome("report", "fail", f"报告写出失败: {error}"))
        return False


def _execute_selftest(
    app: QApplication,
    mode: str,
    report_path: Path | None,
    *,
    close_wait_s: float = _CLOSE_WAIT_S,
) -> int:
    """同步执行全部场景(等待经嵌套事件循环),返回进程退出码。

    close_wait_s 仅是"先写失败报告"的期限:超时后仍保留窗口/worker 与事件
    循环继续协作等真实 finished/关闭,绝不以返回截断 QThread 生命周期;进程
    级超时由外层 product_selftest 报告 owned PID 并保留现场(不杀)。
    """
    allowed, reason = ensure_selftest_data_root()
    if not allowed:
        logger.error("selftest 数据根门禁拒绝: %s", reason)
        data_root = current_data_root().resolve()
        if report_path is not None and not report_path.resolve().is_relative_to(data_root):
            _write_report(
                report_path,
                mode,
                [ScenarioOutcome("scope", "fail", reason)],
                EXIT_FAIL,
            )
        return EXIT_FAIL
    workdir = current_data_root().parent / "selftest-work"
    workdir.mkdir(parents=True, exist_ok=True)
    coordinator = SelftestUpdateCoordinator()
    previous_quit_on_close = app.quitOnLastWindowClosed()
    app.setQuitOnLastWindowClosed(False)  # 退出码由场景结果决定,不经窗口关闭
    boxes = _MessageBoxScript()
    outcomes: list[ScenarioOutcome] = []
    try:
        with boxes:
            from file_toolbox.gui.main_window import MainWindow

            window = MainWindow(coordinator=coordinator)
            window.show()
            _settle(app)
            ctx = _SelftestContext(
                app=app,
                window=window,
                workdir=workdir,
                coordinator=coordinator,
                boxes=boxes,
                outcomes=outcomes,
            )
            try:
                for name, scenario in _scenarios_for(mode):
                    outcomes.append(_run_scenario(name, scenario, ctx))
            finally:
                # 真实关闭路径收尾(F1):任务均已等真实 finished;若仍有收尾中的
                # worker,closeEvent 协作延迟——超 close_wait_s 先写失败报告,但
                # 保留窗口/worker 与事件循环继续等真实关闭,绝不主动退出/强杀。
                window.close()
                exceeded = False
                while window.isVisible():
                    try:
                        _wait_until(app, lambda: not window.isVisible(), close_wait_s)
                    except TimeoutError:
                        if not exceeded:
                            exceeded = True
                            outcomes.append(
                                ScenarioOutcome(
                                    "close",
                                    "fail",
                                    f"窗口关闭超出协作收尾期限({close_wait_s}s),已先落失败报告并继续等待真实收尾",
                                )
                            )
                            # 写失败只记为额外失败证据,不截断对真实收尾的等待(F1)
                            _write_report_best_effort(
                                report_path or _default_report_path(),
                                mode,
                                outcomes,
                                _aggregate(outcomes),
                                boxes,
                            )
    finally:
        coordinator.release.set()
        app.setQuitOnLastWindowClosed(previous_quit_on_close)
    exit_code = _aggregate(outcomes)
    # 终版报告写失败同样保留为非零结果,不让 driver 以异常逃出(F1)
    if not _write_report_best_effort(
        report_path or _default_report_path(), mode, outcomes, exit_code, boxes
    ):
        return EXIT_FAIL
    return exit_code


def make_selftest_driver(mode: str, report_path: Path | None) -> Callable[[QApplication], int]:
    """构造 run_gui 的 selftest driver(真实启动链内调用,返回进程退出码)。"""

    def driver(app: QApplication) -> int:
        return _execute_selftest(app, mode, report_path)

    return driver
