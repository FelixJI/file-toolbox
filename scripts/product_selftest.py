"""项目专属成品自测:解包便携 zip → 启动真实 EXE --selftest → 独立读回验证。

与 scripts/release_smoke.py 的关系:release_smoke 只验证资产结构;本脚本在
同一构建上下文(AUTOMATION_ARTIFACTS_DIR/AUTOMATION_VERSION)上追加真实
业务入口验证——实际启动便携布局中的 current/FileToolbox.exe,经真实 Qt
事件循环/页面/worker 完成代表业务,并由本脚本在 EXE 进程外独立读回结果
(PDF 页数、docx XML、xlsx 工作簿、考勤输出、历史/撤销),不把 EXE 存活
或自述当作业务证据。

运行目录:固定父目录 ``<repo>/build/selftest/`` 下每次运行新建 ``run-*``
子目录,只保留不清理(本脚本不删除任何目录)。

超时契约:EXE 内部任务按协作取消/真实 finished 收尾;进程级超时只 fail
closed——不终止进程(full 模式可能持有真实 Office 会话),保留现场并报告
本次自建子进程的 PID 与路径,不继续第二次启动抢占同一数据根。

用法(CI/本地,cwd 为仓库根):
    uv run --all-extras python scripts/product_selftest.py --mode pure
    uv run --all-extras python scripts/product_selftest.py --mode full

退出码:0=通过;1=失败;3=EVIDENCE_MISSING(full 模式本机无 Office 等,
不把缺失 skip 成 PASS);4=超时,现场已保留。
"""

from __future__ import annotations

import argparse
import json
import os
import subprocess
import sys
import zipfile
from datetime import UTC, datetime
from pathlib import Path
from typing import TypedDict

_ROOT = Path(__file__).resolve().parents[1]
SELFTEST_RUN_ROOT = _ROOT / "build" / "selftest"

EXIT_PASS = 0
EXIT_FAIL = 1
EXIT_EVIDENCE_MISSING = 3
EXIT_TIMEOUT = 4

_LAUNCH_TIMEOUT_S = {"pure": 900, "full": 2700, "reopen": 300}

# 报告工件的实际 JSON 结构(与 gui.selftest_driver 的 ArtifactValue 一致)
ArtifactJSON = str | list[str] | dict[str, str] | bool | int


class ScenarioEntry(TypedDict):
    name: str
    status: str
    detail: str
    duration_s: float
    artifacts: dict[str, ArtifactJSON]


class SelftestReport(TypedDict):
    mode: str
    exit_code: int
    data_root: str
    history_dir: str
    workdir: str
    evidence_missing: bool
    scenarios: list[ScenarioEntry]


# 外层独立锁定的必需场景集合(与 driver 契约对应;缺失/多余/重复/未过均失败)
_PURE_SCENARIO_NAMES = (
    "pages",
    "rename",
    "pdf_pure",
    "pdf_cancel",
    "markdown",
    "replace_pure",
    "update_mutex",
)
_FULL_SCENARIO_NAMES = (*_PURE_SCENARIO_NAMES, "pdf_office", "replace_office", "attendance")
_REOPEN_SCENARIO_NAMES = ("reopen_history_undo",)


def _expected_scenario_names(mode: str) -> tuple[str, ...]:
    if mode == "full":
        return _FULL_SCENARIO_NAMES
    if mode == "reopen":
        return _REOPEN_SCENARIO_NAMES
    return _PURE_SCENARIO_NAMES


def _resolve_inputs(args: argparse.Namespace) -> tuple[Path, str]:
    """解析便携 zip 与版本:显式参数优先,CI 环境变量兜底。"""
    version = args.version or os.environ.get("AUTOMATION_VERSION")
    artifacts = Path(args.artifacts_dir or os.environ.get("AUTOMATION_ARTIFACTS_DIR") or "")
    zip_path = Path(args.zip) if args.zip else None
    if version is None and zip_path is None:
        raise SystemExit(
            "::error::缺少版本/zip:传 --zip 与 --version,或设置 "
            "AUTOMATION_ARTIFACTS_DIR 与 AUTOMATION_VERSION(与 release_smoke 相同)"
        )
    if zip_path is None:
        if not version or not str(artifacts):
            raise SystemExit(f"::error::参数不完整: version={version!r} artifacts={artifacts!r}")
        zip_path = artifacts / f"FileToolbox-v{version}-win-x64.zip"
    if not zip_path.is_file():
        raise SystemExit(f"::error::便携 zip 不存在: {zip_path}")
    return zip_path, version or zip_path.stem


def _new_run_dir() -> Path:
    """在固定父目录下新建本次运行子目录;只新建,不删除任何旧目录。"""
    SELFTEST_RUN_ROOT.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ")
    run_dir = SELFTEST_RUN_ROOT / f"run-{stamp}-{os.getpid()}"
    run_dir.mkdir(parents=True)
    return run_dir


def _safe_extract(portable_zip: Path, destination: Path) -> None:
    """解包便携 zip 并防越界:拒绝绝对路径/../ 穿透与目标外路径。"""
    destination.mkdir(parents=True, exist_ok=True)
    root = destination.resolve()
    with zipfile.ZipFile(portable_zip) as package:
        for name in package.namelist():
            target = (destination / name).resolve()
            if root != target and root not in target.parents:
                raise SystemExit(f"::error::zip 条目越界,拒绝解包: {name!r}")
        package.extractall(destination)


def _qt_offscreen_available(portable: Path) -> bool:
    """探测解包产物内是否携带 Qt offscreen 平台插件(证据驱动,不猜测)。"""
    return any(portable.rglob("qoffscreen*.dll"))


def _launch_exe(
    exe: Path,
    mode: str,
    report: Path,
    *,
    visible: bool,
    timeout_s: int,
    stdout_log: Path,
) -> int:
    """启动真实 EXE 执行 --selftest;超时 fail closed 保留现场,不终止进程。

    平台选择:--visible 用原生窗口;默认在产物携带 offscreen 插件时离屏运行
    (CI 无桌面/本地后台更确定),未携带则回退原生平台(Windows runner 有桌面)。
    """
    command = [str(exe), "--selftest", mode, "--selftest-report", str(report)]
    env = dict(os.environ)
    env.pop("QT_QPA_PLATFORM", None)
    if not visible and _qt_offscreen_available(exe.parent.parent):
        env["QT_QPA_PLATFORM"] = "offscreen"
    creationflags = subprocess.CREATE_NO_WINDOW if sys.platform == "win32" else 0
    with stdout_log.open("wb") as log:
        process = subprocess.Popen(  # noqa: S603 - 固定路径的自建产物
            command,
            stdout=log,
            stderr=subprocess.STDOUT,
            env=env,
            creationflags=creationflags,
            cwd=str(exe.parent.parent),
        )
        try:
            return process.wait(timeout=timeout_s)
        except subprocess.TimeoutExpired:
            print(
                f"::error::EXE selftest 超时({timeout_s}s),不终止进程(可能持有真实 "
                f"Office 会话/未收尾业务):owned pid={process.pid} "
                f"report={report} log={stdout_log} cmd={' '.join(command)}; "
                "现场保留,交主代理处置,不继续第二次启动",
                file=sys.stderr,
            )
            return EXIT_TIMEOUT


def _load_report(path: Path) -> SelftestReport:
    if not path.is_file():
        raise FileNotFoundError(f"EXE 未写出报告: {path}")
    raw = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(raw, dict) or not isinstance(raw.get("scenarios"), list):
        raise ValueError(f"报告结构无效: {path}")
    evidence_missing = raw.get("evidence_missing")
    if not isinstance(evidence_missing, bool):
        raise ValueError(f"报告 evidence_missing 字段无效: {evidence_missing!r}")
    scenarios: list[ScenarioEntry] = []
    for entry in raw["scenarios"]:
        if not isinstance(entry, dict) or not isinstance(entry.get("name"), str):
            raise ValueError(f"报告场景结构无效: {entry!r}")
        artifacts_raw = entry.get("artifacts", {})
        if not isinstance(artifacts_raw, dict):
            raise ValueError(f"报告场景工件结构无效: {entry.get('name')!r}")
        for value in artifacts_raw.values():
            if isinstance(value, dict):
                if not all(isinstance(k, str) and isinstance(v, str) for k, v in value.items()):
                    raise ValueError(f"报告工件映射值非字符串: {entry.get('name')!r}")
            elif isinstance(value, list):
                if not all(isinstance(item, str) for item in value):
                    raise ValueError(f"报告工件列表值非字符串: {entry.get('name')!r}")
            elif not isinstance(value, str | bool | int):
                raise ValueError(f"报告工件类型不符: {entry.get('name')!r}")
        scenarios.append(
            ScenarioEntry(
                name=entry["name"],
                status=str(entry.get("status", "")),
                detail=str(entry.get("detail", "")),
                duration_s=float(entry.get("duration_s", 0.0)),
                artifacts=artifacts_raw,
            )
        )
    return SelftestReport(
        mode=str(raw.get("mode", "")),
        exit_code=int(raw.get("exit_code", -1)),
        data_root=str(raw.get("data_root", "")),
        history_dir=str(raw.get("history_dir", "")),
        workdir=str(raw.get("workdir", "")),
        evidence_missing=evidence_missing,
        scenarios=scenarios,
    )


def _scenario(report: SelftestReport, name: str) -> ScenarioEntry | None:
    return next((item for item in report["scenarios"] if item["name"] == name), None)


def _validate_report(
    report: SelftestReport, *, expected_mode: str, child_code: int, portable_root: Path
) -> None:
    """报告统一 fail-closed 契约(F3):结构/退出码/场景集合/数据根全部核对。

    任何不一致(含报告 exit_code 与子进程不符、场景缺失/重复/多余/未通过、
    数据根不在本次解包的 portable 根内)都拋 AssertionError,不进入读回。
    child_code==EXIT_EVIDENCE_MISSING 时允许场景为 evidence_missing(P3:
    真实缺证据报告可验后透传),其余非 pass 状态仍一律拒绝。
    """
    if report["mode"] != expected_mode:
        raise AssertionError(f"报告 mode 不符: {report['mode']!r} ≠ {expected_mode!r}")
    if report["exit_code"] != child_code:
        raise AssertionError(
            f"报告 exit_code 与子进程退出码不一致: {report['exit_code']} ≠ {child_code}"
        )
    # 顶层缺证据标志必须与退出码/场景状态一致(F3):全 pass/exit0 却声明
    # evidence_missing=true 属于矛盾报告,fail closed 拒绝进入读回。
    if report["evidence_missing"] is not (report["exit_code"] == EXIT_EVIDENCE_MISSING):
        raise AssertionError(
            f"报告 evidence_missing 标志与退出码矛盾: {report['evidence_missing']} "
            f"vs exit_code={report['exit_code']}"
        )
    has_missing_scenario = any(item["status"] == "evidence_missing" for item in report["scenarios"])
    if has_missing_scenario and report["exit_code"] == EXIT_PASS:
        raise AssertionError("存在 evidence_missing 场景但 exit_code 为 PASS,状态矛盾")
    if not has_missing_scenario and report["exit_code"] == EXIT_EVIDENCE_MISSING:
        raise AssertionError("exit_code 为 EVIDENCE_MISSING 但没有 evidence_missing 场景,状态矛盾")
    names = [item["name"] for item in report["scenarios"]]
    if len(names) != len(set(names)):
        raise AssertionError(f"场景重复: {names}")
    expected = _expected_scenario_names(expected_mode)
    if sorted(names) != sorted(expected):
        raise AssertionError(f"场景集合不符: {sorted(names)} ≠ {sorted(expected)}")
    not_passed = [
        (item["name"], item["status"])
        for item in report["scenarios"]
        if item["status"] != "pass"
        and not (child_code == EXIT_EVIDENCE_MISSING and item["status"] == "evidence_missing")
    ]
    if not_passed:
        raise AssertionError(f"存在未通过场景: {not_passed}")
    # 路径用 Path 精确绑定(相等/子树),不接受同前缀兄弟目录
    portable = portable_root.resolve()
    data_root = Path(report["data_root"]).resolve()
    if data_root != portable / ".file_toolbox":
        raise AssertionError(f"报告 data_root 未精确绑定本次解包根: {data_root}")
    if Path(report["history_dir"]).resolve() != data_root / "history":
        raise AssertionError(f"报告 history_dir 未绑定 data_root/history: {report['history_dir']}")
    if Path(report["workdir"]).resolve() != portable / "selftest-work":
        raise AssertionError(f"报告 workdir 未绑定 portable/selftest-work: {report['workdir']}")


def _assert_artifacts_under(scenario: ScenarioEntry, base: Path, keys: tuple[str, ...]) -> None:
    """指定路径型工件必须全部位于 base 子树内(本次场景范围,不接受外部路径)。

    只检查声明的路径键——工件中的非路径字段(如 history_tool 标签、裸文件
    名列表)不在此列,避免误判;未声明的键由具体读回逻辑负责。
    """
    base_resolved = base.resolve()
    for key in keys:
        value = scenario["artifacts"].get(key)
        paths: list[str]
        if isinstance(value, str):
            paths = [value]
        elif isinstance(value, list):
            paths = value
        elif isinstance(value, dict):
            paths = list(value.values())
        else:
            continue  # 缺键/bool/int 工件不是路径
        for item in paths:
            if not Path(item).resolve().is_relative_to(base_resolved):
                raise AssertionError(f"场景 {scenario['name']} 工件 {key} 越出本次运行范围: {item}")


def _require_scenario(report: SelftestReport, name: str) -> ScenarioEntry:
    scenario = _scenario(report, name)
    if scenario is None:
        raise AssertionError(f"报告缺少场景 {name}: {[s['name'] for s in report['scenarios']]}")
    if scenario["status"] != "pass":
        raise AssertionError(f"场景 {name} 未通过: {scenario['status']} {scenario['detail']}")
    return scenario


def _artifact_str(scenario: ScenarioEntry, key: str) -> str:
    value = scenario["artifacts"].get(key)
    if not isinstance(value, str):
        raise AssertionError(f"场景 {scenario['name']} 工件 {key} 不是路径字符串: {value!r}")
    return value


def _artifact_str_list(scenario: ScenarioEntry, key: str) -> list[str]:
    value = scenario["artifacts"].get(key)
    if not isinstance(value, list) or not all(isinstance(item, str) for item in value):
        raise AssertionError(f"场景 {scenario['name']} 工件 {key} 不是字符串列表: {value!r}")
    return value


def _artifact_str_map(scenario: ScenarioEntry, key: str) -> dict[str, str]:
    value = scenario["artifacts"].get(key)
    if not isinstance(value, dict) or not all(
        isinstance(k, str) and isinstance(v, str) for k, v in value.items()
    ):
        raise AssertionError(f"场景 {scenario['name']} 工件 {key} 不是字符串映射: {value!r}")
    return value


def _assert_pdf_pages(path: Path, minimum: int = 1) -> int:
    from pypdf import PdfReader

    try:
        pages = len(PdfReader(str(path)).pages)
    except Exception as error:
        raise AssertionError(f"PDF 无法解析: {path} ({error})") from error
    if pages < minimum:
        raise AssertionError(f"PDF 页数异常: {path} → {pages}")
    return pages


def _assert_docx_contains(path: Path, needle: str, *, absent: str | None = None) -> None:
    with zipfile.ZipFile(path) as package:
        xml = package.read("word/document.xml").decode("utf-8")
    if needle not in xml:
        raise AssertionError(f"docx 缺少文本 {needle!r}: {path}")
    if absent is not None and absent in xml:
        raise AssertionError(f"docx 仍含旧文本 {absent!r}: {path}")


def _verify_run1(report: SelftestReport, mode: str) -> list[str]:
    """EXE 进程外独立读回验证;返回核对清单(异常即失败)。"""
    checks: list[str] = []
    workdir = Path(report["workdir"])
    rename = _require_scenario(report, "rename")
    _assert_artifacts_under(rename, workdir, ("rename_map", "workdir"))
    rename_map = _artifact_str_map(rename, "rename_map")
    if not rename_map:
        raise AssertionError("rename.rename_map 为空")
    for old, new in rename_map.items():
        if not Path(new).is_file():
            raise AssertionError(f"重命名文件缺失: {new}")
        checks.append(f"rename: {Path(old).name} → {Path(new).name}")
    pdf_pure = _require_scenario(report, "pdf_pure")
    _assert_artifacts_under(pdf_pure, workdir, ("pdfs", "pdf_copy"))
    pdfs = _artifact_str_list(pdf_pure, "pdfs")
    if not pdfs:
        raise AssertionError("pdf_pure.pdfs 为空")
    for pdf in [*pdfs, _artifact_str(pdf_pure, "pdf_copy")]:
        checks.append(f"pdf 页数 {Path(pdf).name}: {_assert_pdf_pages(Path(pdf))}")
    markdown = _require_scenario(report, "markdown")
    _assert_artifacts_under(markdown, workdir, ("docx_output", "tables_output", "document_output"))
    _assert_docx_contains(Path(_artifact_str(markdown, "docx_output")), "自测标题")
    checks.append("markdown docx XML")
    from openpyxl import load_workbook

    tables = load_workbook(_artifact_str(markdown, "tables_output"), read_only=True)
    try:
        if "表1" not in tables.sheetnames:
            raise AssertionError("tables 工作簿缺 表1")
    finally:
        tables.close()
    document = load_workbook(_artifact_str(markdown, "document_output"), read_only=True)
    try:
        if "正文" not in document.sheetnames:
            raise AssertionError("document 工作簿缺 正文")
    finally:
        document.close()
    checks.append("markdown 两 xlsx 工作簿")
    replace_pure = _require_scenario(report, "replace_pure")
    _assert_artifacts_under(replace_pure, workdir, ("txt", "md"))
    for key in ("txt", "md"):
        path = Path(_artifact_str(replace_pure, key))
        # 产品 txt 写出为 \r\r\n(既有行为),按内容而非行尾比对
        content = path.read_text(encoding="utf-8").replace("\r", "").replace("\n", "")
        if content != "nihao selftest nihao":
            raise AssertionError(f"替换读回不符: {path}")
    checks.append("replace txt/md")
    mutex = _require_scenario(report, "update_mutex")
    if mutex["artifacts"].get("locked_then_unlocked") is not True:
        raise AssertionError("更新互斥证据缺失")
    checks.append("update mutex")
    history_file = Path(report["history_dir"]) / "rename.jsonl"
    lines = [line for line in history_file.read_text(encoding="utf-8").splitlines() if line.strip()]
    if not any("rename_map" in line for line in lines):
        raise AssertionError(f"rename 历史缺失: {history_file}")
    checks.append("history jsonl")
    if mode == "full":
        pdf_office = _require_scenario(report, "pdf_office")
        # 两条 Office PDF 工件缺一不可(F3):必须在对应场景子目录 workdir/pdf-office
        # 内(不接受指向 pure 场景等其它路径的复用),且解析后两路径必须不同
        # (同一文件重复不构成两条证据);各自读回有效 PDF 页数。
        _assert_artifacts_under(pdf_office, workdir / "pdf-office", ("docx", "xlsx"))
        docx_pdf = Path(_artifact_str(pdf_office, "docx"))
        xlsx_pdf = Path(_artifact_str(pdf_office, "xlsx"))
        if docx_pdf.resolve() == xlsx_pdf.resolve():
            raise AssertionError(f"office pdf docx/xlsx 指向同一文件: {docx_pdf}")
        checks.append(f"office pdf docx: {_assert_pdf_pages(docx_pdf)}")
        checks.append(f"office pdf xlsx: {_assert_pdf_pages(xlsx_pdf)}")
        replace_office = _require_scenario(report, "replace_office")
        _assert_artifacts_under(replace_office, workdir, ("docx",))
        _assert_docx_contains(
            Path(_artifact_str(replace_office, "docx")),
            "office-replaced",
            absent="selftest-body",
        )
        checks.append("replace office docx")
        attendance = _require_scenario(report, "attendance")
        _assert_artifacts_under(
            attendance,
            workdir,
            ("source", "source_roster", "template", "roster", "plain_output", "roster_output"),
        )
        plain_path = _artifact_str(attendance, "plain_output")
        plain = load_workbook(plain_path, read_only=True)
        try:
            for sheet in ("甲明细", "乙明细"):
                if sheet not in plain.sheetnames:
                    raise AssertionError(f"考勤普通输出缺工作表 {sheet}")
        finally:
            plain.close()
        roster = load_workbook(_artifact_str(attendance, "roster_output"), read_only=True)
        try:
            if "丙明细" not in roster.sheetnames or "丁明细" not in roster.sheetnames:
                raise AssertionError(f"考勤名单输出缺工作表: {roster.sheetnames}")
        finally:
            roster.close()
        checks.append("attendance 输出")
    return checks


def _verify_run2(report: SelftestReport, run1: SelftestReport) -> list[str]:
    """reopen 精确读回:绑定 run1 的具体 rename 记录与新旧路径,不用任意历史行。

    核对:restored 与 run1 rename_map 一致;新路径全部消失;原文件恢复且内容
    为虚构样本内容;run1 记录的精确 record_id 在历史中标记 undone;两次运行
    数据根一致。
    """
    reopen = _require_scenario(report, "reopen_history_undo")
    rename = _require_scenario(run1, "rename")
    _assert_artifacts_under(reopen, Path(report["workdir"]), ("restored",))
    rename_map = _artifact_str_map(rename, "rename_map")
    record_id = rename["artifacts"].get("record_id")
    if not isinstance(record_id, int) or record_id < 1:
        raise AssertionError("run1 rename 场景缺少精确 record_id")
    restored = _artifact_str_list(reopen, "restored")
    if not restored:
        raise AssertionError("reopen.restored 为空")
    if sorted(restored) != sorted(rename_map):
        raise AssertionError(
            f"restored 与 run1 rename_map 不一致: {restored} vs {list(rename_map)}"
        )
    checks: list[str] = []
    for old, new in rename_map.items():
        if Path(new).exists():
            raise AssertionError(f"撤销后新路径仍存在: {new}")
        old_path = Path(old)
        if not old_path.is_file():
            raise AssertionError(f"跨启动撤销后原文件未恢复: {old}")
        if f"content of {old_path.name}" not in old_path.read_text(encoding="utf-8"):
            raise AssertionError(f"原样本内容未恢复: {old}")
        checks.append(f"undo: {old_path.name}")
    if run1["data_root"] != report["data_root"] or run1["workdir"] != report["workdir"]:
        raise AssertionError(
            f"两次运行数据根/工作目录不一致: {run1['data_root']} ≠ {report['data_root']}"
        )
    history_file = Path(report["history_dir"]) / "rename.jsonl"
    target: dict | None = None
    for line in history_file.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        entry = json.loads(line)
        if isinstance(entry, dict) and entry.get("id") == record_id:
            target = entry
    if target is None or target.get("undone") is not True:
        raise AssertionError(f"精确记录未标记 undone(id={record_id}): {history_file}")
    checks.append(f"history undone(record {record_id})")
    return checks


def _summarize(report: SelftestReport) -> str:
    return " ".join(
        f"{item['name']}={item['status']}({item['duration_s']}s)" for item in report["scenarios"]
    )


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--mode", choices=("pure", "full"), default="pure")
    parser.add_argument("--zip", help="便携 zip 路径(默认由 artifacts-dir+version 推导)")
    parser.add_argument("--version", help="版本(默认取 AUTOMATION_VERSION)")
    parser.add_argument("--artifacts-dir", help="构建产物目录(默认取 AUTOMATION_ARTIFACTS_DIR)")
    parser.add_argument(
        "--visible", action="store_true", help="使用真实窗口平台(默认优先离屏,CI 可用)"
    )
    args = parser.parse_args(argv)

    portable_zip, version = _resolve_inputs(args)
    run_dir = _new_run_dir()
    portable = run_dir / "portable"
    print(f"selftest: version={version} zip={portable_zip} run={run_dir}")
    _safe_extract(portable_zip, portable)
    exe = portable / "current" / "FileToolbox.exe"
    if not exe.is_file():
        print(f"::error::便携布局缺少 current/FileToolbox.exe: {exe}", file=sys.stderr)
        return EXIT_FAIL

    report1 = run_dir / "report-run1.json"
    code = _launch_exe(
        exe,
        args.mode,
        report1,
        visible=args.visible,
        timeout_s=_LAUNCH_TIMEOUT_S[args.mode],
        stdout_log=run_dir / "exe-run1.log",
    )
    if code == EXIT_TIMEOUT:
        return EXIT_TIMEOUT  # 保留现场,不继续第二次启动
    try:
        report = _load_report(report1)
    except (FileNotFoundError, ValueError, json.JSONDecodeError) as error:
        print(f"::error::第一次启动报告无效(exit={code}): {error}", file=sys.stderr)
        return EXIT_FAIL
    print(f"run1 scenarios: {_summarize(report)}")
    if code != EXIT_PASS:
        if code != EXIT_EVIDENCE_MISSING:
            return EXIT_FAIL
        # 透传 3 前先验完整报告契约(P3):mode/数据根/退出码/标志/场景集合
        # 一致性全部核对——允许真实 evidence_missing 报告,矛盾报告(如报告
        # 声称 exit_code=0)一律按失败(1),不冒称"缺少证据"。
        try:
            _validate_report(
                report, expected_mode=args.mode, child_code=code, portable_root=portable
            )
        except (AssertionError, OSError, KeyError, ValueError, json.JSONDecodeError) as error:
            print(f"::error::evidence_missing 报告校验失败(按失败处理): {error}", file=sys.stderr)
            return EXIT_FAIL
        return EXIT_EVIDENCE_MISSING
    try:
        _validate_report(report, expected_mode=args.mode, child_code=code, portable_root=portable)
        checks = _verify_run1(report, args.mode)
    except (AssertionError, OSError, KeyError, ValueError, json.JSONDecodeError) as error:
        print(f"::error::run1 报告校验/读回验证失败: {error}", file=sys.stderr)
        return EXIT_FAIL
    print("run1 读回验证: " + "; ".join(checks))

    report2 = run_dir / "report-run2.json"
    code = _launch_exe(
        exe,
        "reopen",
        report2,
        visible=args.visible,
        timeout_s=_LAUNCH_TIMEOUT_S["reopen"],
        stdout_log=run_dir / "exe-run2.log",
    )
    if code == EXIT_TIMEOUT:
        return EXIT_TIMEOUT
    try:
        reopen_report = _load_report(report2)
    except (FileNotFoundError, ValueError, json.JSONDecodeError) as error:
        print(f"::error::第二次启动报告无效(exit={code}): {error}", file=sys.stderr)
        return EXIT_FAIL
    print(f"run2 scenarios: {_summarize(reopen_report)}")
    if code != EXIT_PASS:
        return EXIT_FAIL
    try:
        _validate_report(
            reopen_report, expected_mode="reopen", child_code=code, portable_root=portable
        )
        checks = _verify_run2(reopen_report, report)
    except (AssertionError, OSError, KeyError, ValueError, json.JSONDecodeError) as error:
        print(f"::error::run2 报告校验/读回验证失败: {error}", file=sys.stderr)
        return EXIT_FAIL
    print("run2 读回验证: " + "; ".join(checks))

    summary = {
        "mode": args.mode,
        "version": version,
        "zip": str(portable_zip),
        "run_dir": str(run_dir),
        "run1": {"exit_code": report["exit_code"], "scenarios": report["scenarios"]},
        "run2": {
            "exit_code": reopen_report["exit_code"],
            "scenarios": reopen_report["scenarios"],
        },
        "result": "PASS",
    }
    (run_dir / "summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(f"PASS: mode={args.mode} run_dir={run_dir}")
    return EXIT_PASS


if __name__ == "__main__":
    raise SystemExit(main())
