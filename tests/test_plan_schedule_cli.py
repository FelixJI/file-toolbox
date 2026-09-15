"""plan-schedule CLI 测试:预览默认、--yes 生成、参数校验、失败退出码。

全部基于程序化生成的虚构 xlsx,不触发 COM。
"""

from datetime import date

from openpyxl import load_workbook
from typer.testing import CliRunner

from file_toolbox.cli.main import app
from file_toolbox.core.plan_schedule import SHEET_NAME

runner = CliRunner()


def _make_input(make_xlsx, tmp_path, rows):
    return make_xlsx("清单.xlsx", {"S": rows})


def test_preview_lists_items_without_writing(make_xlsx, tmp_path):
    """默认预览:列项点行/覆盖月份/最大并行,不产生输出文件。"""
    src = _make_input(
        make_xlsx,
        tmp_path,
        [
            ["项点名称", "起始日期", "终止日期"],
            ["合5第8列", date(2026, 9, 17), date(2026, 9, 21)],
            ["杭9第2列", date(2026, 9, 21), date(2026, 9, 27)],
        ],
    )

    r = runner.invoke(app, ["plan-schedule", str(src)])

    assert r.exit_code == 0
    assert "共 2 个项点" in r.output
    assert "合5第8列  2026-09-17 ~ 2026-09-21 (5天)" in r.output
    assert "覆盖 2026年9月 ~ 2026年9月(1 个月)" in r.output
    assert "最大并行 2 个项点(2026-09-21)" in r.output
    assert "预览模式" in r.output
    assert not (tmp_path / "计划排布.xlsx").exists()


def test_preview_marks_invalid_rows(make_xlsx, tmp_path):
    """预览标出无效行(仍有有效项点时退出码 0)。"""
    src = _make_input(
        make_xlsx,
        tmp_path,
        [
            ["项点名称", "起始日期", "终止日期"],
            ["倒挂", date(2026, 9, 2), date(2026, 9, 1)],
            ["正常", date(2026, 9, 17), date(2026, 9, 21)],
        ],
    )

    r = runner.invoke(app, ["plan-schedule", str(src)])

    assert r.exit_code == 0
    assert "第2行 [无效]" in r.output
    assert "共 1 个项点" in r.output


def test_preview_without_items_errors(make_xlsx, tmp_path):
    """全部行无效 → 预览报错退出 1。"""
    src = _make_input(
        make_xlsx,
        tmp_path,
        [["项点名称", "起始日期", "终止日期"], ["倒挂", date(2026, 9, 2), date(2026, 9, 1)]],
    )

    r = runner.invoke(app, ["plan-schedule", str(src)])

    assert r.exit_code == 1
    assert "没有可解析的项点" in r.output


def test_execute_writes_default_output(make_xlsx, tmp_path):
    """--yes:默认输出到清单目录,工作表名正确。"""
    src = _make_input(
        make_xlsx,
        tmp_path,
        [["项点名称", "起始日期", "终止日期"], ["A", date(2026, 9, 17), date(2026, 9, 21)]],
    )

    r = runner.invoke(app, ["plan-schedule", str(src), "--yes"])

    assert r.exit_code == 0
    out = tmp_path / "计划排布.xlsx"
    assert out.is_file()
    assert load_workbook(out).sheetnames == [SHEET_NAME]
    assert "完成: 1 个项点" in r.output
    assert "[1/1] 生成 2026年9月" in r.output


def test_execute_custom_output_and_year(make_xlsx, tmp_path):
    """-o 自定义输出;--year 为缺年串补全年份。"""
    src = _make_input(
        make_xlsx,
        tmp_path,
        [["项点名称", "起始日期", "终止日期"], ["A", "9-17", "9-21"]],
    )
    out = tmp_path / "sub" / "自定义.xlsx"

    r = runner.invoke(app, ["plan-schedule", str(src), "--yes", "-o", str(out), "--year", "2027"])

    assert r.exit_code == 0
    assert out.is_file()
    wb = load_workbook(out)
    assert wb[SHEET_NAME]["B1"].value == "2027年9月"


def test_execute_all_invalid_exits_nonzero(make_xlsx, tmp_path):
    """全部行无效 → 退出码 1 并给出原因。"""
    src = _make_input(
        make_xlsx,
        tmp_path,
        [["项点名称", "起始日期", "终止日期"], ["倒挂", date(2026, 9, 2), date(2026, 9, 1)]],
    )

    r = runner.invoke(app, ["plan-schedule", str(src), "--yes"])

    assert r.exit_code == 1
    assert "没有可排布的项点" in r.output


def test_missing_input_errors(tmp_path):
    r = runner.invoke(app, ["plan-schedule", str(tmp_path / "missing.xlsx")])
    assert r.exit_code == 1
    assert "不存在" in r.output


def test_unsupported_suffix_errors(tmp_path):
    f = tmp_path / "清单.csv"
    f.write_text("x")
    r = runner.invoke(app, ["plan-schedule", str(f)])
    assert r.exit_code == 1
    assert "不支持的格式" in r.output


def test_no_argument_errors():
    r = runner.invoke(app, ["plan-schedule"])
    assert r.exit_code == 1
    assert "缺少项点清单文件参数" in r.output


def test_bad_header_errors(make_xlsx, tmp_path):
    src = _make_input(make_xlsx, tmp_path, [["随便", "表头"]])
    r = runner.invoke(app, ["plan-schedule", str(src)])
    assert r.exit_code == 1
    assert "未找到表头" in r.output


def test_output_auto_numbered_when_exists(make_xlsx, tmp_path):
    """输出已存在时自动加序号,原文件不被覆盖。"""
    src = _make_input(
        make_xlsx,
        tmp_path,
        [["项点名称", "起始日期", "终止日期"], ["A", date(2026, 9, 17), date(2026, 9, 21)]],
    )
    out = tmp_path / "计划排布.xlsx"
    out.write_text("precious")

    r = runner.invoke(app, ["plan-schedule", str(src), "--yes", "-o", str(out)])

    assert r.exit_code == 0
    assert out.read_text(encoding="utf-8") == "precious"
    assert (tmp_path / "计划排布_1.xlsx").is_file()
