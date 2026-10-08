"""scripts/product_selftest.py 契约:解包防越界、报告结构校验、读回验证、
超时不杀进程 fail closed、EVIDENCE_MISSING 透传、main 端到端(假启动器)。

不启动真实 EXE(真实成品验证由 .ci/project.json 的 release_smoke 步骤与
本机 full 运行承担);启动器以 monkeypatch 替身注入,其余全部真实。
"""

import json
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest

# 让 tests 能 import scripts 包
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import scripts.product_selftest as ps  # noqa: E402

# ---------------------------------------------------------------------------
# 输入解析与运行目录
# ---------------------------------------------------------------------------


def test_resolve_inputs_from_env(monkeypatch, tmp_path):
    (tmp_path / "FileToolbox-v9.9.9-win-x64.zip").write_bytes(b"zip")
    monkeypatch.setenv("AUTOMATION_ARTIFACTS_DIR", str(tmp_path))
    monkeypatch.setenv("AUTOMATION_VERSION", "9.9.9")
    zip_path, version = ps._resolve_inputs(_args(zip=None, version=None, artifacts_dir=None))
    assert version == "9.9.9"
    assert zip_path == tmp_path / "FileToolbox-v9.9.9-win-x64.zip"


def _args(**kwargs):
    import argparse

    return argparse.Namespace(**kwargs)


def test_resolve_inputs_missing_zip_fails_closed(monkeypatch, tmp_path):
    monkeypatch.setenv("AUTOMATION_ARTIFACTS_DIR", str(tmp_path))
    monkeypatch.setenv("AUTOMATION_VERSION", "9.9.9")
    with pytest.raises(SystemExit, match="便携 zip 不存在"):
        ps._resolve_inputs(_args(zip=None, version=None, artifacts_dir=None))


def test_resolve_inputs_requires_version_or_zip(monkeypatch):
    """缺版本/zip 时 fail closed;须隔离继承的 CI 真实环境(F9)。

    release_smoke/CI 进程会导出 AUTOMATION_VERSION/AUTOMATION_ARTIFACTS_DIR;
    不清理时本测试会走到“便携 zip 不存在”分支,预期“缺少版本”断言失真。
    """
    monkeypatch.delenv("AUTOMATION_VERSION", raising=False)
    monkeypatch.delenv("AUTOMATION_ARTIFACTS_DIR", raising=False)
    with pytest.raises(SystemExit, match="缺少版本"):
        ps._resolve_inputs(_args(zip=None, version=None, artifacts_dir=None))


def test_new_run_dir_creates_without_deleting(monkeypatch, tmp_path):
    monkeypatch.setattr(ps, "SELFTEST_RUN_ROOT", tmp_path)
    existing = tmp_path / "run-20200101T000000Z-1"
    existing.mkdir()
    marker = existing / "keep.txt"
    marker.write_text("x", encoding="utf-8")
    run_dir = ps._new_run_dir()
    assert run_dir.is_dir() and run_dir.name.startswith("run-")
    assert marker.is_file()  # 旧目录不清理、不删除


# ---------------------------------------------------------------------------
# 解包防越界
# ---------------------------------------------------------------------------


def _make_zip(path: Path, entries: dict[str, bytes]) -> None:
    with zipfile.ZipFile(path, "w") as package:
        for name, payload in entries.items():
            package.writestr(name, payload)


def test_safe_extract_normal_layout(tmp_path):
    portable_zip = tmp_path / "portable.zip"
    _make_zip(portable_zip, {"current/FileToolbox.exe": b"MZ", "current/data.txt": b"x"})
    destination = tmp_path / "out"
    ps._safe_extract(portable_zip, destination)
    assert (destination / "current" / "FileToolbox.exe").read_bytes() == b"MZ"


def test_safe_extract_rejects_zip_slip(tmp_path):
    portable_zip = tmp_path / "evil.zip"
    _make_zip(portable_zip, {"../escape.txt": b"x"})
    with pytest.raises(SystemExit, match="越界"):
        ps._safe_extract(portable_zip, tmp_path / "out")


def test_qt_offscreen_available_probe(tmp_path):
    assert ps._qt_offscreen_available(tmp_path) is False
    (tmp_path / "platforms").mkdir()
    (tmp_path / "platforms" / "qoffscreen.dll").write_bytes(b"d")
    assert ps._qt_offscreen_available(tmp_path) is True


# ---------------------------------------------------------------------------
# 报告结构校验与工件类型收紧
# ---------------------------------------------------------------------------


def _valid_report(tmp_path: Path, **overrides: object) -> Path:
    payload = {
        "mode": "pure",
        "exit_code": 0,
        "data_root": str(tmp_path / ".file_toolbox"),
        "history_dir": str(tmp_path / ".file_toolbox" / "history"),
        "scenarios": [
            {"name": "pages", "status": "pass", "detail": "", "duration_s": 0.1, "artifacts": {}}
        ],
        "evidence_missing": False,
        **overrides,
    }
    path = tmp_path / "report.json"
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    return path


def test_load_report_valid_and_invalid(tmp_path):
    report = ps._load_report(_valid_report(tmp_path))
    assert report["mode"] == "pure"
    assert report["scenarios"][0]["name"] == "pages"
    broken = tmp_path / "broken.json"
    broken.write_text(json.dumps({"scenarios": "not-a-list"}), encoding="utf-8")
    with pytest.raises(ValueError, match="报告结构无效"):
        ps._load_report(broken)
    bad_artifact_dir = tmp_path / "sub"
    bad_artifact_dir.mkdir()
    bad_artifact = _valid_report(
        bad_artifact_dir,
        scenarios=[
            {
                "name": "x",
                "status": "pass",
                "detail": "",
                "duration_s": 0.0,
                "artifacts": {"pdf": {"nested": {"deep": True}}},
            }
        ],
    )
    with pytest.raises(ValueError, match="报告工件"):
        ps._load_report(bad_artifact)


def test_artifact_accessors_reject_wrong_types():
    scenario = ps.ScenarioEntry(
        name="x",
        status="pass",
        detail="",
        duration_s=0.0,
        artifacts={"path": "a.pdf", "list": ["b"], "map": {"k": "v"}, "n": 3},
    )
    assert ps._artifact_str(scenario, "path") == "a.pdf"
    assert ps._artifact_str_list(scenario, "list") == ["b"]
    assert ps._artifact_str_map(scenario, "map") == {"k": "v"}
    with pytest.raises(AssertionError, match="不是路径字符串"):
        ps._artifact_str(scenario, "list")
    with pytest.raises(AssertionError, match="不是字符串列表"):
        ps._artifact_str_list(scenario, "n")


def _scenario_dict(name: str, status: str, detail: str = "", artifacts: object = None) -> dict:
    return {
        "name": name,
        "status": status,
        "detail": detail,
        "duration_s": 0.0,
        "artifacts": artifacts or {},
    }


def _require_scenario_reports_missing_and_failing():
    report = ps.SelftestReport(
        mode="pure",
        exit_code=1,
        data_root="",
        history_dir="",
        scenarios=[
            _scenario_dict("only", "pass"),
            _scenario_dict("bad", "fail", "boom"),
        ],
    )
    with pytest.raises(AssertionError, match="缺少场景 missing"):
        ps._require_scenario(report, "missing")
    with pytest.raises(AssertionError, match="未通过"):
        ps._require_scenario(report, "bad")


# ---------------------------------------------------------------------------
# 读回验证(合成 fixture:真实 pypdf/openpyxl/zip 文件)
# ---------------------------------------------------------------------------


def _fixture_outputs(root: Path) -> dict[str, str]:
    """构造与 driver 报告工件对应的真实输出文件(虚构内容)。"""
    from openpyxl import Workbook
    from pypdf import PdfWriter

    root.mkdir(parents=True, exist_ok=True)
    pdf = root / "pure-image-0.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with pdf.open("wb") as stream:
        writer.write(stream)
    pdf_copy = root / "pure-image-0_1.pdf"
    writer2 = PdfWriter()
    writer2.add_blank_page(width=200, height=200)
    with pdf_copy.open("wb") as stream:
        writer2.write(stream)
    docx = root / "sample-docx.docx"
    with zipfile.ZipFile(docx, "w") as package:
        package.writestr("word/document.xml", "<w:t>自测标题 selftest-body</w:t>")
    tables = root / "sample-tables.xlsx"
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.create_sheet("表1")["A1"] = "名称"
    workbook.save(tables)
    workbook.close()
    document = root / "sample-document.xlsx"
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.create_sheet("正文")
    workbook.save(document)
    workbook.close()
    txt = root / "notes.txt"
    txt.write_text("nihao selftest nihao\r\r\n", encoding="utf-8")
    history_dir = root / ".file_toolbox" / "history"
    history_dir.mkdir(parents=True)
    (history_dir / "rename.jsonl").write_text(
        json.dumps(
            {"id": 1, "data": {"rename_map": {"a.txt": "b.txt"}}, "undone": True},
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )
    return {
        "pdf": str(pdf),
        "pdf_copy": str(pdf_copy),
        "docx": str(docx),
        "tables": str(tables),
        "document": str(document),
        "txt": str(txt),
    }


def _pure_report(portable: Path, files: dict[str, str]) -> ps.SelftestReport:
    def entry(name: str, artifacts: dict[str, ps.ArtifactJSON]) -> dict:
        return {
            "name": name,
            "status": "pass",
            "detail": "",
            "duration_s": 0.0,
            "artifacts": artifacts,
        }

    scene = portable / "selftest-work" / "scene"
    return ps.SelftestReport(
        mode="pure",
        exit_code=0,
        data_root=str(portable / ".file_toolbox"),
        history_dir=str(portable / ".file_toolbox" / "history"),
        workdir=str(portable / "selftest-work"),
        evidence_missing=False,
        scenarios=[
            entry("pages", {}),
            entry(
                "rename",
                {
                    "rename_map": {str(scene / "a.txt"): str(scene / "b.txt")},
                    "record_id": 1,
                    "workdir": str(portable / "selftest-work" / "rename"),
                    "history_tool": "rename",
                },
            ),
            entry("pdf_pure", {"pdfs": [files["pdf"]], "pdf_copy": files["pdf_copy"]}),
            entry(
                "markdown",
                {
                    "docx_output": files["docx"],
                    "tables_output": files["tables"],
                    "document_output": files["document"],
                },
            ),
            entry("replace_pure", {"txt": files["txt"], "md": files["txt"]}),
            entry("update_mutex", {"locked_then_unlocked": True, "download_calls": 1}),
            entry("pdf_cancel", {"outputs": [files["pdf"]]}),
        ],
    )


def _run1_scene(portable: Path) -> tuple[dict[str, str], ps.SelftestReport]:
    """在 portable/selftest-work/scene 内建立 run1 全套虚构产物与历史。"""
    scene = portable / "selftest-work" / "scene"
    files = _fixture_outputs(scene)
    history_dir = portable / ".file_toolbox" / "history"
    history_dir.mkdir(parents=True, exist_ok=True)
    (scene / "a.txt").write_text("content of a.txt\n", encoding="utf-8")
    (scene / "b.txt").write_text("content of a.txt\n", encoding="utf-8")
    history_dir = portable / ".file_toolbox" / "history"
    (history_dir / "rename.jsonl").write_text(
        json.dumps(
            {
                "id": 1,
                "data": {"rename_map": {str(scene / "a.txt"): str(scene / "b.txt")}},
                "undone": False,
            },
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )
    return files, _pure_report(portable, files)


def _blank_pdf(path: Path) -> str:
    """写一个真实可解析的空白单页 PDF,返回路径字符串。"""
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with path.open("wb") as stream:
        writer.write(stream)
    return str(path)


def _full_scene(portable: Path) -> tuple[dict[str, str], ps.SelftestReport]:
    """在 pure 基础上补齐 full 三场景的虚构产物(docx/xlsx 两条 Office PDF)。"""
    from openpyxl import Workbook

    files, report = _run1_scene(portable)
    office = portable / "selftest-work" / "pdf-office"
    office.mkdir(parents=True, exist_ok=True)
    docx_pdf = _blank_pdf(office / "office-docx-source.pdf")
    xlsx_pdf = _blank_pdf(office / "office-xlsx-source.pdf")
    replaced = office / "replaced.docx"
    with zipfile.ZipFile(replaced, "w") as package:
        package.writestr("word/document.xml", "<w:t>office-replaced</w:t>")

    def _workbook(path: Path, *sheets: str) -> str:
        workbook = Workbook()
        workbook.remove(workbook.active)
        for sheet in sheets:
            workbook.create_sheet(sheet)
        workbook.save(path)
        workbook.close()
        return str(path)

    plain_output = _workbook(office / "plain.xlsx", "甲明细", "乙明细")
    roster_output = _workbook(office / "roster.xlsx", "丙明细", "丁明细")
    placeholders = {
        key: str(office / f"{key}.bin") for key in ("source", "source_roster", "template", "roster")
    }
    for value in placeholders.values():
        Path(value).write_bytes(b"fixture")
    report["mode"] = "full"
    report["scenarios"].extend(
        [
            _scenario_dict("pdf_office", "pass", artifacts={"docx": docx_pdf, "xlsx": xlsx_pdf}),
            _scenario_dict("replace_office", "pass", artifacts={"docx": str(replaced)}),
            _scenario_dict(
                "attendance",
                "pass",
                artifacts={
                    **placeholders,
                    "plain_output": plain_output,
                    "roster_output": roster_output,
                },
            ),
        ]
    )
    return files, report


def _full_missing_scene(portable: Path) -> ps.SelftestReport:
    """full 一致性缺证据报告:纯场景 pass,Office 三场景 evidence_missing,exit 3。"""
    _files, report = _run1_scene(portable)
    report["mode"] = "full"
    report["exit_code"] = ps.EXIT_EVIDENCE_MISSING
    report["evidence_missing"] = True
    report["scenarios"].extend(
        _scenario_dict(name, "evidence_missing", "Office 预筛不可用")
        for name in ("pdf_office", "replace_office", "attendance")
    )
    return report


def _reopen_report_after_undo(portable: Path) -> ps.SelftestReport:
    """run2:撤销已发生(新名消失/原名恢复/历史标记 undone)后的报告。"""
    scene = portable / "selftest-work" / "scene"
    (scene / "b.txt").unlink(missing_ok=True)
    (scene / "a.txt").write_text("content of a.txt\n", encoding="utf-8")
    history = portable / ".file_toolbox" / "history" / "rename.jsonl"
    history.write_text(
        json.dumps(
            {
                "id": 1,
                "data": {"rename_map": {str(scene / "a.txt"): str(scene / "b.txt")}},
                "undone": True,
            },
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )
    return ps.SelftestReport(
        mode="reopen",
        exit_code=0,
        data_root=str(portable / ".file_toolbox"),
        history_dir=str(portable / ".file_toolbox" / "history"),
        workdir=str(portable / "selftest-work"),
        evidence_missing=False,
        scenarios=[
            _scenario_dict(
                "reopen_history_undo", "pass", artifacts={"restored": [str(scene / "a.txt")]}
            )
        ],
    )


def test_verify_run1_reads_back_real_outputs(tmp_path):
    portable = tmp_path / "portable"
    files, report = _run1_scene(portable)
    checks = ps._verify_run1(report, "pure")
    assert any("pdf 页数" in check for check in checks)
    assert "update mutex" in checks


def test_verify_run1_fails_on_missing_rename_output(tmp_path):
    portable = tmp_path / "portable"
    _files, report = _run1_scene(portable)
    (portable / "selftest-work" / "scene" / "b.txt").unlink()
    with pytest.raises(AssertionError, match="重命名文件缺失"):
        ps._verify_run1(report, "pure")


# ---------------------------------------------------------------------------
# F3 残留:full 读回必须要求 docx/xlsx 两条 Office PDF 工件,拒绝空/缺一/越界/无效
# ---------------------------------------------------------------------------


def test_verify_run1_full_requires_both_office_pdfs(tmp_path):
    """full 正例:docx 与 xlsx 两条 Office PDF 都必须读回页数。"""
    portable = tmp_path / "portable"
    _files, report = _full_scene(portable)
    checks = ps._verify_run1(report, "full")
    assert any("office pdf docx" in check for check in checks)
    assert any("office pdf xlsx" in check for check in checks)


def test_verify_run1_full_rejects_empty_office_artifacts(tmp_path):
    """负例:pdf_office.artifacts 为空 dict 时 full 读回必须失败(F3 残留)。"""
    portable = tmp_path / "portable"
    _files, report = _full_scene(portable)
    ps._scenario(report, "pdf_office")["artifacts"] = {}
    with pytest.raises(AssertionError, match="不是路径字符串"):
        ps._verify_run1(report, "full")


def test_verify_run1_full_rejects_single_office_artifact(tmp_path):
    """负例:仅剩 docx 一条 Office PDF 时 full 读回必须失败(缺 xlsx)。"""
    portable = tmp_path / "portable"
    _files, report = _full_scene(portable)
    artifacts = ps._scenario(report, "pdf_office")["artifacts"]
    del artifacts["xlsx"]
    with pytest.raises(AssertionError, match="不是路径字符串"):
        ps._verify_run1(report, "full")


def test_verify_run1_full_rejects_office_pdf_outside_workdir(tmp_path):
    """负例:Office PDF 路径越出本次运行 workdir 时必须失败。"""
    portable = tmp_path / "portable"
    _files, report = _full_scene(portable)
    outside = _blank_pdf(tmp_path / "outside.pdf")
    ps._scenario(report, "pdf_office")["artifacts"]["xlsx"] = outside
    with pytest.raises(AssertionError, match="越出本次运行范围"):
        ps._verify_run1(report, "full")


def test_verify_run1_full_rejects_invalid_office_pdf(tmp_path):
    """负例:工件指向无法解析的伪 PDF 时必须失败(有效 PDF 校验)。"""
    portable = tmp_path / "portable"
    _files, report = _full_scene(portable)
    fake = portable / "selftest-work" / "pdf-office" / "fake.pdf"
    fake.write_bytes(b"not a pdf")
    ps._scenario(report, "pdf_office")["artifacts"]["xlsx"] = str(fake)
    with pytest.raises(AssertionError, match="PDF 无法解析"):
        ps._verify_run1(report, "full")


def test_verify_run1_full_rejects_reused_pure_pdf(tmp_path):
    """负例:docx/xlsx 复用 pure 场景的 PDF(不在 pdf-office 子目录)→ 失败。

    仅约束整个 run 时同一路径既当 pure 证据又当 office 证据可混过;收紧到
    对应场景子目录后,跨场景复用必须被拒绝。
    """
    portable = tmp_path / "portable"
    files, report = _full_scene(portable)
    artifacts = ps._scenario(report, "pdf_office")["artifacts"]
    artifacts["docx"] = files["pdf"]
    artifacts["xlsx"] = files["pdf"]
    with pytest.raises(AssertionError, match="越出本次运行范围"):
        ps._verify_run1(report, "full")


def test_verify_run1_full_rejects_duplicate_office_pdf_path(tmp_path):
    """负例:docx/xlsx 都指向 pdf-office 内同一文件 → 失败(不是两条证据)。"""
    portable = tmp_path / "portable"
    _files, report = _full_scene(portable)
    artifacts = ps._scenario(report, "pdf_office")["artifacts"]
    artifacts["xlsx"] = artifacts["docx"]
    with pytest.raises(AssertionError, match="指向同一文件"):
        ps._verify_run1(report, "full")


# ---------------------------------------------------------------------------
# F3 残留:报告顶层 evidence_missing 标志与 exit/status 不得矛盾
# ---------------------------------------------------------------------------


def test_load_report_rejects_invalid_evidence_missing_flag(tmp_path):
    """顶层 evidence_missing 非 bool → 拒绝整份报告。"""
    report_path = _valid_report(tmp_path, evidence_missing="yes")
    with pytest.raises(ValueError, match="evidence_missing"):
        ps._load_report(report_path)


def test_validate_report_accepts_consistent_report(tmp_path):
    portable = tmp_path / "portable"
    _files, report = _run1_scene(portable)
    ps._validate_report(report, expected_mode="pure", child_code=0, portable_root=portable)


def test_validate_report_rejects_flag_exit_contradiction(tmp_path):
    """负例:全 pass/exit0 但顶层 evidence_missing=true → 矛盾拒绝。"""
    portable = tmp_path / "portable"
    _files, report = _run1_scene(portable)
    report["evidence_missing"] = True
    with pytest.raises(AssertionError, match="矛盾"):
        ps._validate_report(report, expected_mode="pure", child_code=0, portable_root=portable)


def test_validate_report_rejects_missing_scenario_with_pass_exit(tmp_path):
    """负例:存在 evidence_missing 场景但 exit_code 仍为 PASS → 矛盾拒绝。"""
    portable = tmp_path / "portable"
    _files, report = _run1_scene(portable)
    ps._scenario(report, "pages")["status"] = "evidence_missing"
    with pytest.raises(AssertionError, match="evidence_missing"):
        ps._validate_report(report, expected_mode="pure", child_code=0, portable_root=portable)


def test_verify_run2_binds_exact_record_and_content(tmp_path):
    portable = tmp_path / "portable"
    _files, run1 = _run1_scene(portable)
    reopen = _reopen_report_after_undo(portable)
    checks = ps._verify_run2(reopen, run1)
    assert any(check.startswith("undo:") for check in checks)
    assert any("history undone" in check for check in checks)
    # 原样本内容未恢复 → 失败(不用任意 undone 行替代)
    scene = portable / "selftest-work" / "scene"
    reopen_tampered = _reopen_report_after_undo(portable)
    (scene / "a.txt").write_text("被篡改的内容\n", encoding="utf-8")
    with pytest.raises(AssertionError, match="原样本内容未恢复"):
        ps._verify_run2(reopen_tampered, run1)


# ---------------------------------------------------------------------------
# 启动器:超时不杀进程、fail closed 不继续第二次启动
# ---------------------------------------------------------------------------


class _FakeProcess:
    def __init__(self) -> None:
        self.pid = 4242
        self.killed = False

    def wait(self, timeout: float | None = None) -> int:
        raise subprocess.TimeoutExpired(cmd=["FileToolbox.exe"], timeout=timeout)

    def kill(self) -> None:  # 新契约:超时路径不得调用
        self.killed = True


def test_launch_exe_timeout_keeps_process_alive(monkeypatch, tmp_path, capsys):
    """超时契约:不 kill(可能持有真实 Office 会话),报告 owned pid,返回 4。"""
    fake = _FakeProcess()
    monkeypatch.setattr(subprocess, "Popen", lambda *a, **k: fake)
    code = ps._launch_exe(
        tmp_path / "FileToolbox.exe",
        "full",
        tmp_path / "report.json",
        visible=False,
        timeout_s=1,
        stdout_log=tmp_path / "log.txt",
    )
    assert code == ps.EXIT_TIMEOUT
    assert fake.killed is False
    err = capsys.readouterr().err
    assert "pid=4242" in err and "不终止进程" in err


def _install_fake_launcher(monkeypatch, codes: list[int], reports: list[ps.SelftestReport]):
    """替身启动器:按序写报告并返回预置退出码;记录调用模式。"""
    calls: list[str] = []

    def fake_launch(exe, mode, report, *, visible, timeout_s, stdout_log):
        calls.append(mode)
        current = reports[min(len(calls) - 1, len(reports) - 1)]
        report.write_text(json.dumps(current, ensure_ascii=False), encoding="utf-8")
        return codes[min(len(calls) - 1, len(codes) - 1)]

    monkeypatch.setattr(ps, "_launch_exe", fake_launch)
    return calls


def _portable_zip_with_exe(tmp_path: Path) -> Path:
    zip_path = tmp_path / "FileToolbox-v0.0.1-win-x64.zip"
    _make_zip(zip_path, {"current/FileToolbox.exe": b"MZ"})
    return zip_path


def test_main_end_to_end_pass(monkeypatch, tmp_path):
    monkeypatch.setattr(ps, "SELFTEST_RUN_ROOT", tmp_path)
    calls: list[str] = []

    def fake_launch(exe, mode, report, *, visible, timeout_s, stdout_log):
        calls.append(mode)
        portable = exe.parent.parent
        if mode == "pure":
            _files, run1 = _run1_scene(portable)
            report.write_text(json.dumps(run1, ensure_ascii=False), encoding="utf-8")
            return 0
        reopen = _reopen_report_after_undo(portable)
        report.write_text(json.dumps(reopen, ensure_ascii=False), encoding="utf-8")
        return 0

    monkeypatch.setattr(ps, "_launch_exe", fake_launch)
    monkeypatch.chdir(tmp_path)
    code = ps.main(["--zip", str(_portable_zip_with_exe(tmp_path)), "--mode", "pure"])
    assert code == ps.EXIT_PASS
    assert calls == ["pure", "reopen"]
    summary = json.loads(next(tmp_path.glob("run-*/summary.json")).read_text(encoding="utf-8"))
    assert summary["result"] == "PASS"


@pytest.mark.parametrize(
    ("mutation", "match"),
    [
        pytest.param(
            lambda report: report["scenarios"].pop(
                next(i for i, s in enumerate(report["scenarios"]) if s["name"] == "pages")
            ),
            "场景集合不符",
            id="missing-pages",
        ),
        pytest.param(
            lambda report: report["scenarios"].pop(
                next(i for i, s in enumerate(report["scenarios"]) if s["name"] == "pdf_cancel")
            ),
            "场景集合不符",
            id="missing-pdf-cancel",
        ),
        pytest.param(
            lambda report: report.__setitem__("exit_code", 1),
            "exit_code 与子进程",
            id="report-exit-mismatch",
        ),
        pytest.param(
            lambda report: report["scenarios"].append(
                {
                    "name": "close",
                    "status": "fail",
                    "detail": "x",
                    "duration_s": 0.0,
                    "artifacts": {},
                }
            ),
            "场景集合不符|存在未通过场景",
            id="close-fail",
        ),
        pytest.param(
            lambda report: report.__setitem__(
                "data_root",
                str(Path(report["data_root"]).parent / "portable-sibling" / ".file_toolbox"),
            ),
            "data_root 未精确绑定",
            id="wrong-data-root",
        ),
        pytest.param(
            lambda report: next(s for s in report["scenarios"] if s["name"] == "rename")[
                "artifacts"
            ].__setitem__("rename_map", {}),
            "rename_map 为空",
            id="empty-rename-map",
        ),
        pytest.param(
            lambda report: next(s for s in report["scenarios"] if s["name"] == "pdf_pure")[
                "artifacts"
            ].__setitem__("pdfs", []),
            "pdfs 为空",
            id="empty-pdfs",
        ),
        pytest.param(
            lambda report: next(s for s in report["scenarios"] if s["name"] == "pdf_pure")[
                "artifacts"
            ].__setitem__("pdf_copy", str(Path("C:/elsewhere/evil.pdf"))),
            "越出本次运行范围",
            id="artifact-outside-scope",
        ),
    ],
)
def test_main_fails_closed_on_invalid_run1_report(monkeypatch, tmp_path, mutation, match):
    """F3 负例:结构/路径/集合任何遗假 → 非零,不写 PASS 摘要。"""
    monkeypatch.setattr(ps, "SELFTEST_RUN_ROOT", tmp_path)

    def fake_launch(exe, mode, report, *, visible, timeout_s, stdout_log):
        portable = exe.parent.parent
        _files, run1 = _run1_scene(portable)
        mutation(run1)
        report.write_text(json.dumps(run1, ensure_ascii=False), encoding="utf-8")
        return 0

    monkeypatch.setattr(ps, "_launch_exe", fake_launch)
    monkeypatch.chdir(tmp_path)
    code = ps.main(["--zip", str(_portable_zip_with_exe(tmp_path)), "--mode", "pure"])
    assert code == ps.EXIT_FAIL
    assert not list(tmp_path.glob("run-*/summary.json"))


def test_main_timeout_stops_before_reopen(monkeypatch, tmp_path, capsys):
    monkeypatch.setattr(ps, "SELFTEST_RUN_ROOT", tmp_path)
    empty_report = ps.SelftestReport(
        mode="pure",
        exit_code=0,
        data_root=str(tmp_path),
        history_dir=str(tmp_path),
        scenarios=[_scenario_dict("pages", "pass")],
    )
    calls = _install_fake_launcher(monkeypatch, [ps.EXIT_TIMEOUT], [empty_report])
    monkeypatch.chdir(tmp_path)
    code = ps.main(["--zip", str(_portable_zip_with_exe(tmp_path)), "--mode", "pure"])
    assert code == ps.EXIT_TIMEOUT
    assert calls == ["pure"]  # 不继续第二次启动抢占数据根


def test_main_evidence_missing_propagates(monkeypatch, tmp_path):
    """exit 3 + 一致的缺证据报告(mode/根/标志/场景全核对)→ 如实透传 3(P3)。"""
    monkeypatch.setattr(ps, "SELFTEST_RUN_ROOT", tmp_path)
    calls: list[str] = []

    def fake_launch(exe, mode, report, *, visible, timeout_s, stdout_log):
        calls.append(mode)
        run1 = _full_missing_scene(exe.parent.parent)
        report.write_text(json.dumps(run1, ensure_ascii=False), encoding="utf-8")
        return ps.EXIT_EVIDENCE_MISSING

    monkeypatch.setattr(ps, "_launch_exe", fake_launch)
    monkeypatch.chdir(tmp_path)
    code = ps.main(["--zip", str(_portable_zip_with_exe(tmp_path)), "--mode", "full"])
    assert code == ps.EXIT_EVIDENCE_MISSING
    assert calls == ["full"]


def test_main_evidence_missing_contradictory_report_fails(monkeypatch, tmp_path):
    """P3 回归:exit 3 但报告声称 exit_code=0/全 pass(矛盾)→ 按失败(1),
    不冒称"缺少证据"透传。"""
    monkeypatch.setattr(ps, "SELFTEST_RUN_ROOT", tmp_path)

    def fake_launch(exe, mode, report, *, visible, timeout_s, stdout_log):
        _files, run1 = _run1_scene(exe.parent.parent)  # exit_code=0/全 pass 的矛盾报告
        run1["mode"] = "full"
        report.write_text(json.dumps(run1, ensure_ascii=False), encoding="utf-8")
        return ps.EXIT_EVIDENCE_MISSING

    monkeypatch.setattr(ps, "_launch_exe", fake_launch)
    monkeypatch.chdir(tmp_path)
    code = ps.main(["--zip", str(_portable_zip_with_exe(tmp_path)), "--mode", "full"])
    assert code == ps.EXIT_FAIL


def test_validate_report_accepts_consistent_evidence_missing(tmp_path):
    """P3 单元:一致的缺证据报告(场景 evidence_missing + exit 3)通过校验。"""
    portable = tmp_path / "portable"
    report = _full_missing_scene(portable)
    ps._validate_report(
        report, expected_mode="full", child_code=ps.EXIT_EVIDENCE_MISSING, portable_root=portable
    )
