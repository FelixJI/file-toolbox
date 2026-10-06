"""markdown_convert 核心测试:xlsx(openpyxl 真读)、docx(zip/XML)、错误、取消、冲突、历史。

全部基于程序化生成的虚构 Markdown;docx 走真实内置 Pandoc 子进程。
"""

import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from file_toolbox.common.history import JsonHistoryStore
from file_toolbox.core import markdown_convert
from file_toolbox.core.markdown_convert import (
    EXCEL_CELL_MAX_CHARS,
    SUPPORTED_SUFFIXES,
    ConversionItem,
    ConversionResult,
    MarkdownConvertService,
)

MD_FULL = """# 标题一

开头一段,带 [链接](https://example.com/a)。

| 名称 | 编号 | 备注 |
|---|---|---|
| 苹果 | 007 | 甜 |
| 香蕉\\|梨 | =SUM(1,2) |  |

- 列表项一
- 列表项二

```text
| 伪表 | 不是表格 |
|---|---|
```

结尾段落。
"""


def _md(tmp_path: Path, name: str, content: str) -> Path:
    path = tmp_path / name
    path.write_text(content, encoding="utf-8")
    return path


def _cells(sheet) -> list[list[str]]:
    return [[str(c.value) if c.value is not None else "" for c in row] for row in sheet.iter_rows()]


# ==================== 契约 ====================


def test_contract_constants_and_result_shapes():
    assert SUPPORTED_SUFFIXES == (".md", ".markdown")
    item = ConversionItem(source=Path("a.md"))
    assert item.output is None and item.error == "" and item.skipped is False
    assert ConversionResult().success is False
    assert ConversionResult([ConversionItem(Path("a.md"), Path("a.docx"))]).success is True
    assert ConversionResult([ConversionItem(Path("a.md"))], cancelled=True).success is False


def test_invalid_target_or_excel_mode_rejected(tmp_path):
    src = _md(tmp_path, "a.md", "# t")
    with pytest.raises(ValueError, match="无效的目标格式"):
        MarkdownConvertService().convert([src], None, target="pdf")
    with pytest.raises(ValueError, match="无效的 Excel 模式"):
        MarkdownConvertService().convert([src], None, target="docx", excel_mode="bad")


# ==================== xlsx ====================


def test_xlsx_tables_mode_one_sheet_per_table(tmp_path):
    src = _md(tmp_path, "a.md", MD_FULL)
    result = MarkdownConvertService().convert([src], None, target="xlsx", excel_mode="tables")

    assert result.success
    output = result.items[0].output
    assert output == tmp_path / "a.xlsx"
    wb = load_workbook(output)
    assert wb.sheetnames == ["表1"]  # 代码块伪表格不产生第二个 sheet
    cells = _cells(wb["表1"])
    assert cells[0] == ["名称", "编号", "备注"]
    assert cells[1] == ["苹果", "007", "甜"]
    assert cells[2] == ["香蕉|梨", "=SUM(1,2)", ""]  # 转义竖线/公式样文本/空单元格
    assert wb["表1"]["B2"].number_format == "@"
    wb.close()


def test_xlsx_document_mode_body_sheet_in_reading_order(tmp_path):
    src = _md(tmp_path, "a.md", MD_FULL)
    result = MarkdownConvertService().convert([src], None, target="xlsx", excel_mode="document")

    assert result.success
    wb = load_workbook(result.items[0].output)
    assert wb.sheetnames == ["正文", "表1"]
    rows = _cells(wb["正文"])
    assert rows[0] == ["类型", "内容"]
    kinds = [row[0] for row in rows[1:]]
    assert kinds == ["标题1", "段落", "表格", "列表", "列表", "代码", "段落"]
    table_ref = rows[1:][2]
    assert table_ref == ["表格", "内容见工作表「表1」"]
    assert rows[1:][5] == ["代码", "| 伪表 | 不是表格 |\n|---|---|"]
    assert rows[1:][0] == ["标题1", "标题一"]
    wb.close()


def test_xlsx_document_mode_without_tables_still_writes_body(tmp_path):
    src = _md(tmp_path, "a.md", "# 只有标题\n\n纯正文。\n")
    result = MarkdownConvertService().convert([src], None, target="xlsx", excel_mode="document")

    assert result.success
    wb = load_workbook(result.items[0].output)
    assert wb.sheetnames == ["正文"]
    assert _cells(wb["正文"])[1] == ["标题1", "只有标题"]
    wb.close()


def test_xlsx_tables_mode_skips_file_without_tables(tmp_path):
    src = _md(tmp_path, "a.md", "# 无表格\n\n只有正文。\n")
    result = MarkdownConvertService().convert([src], None, target="xlsx", excel_mode="tables")

    assert result.success is False
    item = result.items[0]
    assert item.skipped and item.output is None
    assert "没有表格" in item.error
    assert not (tmp_path / "a.xlsx").exists()


def test_xlsx_cell_over_excel_limit_fails_explicitly(tmp_path, monkeypatch):
    monkeypatch.setattr(markdown_convert, "EXCEL_CELL_MAX_CHARS", 10)
    src = _md(tmp_path, "a.md", f"# {'x' * 11}\n")
    result = MarkdownConvertService().convert([src], None, target="xlsx", excel_mode="document")

    assert result.success is False
    assert "超过 Excel 上限" in result.items[0].error
    assert "拒绝截断" in result.items[0].error
    assert not (tmp_path / "a.xlsx").exists()


def test_xlsx_markdown_suffix_and_output_dir(tmp_path):
    src = _md(tmp_path, "doc.markdown", "| a | b |\n|---|---|\n| 1 | 2 |\n")
    out_dir = tmp_path / "out"
    result = MarkdownConvertService().convert([src], out_dir, target="xlsx")

    assert result.success
    output = result.items[0].output
    assert output == out_dir / "doc.xlsx"
    wb = load_workbook(output)
    assert _cells(wb["表1"]) == [["a", "b"], ["1", "2"]]
    wb.close()


def test_xlsx_cell_limit_is_real_excel_bound(tmp_path):
    """契约上限是 Excel 真实边界 32767(防将来改错常量)。"""
    assert EXCEL_CELL_MAX_CHARS == 32767


def test_xlsx_column_limit_guard(tmp_path, monkeypatch):
    """openpyxl 可写出超 16384 列的无效工作簿,这里显式拦截。"""
    monkeypatch.setattr(markdown_convert, "EXCEL_MAX_COLUMNS", 2)
    src = _md(tmp_path, "a.md", "| a | b | c |\n|---|---|---|\n| 1 | 2 | 3 |\n")
    result = MarkdownConvertService().convert([src], None, target="xlsx")

    assert result.success is False
    assert "列数 3 超过 Excel 上限 2" in result.items[0].error
    assert not (tmp_path / "a.xlsx").exists()


# ==================== docx(真实内置 Pandoc) ====================


def test_docx_real_conversion_keeps_structure_and_math(tmp_path):
    src = _md(tmp_path, "a.md", MD_FULL + "\n数学 $x^2+y$ 收尾。\n")
    result = MarkdownConvertService().convert([src], None, target="docx")

    assert result.success
    output = result.items[0].output
    assert output == tmp_path / "a.docx"
    assert output.read_bytes()[:2] == b"PK"
    assert zipfile.is_zipfile(output)
    with zipfile.ZipFile(output) as zf:
        xml = zf.read("word/document.xml").decode("utf-8")
        assert "标题一" in xml
        rels = zf.read("word/_rels/document.xml.rels").decode("utf-8")
        assert "https://example.com/a" in rels  # 超链接目标
        assert "oMath" in xml  # 数学公式保留(OMML)
        assert "列表项一" in xml
        assert "伪表" in xml  # 代码块保留


def test_docx_missing_image_fails_explicitly(tmp_path):
    src = _md(tmp_path, "a.md", "![图](missing.png)\n")
    result = MarkdownConvertService().convert([src], None, target="docx")

    assert result.success is False
    item = result.items[0]
    assert item.output is None and not item.skipped
    assert "Pandoc 转换失败" in item.error
    assert "Could not fetch resource" in item.error
    assert not (tmp_path / "a.docx").exists()


def test_docx_pandoc_start_failure_recorded(tmp_path, monkeypatch):
    fake = tmp_path / "pandoc.exe"
    fake.write_bytes(b"not an executable")
    monkeypatch.setattr(markdown_convert, "_locate_pandoc", lambda: fake)
    src = _md(tmp_path, "a.md", "# t")
    result = MarkdownConvertService().convert([src], None, target="docx")

    assert result.success is False
    assert "无法启动内置 Pandoc" in result.items[0].error


def test_docx_pandoc_missing_package_or_path(tmp_path, monkeypatch):
    src = _md(tmp_path, "a.md", "# t")
    monkeypatch.setattr(markdown_convert, "_find_spec", lambda name: None)
    with pytest.raises(ImportError, match="pypandoc"):
        MarkdownConvertService().convert([src], None, target="docx")

    monkeypatch.setattr(markdown_convert, "is_packaged_runtime", lambda: True)
    with pytest.raises(ImportError, match="未找到内置 Pandoc"):
        MarkdownConvertService().convert([src], None, target="docx")


# ==================== 批量错误 / 取消 / 冲突 ====================


def test_batch_records_per_file_errors_and_continues(tmp_path):
    good = _md(tmp_path, "a.md", "| a |\n|---|\n| 1 |\n")
    unsupported = tmp_path / "b.txt"
    unsupported.write_text("x", encoding="utf-8")
    undecodable = tmp_path / "c.md"
    undecodable.write_bytes("中文".encode("gbk"))
    missing = tmp_path / "d.md"

    result = MarkdownConvertService().convert(
        [good, unsupported, undecodable, missing], None, target="xlsx"
    )

    assert result.success  # 有效文件仍产出
    assert result.items[0].output == tmp_path / "a.xlsx"
    assert "不支持的格式 .txt" in result.items[1].error
    assert "无法按 UTF-8 解码" in result.items[2].error
    assert "无法读取源文件" in result.items[3].error
    assert all(not item.skipped for item in result.items[1:])


def test_cancel_before_any_file(tmp_path):
    a = _md(tmp_path, "a.md", "| a |\n|---|\n| 1 |\n")
    b = _md(tmp_path, "b.md", "| a |\n|---|\n| 2 |\n")
    result = MarkdownConvertService().convert(
        [a, b], None, target="xlsx", cancel_check=lambda: True
    )

    assert result.cancelled and not result.success
    assert all(item.skipped and item.output is None for item in result.items)
    assert not any(tmp_path.glob("*.xlsx"))


def test_cancel_at_commit_boundary_keeps_no_partial_output(tmp_path, monkeypatch):
    a = _md(tmp_path, "a.md", "| a |\n|---|\n| 1 |\n")
    b = _md(tmp_path, "b.md", "| a |\n|---|\n| 2 |\n")
    calls = {"count": 0}

    def cancel() -> bool:
        calls["count"] += 1
        return calls["count"] > 1  # 第 1 次在文件间放行,第 2 次在提交前取消

    result = MarkdownConvertService().convert([a, b], None, target="xlsx", cancel_check=cancel)

    assert result.cancelled and not result.success
    assert result.items[0].skipped and "未写出输出" in result.items[0].error
    assert result.items[1].skipped and "未处理" in result.items[1].error
    assert not any(tmp_path.glob("*.xlsx"))  # 提交前取消不落半成品


def test_cancel_after_first_file_keeps_completed_outputs(tmp_path):
    a = _md(tmp_path, "a.md", "| a |\n|---|\n| 1 |\n")
    b = _md(tmp_path, "b.md", "| a |\n|---|\n| 2 |\n")
    state = {"file": 0}

    def cancel() -> bool:
        # 以首个输出已提交落盘为准:提交前检查与下一文件入口检查都能看到它,
        # 取消恰好发生在第二个文件入口(不与提交前检查混淆)。
        return (tmp_path / "a.xlsx").exists()

    def progress(current: int, total: int, message: str) -> None:
        state["file"] = current

    result = MarkdownConvertService().convert(
        [a, b], None, target="xlsx", cancel_check=cancel, progress_callback=progress
    )

    assert result.cancelled and not result.success  # 契约:取消时 success=False,但已完成输出保留
    assert result.items[0].output == tmp_path / "a.xlsx"
    assert (tmp_path / "a.xlsx").is_file()
    assert result.items[1].skipped and result.items[1].output is None
    assert "未处理" in result.items[1].error
    assert not (tmp_path / "b.xlsx").exists()


def test_output_conflict_auto_numbered_never_overwrites(tmp_path):
    src = _md(tmp_path, "a.md", "| a |\n|---|\n| 1 |\n")
    existing = tmp_path / "a.xlsx"
    existing.write_text("precious", encoding="utf-8")

    result = MarkdownConvertService().convert([src], None, target="xlsx")

    assert result.success
    assert existing.read_text(encoding="utf-8") == "precious"
    assert result.items[0].output == tmp_path / "a_1.xlsx"


# ==================== 历史 ====================


def test_history_recorded_on_success_shape_matches_gui(tmp_path):
    store = JsonHistoryStore(tmp_path / "h")
    src = _md(tmp_path, "a.md", "| a |\n|---|\n| 1 |\n")
    MarkdownConvertService(history_store=store).convert([src], None, target="xlsx")

    records = store.get_records("markdown_convert")
    assert len(records) == 1
    data = records[0]["data"]
    assert data["success"] == 1
    assert data["file_count"] == 1
    assert data["target"] == "xlsx"
    assert data["excel_mode"] == "tables"
    assert data["outputs"] == [str(tmp_path / "a.xlsx")]


def test_history_skipped_without_outputs_records_cancelled_batch(tmp_path):
    store = JsonHistoryStore(tmp_path / "h")
    svc = MarkdownConvertService(history_store=store)
    src = _md(tmp_path, "a.md", "无表格")
    svc.convert([src], None, target="xlsx")  # 全跳过:无输出 → 无记录
    svc.convert([src], None, target="xlsx", cancel_check=lambda: True)  # 首文件前取消
    assert store.get_records("markdown_convert") == []

    a = _md(tmp_path, "done.md", "| a |\n|---|\n| 1 |\n")
    b = _md(tmp_path, "pending.md", "| a |\n|---|\n| 2 |\n")
    done = tmp_path / "done.xlsx"
    svc.convert(
        [a, b],
        None,
        target="xlsx",
        cancel_check=lambda: done.exists(),  # 首个输出提交后取消
    )  # 已完成输出保留,历史如实标记 cancelled

    records = store.get_records("markdown_convert")
    assert len(records) == 1
    data = records[0]["data"]
    assert data["cancelled"] is True
    assert data["success"] == 1
    assert data["file_count"] == 2


def test_history_failure_wrapped_but_result_preserved(tmp_path):
    """历史写失败经 preserve_history_result 包装,已写出的输出保留。"""
    from file_toolbox.common.operation_errors import HistorySaveError

    store = JsonHistoryStore(tmp_path / "h")
    src = _md(tmp_path, "a.md", "| a |\n|---|\n| 1 |\n")
    original = store.add_record

    def broken(tool: str, data: dict[str, object]) -> int:
        raise OSError("history denied")

    store.add_record = broken  # type: ignore[method-assign]
    try:
        with pytest.raises(HistorySaveError) as excinfo:
            MarkdownConvertService(history_store=store).convert([src], None, target="xlsx")
        assert excinfo.value.result.success
        assert (tmp_path / "a.xlsx").is_file()
    finally:
        store.add_record = original  # type: ignore[method-assign]
