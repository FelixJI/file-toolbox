"""计划排布核心测试:解析、排布计算、模板与生成渲染。

全部基于程序化生成的虚构 xlsx(复用 conftest 的 make_xlsx 工厂),不触发 COM。
"""

import calendar
from datetime import date, datetime
from pathlib import Path
from unittest.mock import MagicMock

import pytest
from openpyxl import load_workbook

from file_toolbox.core.plan_schedule import (
    SHEET_NAME,
    PlanScheduleService,
    ScheduleOptions,
    ScheduleResult,
)
from file_toolbox.core.plan_schedule.constants import (
    CELL_NAME,
    DAY_COLUMN_WIDTH,
    ITEM_FILLS,
    NAME_MODE_DAY_WIDTH_MAX,
    WEEKEND_FILL,
)
from file_toolbox.core.plan_schedule.service import _EXCEL_SERIAL_EPOCH
from file_toolbox.core.plan_schedule.types import InvalidRow, MonthPlan, PlanItem


@pytest.fixture
def svc() -> PlanScheduleService:
    return PlanScheduleService()


# ==================== 解析 ====================


def test_parse_reads_datetime_and_date_cells(svc, make_xlsx):
    """datetime/date 单元格直接解析;行号为 1-based Excel 行号。"""
    src = make_xlsx(
        "list.xlsx",
        {
            "Sheet1": [
                ["项点名称", "起始日期", "终止日期"],
                ["A", datetime(2026, 9, 17, 8, 30), date(2026, 9, 21)],
            ]
        },
    )
    items, invalid = svc.parse(src)
    assert invalid == []
    assert len(items) == 1
    assert items[0] == PlanItem("A", date(2026, 9, 17), date(2026, 9, 21), row=2)
    assert items[0].days == 5


def test_parse_string_dates_full_and_no_year(svc, make_xlsx):
    """字符串日期:含年多格式;缺年串按 default_year 补全。"""
    src = make_xlsx(
        "list.xlsx",
        {
            "S": [
                ["项点名称", "起始日期", "终止日期"],
                ["A", "2026-09-17", "2026年9月21日"],
                ["B", "9/22", "9.28"],
            ]
        },
    )
    items, invalid = svc.parse(src, ScheduleOptions(default_year=2026))
    assert invalid == []
    assert items[0].start == date(2026, 9, 17)
    assert items[0].end == date(2026, 9, 21)
    assert items[1] == PlanItem("B", date(2026, 9, 22), date(2026, 9, 28), row=3)


def test_parse_excel_serial_numbers(svc, make_xlsx):
    """裸数字按 Excel 1900 序列日期换算。"""
    serial_start = (date(2026, 9, 17) - _EXCEL_SERIAL_EPOCH).days
    serial_end = (date(2026, 9, 21) - _EXCEL_SERIAL_EPOCH).days
    src = make_xlsx(
        "list.xlsx",
        {"S": [["项点名称", "起始日期", "终止日期"], ["A", serial_start, serial_end]]},
    )
    items, invalid = svc.parse(src)
    assert invalid == []
    assert items[0].start == date(2026, 9, 17)
    assert items[0].end == date(2026, 9, 21)


def test_parse_invalid_rows_keep_going(svc, make_xlsx):
    """行级问题进 invalid 不中断:缺名称/坏日期/终早于起。"""
    src = make_xlsx(
        "list.xlsx",
        {
            "S": [
                ["项点名称", "起始日期", "终止日期"],
                [None, date(2026, 9, 1), date(2026, 9, 2)],  # 缺名称
                ["坏日期", "不清楚", date(2026, 9, 2)],  # 起始无法识别
                ["缺终止", date(2026, 9, 1), None],  # 终止无法识别
                ["倒挂", date(2026, 9, 10), date(2026, 9, 1)],  # 终止早于起始
                ["正常", date(2026, 9, 17), date(2026, 9, 21)],
            ]
        },
    )
    items, invalid = svc.parse(src)
    assert [it.name for it in items] == ["正常"]
    assert [inv.row for inv in invalid] == [2, 3, 4, 5]
    assert invalid[0].error == "缺少项点名称"
    assert "起始日期无法识别" in invalid[1].error
    assert "终止日期无法识别" in invalid[2].error
    assert "早于" in invalid[3].error


def test_parse_skips_blank_rows_and_normalizes_name(svc, make_xlsx):
    """整行空跳过;数字名称转文本(123.0 -> "123")。"""
    src = make_xlsx(
        "list.xlsx",
        {
            "S": [
                ["项点名称", "起始日期", "终止日期"],
                [None, None, None],
                ["", "  ", None],
                [123.0, "2026-09-17", "2026-09-18"],
            ]
        },
    )
    items, invalid = svc.parse(src)
    assert invalid == []
    assert [it.name for it in items] == ["123"]


def test_parse_header_aliases_offset_and_order(svc, make_xlsx):
    """表头别名(名称/开始/结束)、表头不在首行、列序变化均可定位。"""
    src = make_xlsx(
        "list.xlsx",
        {
            "S": [
                ["交付计划清单(标题行)"],
                ["终止日期", "名称", "开始日期"],
                ["2026-09-21", "A", "2026-09-17"],
            ]
        },
    )
    items, invalid = svc.parse(src)
    assert invalid == []
    assert items[0].name == "A"
    assert items[0].start == date(2026, 9, 17)
    assert items[0].end == date(2026, 9, 21)
    assert items[0].row == 3


def test_parse_header_not_found_raises(svc, make_xlsx):
    """缺少任一列表头 → ValueError 说明所需列。"""
    src = make_xlsx("list.xlsx", {"S": [["项点名称", "起始日期", "别的"]]})
    with pytest.raises(ValueError, match="未找到表头"):
        svc.parse(src)


def test_parse_unsupported_suffix_raises(svc, tmp_path):
    f = tmp_path / "list.csv"
    f.write_text("x")
    with pytest.raises(ValueError, match="不支持的格式"):
        svc.parse(f)


def test_parse_missing_file_raises(svc, tmp_path):
    with pytest.raises(ValueError, match="不存在"):
        svc.parse(tmp_path / "missing.xlsx")


def test_parse_unreadable_file_raises(svc, tmp_path):
    bad = tmp_path / "bad.xlsx"
    bad.write_bytes(b"not an excel file")
    with pytest.raises(ValueError, match="无法读取"):
        svc.parse(bad)


def test_parse_date_direct_edge_branches(svc):
    """_parse_date 直接分支:None/bool/空串/未知类型/非正序号 → None。"""
    assert svc._parse_date(None, 2026) is None
    assert svc._parse_date(True, 2026) is None
    assert svc._parse_date("  ", 2026) is None
    assert svc._parse_date(object(), 2026) is None
    assert svc._parse_date(0, 2026) is None
    assert svc._parse_date(-1, 2026) is None


def test_parse_date_out_of_range_numbers_return_none(svc):
    """超出 date/timedelta 范围或非有限的数值 → None(行级无效,不抛 OverflowError)。"""
    assert svc._parse_date(20260917, 2026) is None  # 越界整数,不按 YYYYMMDD 解释
    assert svc._parse_date(20260917.0, 2026) is None
    assert svc._parse_date(10**12, 2026) is None
    assert svc._parse_date(float("inf"), 2026) is None
    assert svc._parse_date(float("-inf"), 2026) is None
    assert svc._parse_date(float("nan"), 2026) is None


def test_generate_continues_after_out_of_range_numeric_row(svc, make_xlsx, tmp_path):
    """越界数值日期(20260917/20260918)只作废该行,后续有效行照常生成输出。"""
    src = make_xlsx(
        "list.xlsx",
        {
            "S": [
                ["项点名称", "起始日期", "终止日期"],
                ["坏数值", 20260917, 20260918],
                ["正常", date(2026, 9, 17), date(2026, 9, 21)],
            ]
        },
    )
    result = svc.generate(src, tmp_path / "out.xlsx")
    assert result.success
    assert result.output is not None and result.output.is_file()
    assert [it.name for it in result.items] == ["正常"]
    assert [inv.row for inv in result.invalid] == [2]
    assert result.invalid[0].error == "起始日期无法识别:20260917"


# ==================== 排布计算 ====================


def _item(name: str, start: date, end: date) -> PlanItem:
    return PlanItem(name, start, end)


def test_plan_empty(svc):
    assert svc.plan([]) == []


def test_plan_months_weekends_and_days(svc):
    """月份覆盖首末;weekends 恰为周六周日;days 与日历一致。"""
    months = svc.plan([_item("A", date(2026, 9, 17), date(2026, 12, 5))])
    assert [(m.year, m.month) for m in months] == [
        (2026, 9),
        (2026, 10),
        (2026, 11),
        (2026, 12),
    ]
    for m in months:
        assert m.days == calendar.monthrange(m.year, m.month)[1]
        expected = {d for d in range(1, m.days + 1) if date(m.year, m.month, d).weekday() >= 5}
        assert set(m.weekends) == expected
        assert all(date(m.year, m.month, d).weekday() >= 5 for d in m.weekends)


def test_plan_cross_month_numbering_continues(svc):
    """项点跨月:9 月块编 1..3,10 月块接续 4..6。"""
    months = svc.plan([_item("A", date(2026, 9, 28), date(2026, 10, 3))])
    assert len(months) == 2
    sept, october = months
    assert sept.items[0][1] == [(28, 1), (29, 2), (30, 3)]
    assert october.items[0][1] == [(1, 4), (2, 5), (3, 6)]


def test_plan_cross_year_month_enumeration(svc):
    """12 月 -> 次年 1 月的跨年枚举。"""
    months = svc.plan([_item("A", date(2026, 12, 30), date(2027, 1, 2))])
    assert [(m.year, m.month) for m in months] == [(2026, 12), (2027, 1)]


def test_plan_only_active_months_and_sorting(svc):
    """项点只出现在其活跃月份块;块内按(开始, 名称)排序。"""
    months = svc.plan(
        [
            _item("B9", date(2026, 9, 21), date(2026, 9, 27)),
            _item("H5", date(2026, 9, 17), date(2026, 9, 21)),
            _item("Z12", date(2026, 10, 1), date(2026, 10, 5)),
        ]
    )
    assert len(months) == 2
    assert [it.name for it, _c in months[0].items] == ["H5", "B9"]
    assert [it.name for it, _c in months[1].items] == ["Z12"]


def test_plan_parallel_counts(svc):
    """并行数:重叠日=2,独占日=1,无项点日不入字典。"""
    months = svc.plan(
        [
            _item("A", date(2026, 9, 17), date(2026, 9, 21)),
            _item("B", date(2026, 9, 21), date(2026, 9, 27)),
        ]
    )
    sept = months[0]
    assert sept.parallel == dict.fromkeys(range(17, 21), 1) | {21: 2} | dict.fromkeys(
        range(22, 28), 1
    )


def test_peak_parallel(svc):
    months = svc.plan(
        [
            _item("A", date(2026, 9, 17), date(2026, 9, 21)),
            _item("B", date(2026, 9, 21), date(2026, 9, 27)),
            _item("C", date(2026, 9, 21), date(2026, 9, 21)),
        ]
    )
    assert PlanScheduleService.peak_parallel(months) == (3, date(2026, 9, 21))
    assert PlanScheduleService.peak_parallel([]) is None


# ==================== 模板 ====================


def test_write_template_round_trips(svc, tmp_path):
    """模板可写盘且能被 parse 解析回 2 个示例项点。"""
    out = svc.write_template(tmp_path / "tpl.xlsx")
    assert out.is_file()
    items, invalid = svc.parse(out)
    assert invalid == []
    assert [it.name for it in items] == ["示例项点A", "示例项点B"]
    assert items[0].start == date(2026, 9, 17)


def test_write_template_never_overwrites(svc, tmp_path):
    first = svc.write_template(tmp_path / "tpl.xlsx")
    precious = tmp_path / "tpl.xlsx"
    precious.write_text("precious")
    second = svc.write_template(precious)
    assert precious.read_text(encoding="utf-8") == "precious"
    assert second != first and second.name == "tpl_1.xlsx"


def test_write_template_normalizes_suffix(svc, tmp_path):
    out = svc.write_template(tmp_path / "tpl.txt")
    assert out.suffix == ".xlsx"


# ==================== 生成:布局渲染 ====================


def _make_input(make_xlsx, rows: list[list[object]]) -> Path:
    return make_xlsx("list.xlsx", {"S": rows})


def test_generate_layout_matches_delivery_template(svc, make_xlsx, tmp_path):
    """生成布局对齐交付计划示例:MONTH/DATE/项点行/并行数 + 第几天编号。"""
    src = _make_input(
        make_xlsx,
        [
            ["项点名称", "起始日期", "终止日期"],
            ["合5第8列", date(2026, 9, 17), date(2026, 9, 21)],
            ["杭9第2列", date(2026, 9, 21), date(2026, 9, 27)],
        ],
    )
    result = svc.generate(src, tmp_path / "out.xlsx")
    assert result.success
    assert result.output is not None and result.output.is_file()

    wb = load_workbook(result.output)
    assert wb.sheetnames == [SHEET_NAME]
    ws = wb[SHEET_NAME]

    # MONTH 行:A1 标签 + 合并单元格 B1:AE1(9 月 30 天)+ 年月文本
    assert ws["A1"].value == "MONTH"
    assert "B1:AE1" in {str(r) for r in ws.merged_cells.ranges}
    assert ws["B1"].value == "2026年9月"

    # DATE 行:1..30
    assert ws["A2"].value == "DATE"
    assert [ws.cell(row=2, column=c).value for c in range(2, 32)] == list(range(1, 31))

    # 项点行:第几天编号落在正确列(9/17 -> R3=1 ... 9/21 -> V3=5)
    assert ws["A3"].value == "合5第8列"
    assert [ws.cell(row=3, column=c).value for c in range(18, 23)] == [1, 2, 3, 4, 5]
    assert ws["A4"].value == "杭9第2列"
    assert ws["V4"].value == 1  # 9/21
    assert ws["AB4"].value == 7  # 9/27

    # 并行数行:重叠日 21 = 2,独占日 17/26 = 1,无项点日为空
    assert ws["A5"].value == "并行数"
    assert ws["R5"].value == 1
    assert ws["V5"].value == 2
    assert ws["AA5"].value == 1
    assert ws["B5"].value is None

    # 活动格实心填色(项点色),两个项点颜色不同
    fill_a = ws["R3"].fill
    fill_b = ws["V4"].fill
    assert fill_a.patternType == "solid"
    assert fill_a.fgColor.rgb == ITEM_FILLS[0]
    assert fill_b.fgColor.rgb == ITEM_FILLS[1]
    assert fill_a.fgColor.rgb != fill_b.fgColor.rgb

    # 列宽:A 列加宽、日期列窄
    assert ws.column_dimensions["A"].width == 14.0
    assert ws.column_dimensions["B"].width == 4.5


def test_generate_marks_weekends(svc, make_xlsx, tmp_path):
    """周末整列灰底优先:DATE 灰底红字;项点/并行数行周末列灰底(含活动值格),
    非周末活动格保留项点填色,工作日无值格无底色。"""
    src = _make_input(
        make_xlsx,
        [["项点名称", "起始日期", "终止日期"], ["A", date(2026, 9, 17), date(2026, 9, 21)]],
    )
    result = svc.generate(src, tmp_path / "out.xlsx")
    assert result.success
    assert result.output is not None
    ws = load_workbook(result.output)[SHEET_NAME]

    saturday = next(d for d in range(17, 22) if date(2026, 9, d).weekday() == 5)  # 19,区间内
    sunday = saturday + 1  # 20
    monday = sunday + 1  # 21,区间最后一天
    # DATE 行:周末表头灰底红字
    sat_cell = ws.cell(row=2, column=1 + saturday)
    assert sat_cell.fill.fgColor.rgb == WEEKEND_FILL
    assert sat_cell.font.color.rgb == "FFC00000"
    # 项点行(第 3 行):周末活动格灰底优先、值保留;非周末活动格保留项点填色
    item_sat = ws.cell(row=3, column=1 + saturday)
    assert item_sat.value == 3  # 9/19 = 项点内第 3 天
    assert item_sat.fill.fgColor.rgb == WEEKEND_FILL
    assert ws.cell(row=3, column=1 + sunday).fill.fgColor.rgb == WEEKEND_FILL
    item_mon = ws.cell(row=3, column=1 + monday)
    assert item_mon.value == 5
    assert item_mon.fill.fgColor.rgb == ITEM_FILLS[0]
    # 项点行:区间外的空周末格仍灰底、无值
    item_empty_sat = ws.cell(row=3, column=1 + saturday + 7)  # 9/26
    assert item_empty_sat.value is None
    assert item_empty_sat.fill.fgColor.rgb == WEEKEND_FILL
    # 并行数行(第 4 行):非零并行数的周末格灰底优先、值保留;周一有值无填色
    parallel_sat = ws.cell(row=4, column=1 + saturday)
    assert parallel_sat.value == 1
    assert parallel_sat.fill.fgColor.rgb == WEEKEND_FILL
    parallel_mon = ws.cell(row=4, column=1 + monday)
    assert parallel_mon.value == 1
    assert parallel_mon.fill.patternType is None


def test_generate_name_mode_writes_item_name(svc, make_xlsx, tmp_path):
    """cell_mode=name:活动日期格写项点名称(第x列/批次)而非天数序号。

    日期列按最长名称自适应加宽(合5第8列 5 字 -> 2*5+1=11),并行数行不变。
    """
    src = _make_input(
        make_xlsx,
        [
            ["项点名称", "起始日期", "终止日期"],
            ["合5第8列", date(2026, 9, 17), date(2026, 9, 21)],
            ["杭9第2列", date(2026, 9, 21), date(2026, 9, 27)],
        ],
    )
    result = svc.generate(src, tmp_path / "out.xlsx", ScheduleOptions(cell_mode=CELL_NAME))
    assert result.success
    assert result.output is not None
    ws = load_workbook(result.output)[SHEET_NAME]

    # 每个活动日格子都写项点名称;非活动日仍为空
    assert [ws.cell(row=3, column=c).value for c in range(18, 23)] == ["合5第8列"] * 5
    assert [ws.cell(row=4, column=c).value for c in range(22, 29)] == ["杭9第2列"] * 7
    assert ws.cell(row=3, column=17).value is None
    # 并行数行不受模式影响
    assert ws["V5"].value == 2
    # 活动格填色保持(区分并行项点);周末活动格灰底优先、名称值保留
    assert ws["R3"].fill.fgColor.rgb == ITEM_FILLS[0]
    assert ws["T3"].value == "合5第8列"  # 9/19 周六
    assert ws["T3"].fill.fgColor.rgb == WEEKEND_FILL
    # 日期列宽自适应:最长名称 5 字 -> 11.0
    assert ws.column_dimensions["B"].width == 11.0


def test_generate_name_mode_width_capped(svc, make_xlsx, tmp_path):
    """name 模式列宽封顶:超长名称不超过 NAME_MODE_DAY_WIDTH_MAX。"""
    long_name = "很长很长的项点名称示例"
    src = _make_input(
        make_xlsx,
        [["项点名称", "起始日期", "终止日期"], [long_name, date(2026, 9, 17), date(2026, 9, 21)]],
    )
    result = svc.generate(src, tmp_path / "o.xlsx", ScheduleOptions(cell_mode=CELL_NAME))
    assert result.success
    assert result.output is not None
    ws = load_workbook(result.output)[SHEET_NAME]
    assert ws.cell(row=3, column=18).value == long_name
    assert ws.column_dimensions["B"].width == NAME_MODE_DAY_WIDTH_MAX
    # index 模式(默认)不受长名称影响,保持窄列
    result2 = svc.generate(src, tmp_path / "o2.xlsx")
    assert result2.success
    assert result2.output is not None
    ws2 = load_workbook(result2.output)[SHEET_NAME]
    assert ws2.cell(row=3, column=18).value == 1
    assert ws2.column_dimensions["B"].width == DAY_COLUMN_WIDTH


def test_generate_two_month_blocks_and_continuity(svc, make_xlsx, tmp_path):
    """跨月两个块:10 月块的 MONTH/DATE 行位置正确且编号接续。"""
    src = _make_input(
        make_xlsx,
        [["项点名称", "起始日期", "终止日期"], ["A", date(2026, 9, 28), date(2026, 10, 3)]],
    )
    result = svc.generate(src, tmp_path / "out.xlsx")
    assert result.success
    assert result.output is not None
    ws = load_workbook(result.output)[SHEET_NAME]

    # 9 月块:MONTH r1 / DATE r2 / A r3 / 并行数 r4
    assert ws["A4"].value == "并行数"
    assert ws.cell(row=3, column=1 + 28).value == 1
    # 10 月块从 r5 开始
    assert ws["A5"].value == "MONTH"
    assert ws["B5"].value == "2026年10月"
    assert ws["A6"].value == "DATE"
    assert ws["A7"].value == "A"
    assert ws.cell(row=7, column=2).value == 4  # 10/1 = 项点内第 4 天
    assert ws["A8"].value == "并行数"


def test_generate_progress_callback_per_month(svc, make_xlsx, tmp_path):
    src = _make_input(
        make_xlsx,
        [["项点名称", "起始日期", "终止日期"], ["A", date(2026, 9, 28), date(2026, 10, 3)]],
    )
    seen: list[tuple[int, int, str]] = []
    result = svc.generate(src, tmp_path / "o.xlsx", progress_callback=lambda *a: seen.append(a))
    assert result.success
    assert seen == [(1, 2, "生成 2026年9月"), (2, 2, "生成 2026年10月")]


# ==================== 生成:输出与错误路径 ====================


def test_generate_output_auto_numbered(svc, make_xlsx, tmp_path):
    """输出已存在时自动加序号,原文件不被覆盖。"""
    src = _make_input(
        make_xlsx,
        [["项点名称", "起始日期", "终止日期"], ["A", date(2026, 9, 17), date(2026, 9, 21)]],
    )
    out = tmp_path / "out.xlsx"
    out.write_text("precious")

    result = svc.generate(src, out)

    assert result.success
    assert out.read_text(encoding="utf-8") == "precious"
    assert result.output == tmp_path / "out_1.xlsx"
    assert result.output.is_file()


def test_generate_normalizes_suffix_and_creates_dirs(svc, make_xlsx, tmp_path):
    """非 .xlsx 后缀归一;输出父目录自动创建。"""
    src = _make_input(
        make_xlsx,
        [["项点名称", "起始日期", "终止日期"], ["A", date(2026, 9, 17), date(2026, 9, 21)]],
    )
    out = tmp_path / "sub" / "deep" / "out.xls"
    result = svc.generate(src, out)
    assert result.success
    assert result.output == tmp_path / "sub" / "deep" / "out.xlsx"
    assert result.output.is_file()


def test_generate_all_invalid_returns_error_with_rows(svc, make_xlsx, tmp_path):
    """全部行无效:不写输出,保留 invalid 供展示。"""
    src = _make_input(
        make_xlsx,
        [["项点名称", "起始日期", "终止日期"], ["倒挂", date(2026, 9, 2), date(2026, 9, 1)]],
    )
    result = svc.generate(src, tmp_path / "o.xlsx")
    assert not result.success
    assert result.output is None
    assert "没有可排布的项点" in result.error_message
    assert [inv.row for inv in result.invalid] == [2]


def test_generate_empty_list_returns_error(svc, make_xlsx, tmp_path):
    """清单只有表头:错误信息标明清单为空。"""
    src = _make_input(make_xlsx, [["项点名称", "起始日期", "终止日期"]])
    result = svc.generate(src, tmp_path / "o.xlsx")
    assert not result.success
    assert "清单为空" in result.error_message


def test_generate_unreadable_input_returns_error(svc, tmp_path):
    """文件级失败(损坏)进 error_message 而非异常。"""
    bad = tmp_path / "bad.xlsx"
    bad.write_bytes(b"nope")
    result = svc.generate(bad, tmp_path / "o.xlsx")
    assert not result.success
    assert "无法读取" in result.error_message


def test_generate_unsupported_suffix_returns_error(svc, tmp_path):
    f = tmp_path / "list.csv"
    f.write_text("x")
    result = svc.generate(f, tmp_path / "o.xlsx")
    assert not result.success
    assert "不支持的格式" in result.error_message


def test_generate_records_history(make_xlsx, tmp_path):
    """注入 history_store:成功后记录一条 plan_schedule 历史(含各计数)。"""
    store = MagicMock()
    src = _make_input(
        make_xlsx,
        [["项点名称", "起始日期", "终止日期"], ["A", date(2026, 9, 17), date(2026, 9, 21)]],
    )
    svc = PlanScheduleService(history_store=store)
    result = svc.generate(src, tmp_path / "o.xlsx")

    assert result.success
    store.add_record.assert_called_once()
    tool, payload = store.add_record.call_args[0]
    assert tool == "plan_schedule"
    assert payload["item_count"] == 1
    assert payload["month_count"] == 1
    assert payload["invalid_count"] == 0
    assert payload["success"] is True


def test_generate_failure_does_not_record_history(make_xlsx, tmp_path):
    """失败(全部行无效)不写历史。"""
    store = MagicMock()
    src = _make_input(
        make_xlsx,
        [["项点名称", "起始日期", "终止日期"], ["倒挂", date(2026, 9, 2), date(2026, 9, 1)]],
    )
    result = PlanScheduleService(history_store=store).generate(src, tmp_path / "o.xlsx")
    assert not result.success
    store.add_record.assert_not_called()


def test_result_success_semantics():
    """ScheduleResult.success 仅看 output 是否写出。"""
    assert ScheduleResult(output=Path("o.xlsx")).success is True
    assert ScheduleResult(error_message="x").success is False


def test_month_plan_defaults():
    """MonthPlan 的 items/parallel 有默认空容器。"""
    mp = MonthPlan(year=2026, month=9, days=30, weekends=frozenset({5}))
    assert mp.items == []
    assert mp.parallel == {}
    assert InvalidRow(2, "err").error == "err"
