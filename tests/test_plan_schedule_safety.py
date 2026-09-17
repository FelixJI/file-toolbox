"""PR #95: 输出真实性、排他写入与线程收尾回归。"""

import time
from datetime import date
from threading import Event

import pytest
from openpyxl import Workbook, load_workbook

from file_toolbox.common.history import JsonHistoryStore
from file_toolbox.core.plan_schedule import PlanScheduleService, ScheduleOptions, ScheduleResult


def input_book(tmp_path, name="A"):
    path = tmp_path / "input.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["项点名称", "起始日期", "终止日期"])
    ws.append([name, date(2026, 9, 17), date(2026, 9, 19)])
    ws["A2"].data_type = "s"
    wb.save(path)
    return path


def test_history_failure_keeps_written_result(tmp_path):
    blocked = tmp_path / "history"
    blocked.write_text("occupied")
    svc = PlanScheduleService(history_store=JsonHistoryStore(blocked))
    result = svc.generate(input_book(tmp_path), tmp_path / "out.xlsx")
    assert result.success and result.output.is_file()
    assert "历史" in result.warning_message
    wb = load_workbook(result.output)
    assert wb.active["A3"].value == "A"
    wb.close()


@pytest.mark.parametrize("mode", ["index", "name"])
@pytest.mark.parametrize("name", ["=1+1", "  项点 A  "])
def test_literal_name_round_trip(tmp_path, mode, name):
    result = PlanScheduleService().generate(
        input_book(tmp_path, name), tmp_path / "out.xlsx", ScheduleOptions(cell_mode=mode)
    )
    wb = load_workbook(result.output)
    try:
        assert wb.active["A3"].value == name
        assert wb.active["A3"].data_type == "s"
        if mode == "name":
            assert wb.active["R3"].value == name
            assert wb.active["R3"].data_type == "s"
    finally:
        wb.close()


@pytest.mark.parametrize("template", [False, True])
def test_racing_writer_is_not_overwritten(tmp_path, monkeypatch, template):
    svc = PlanScheduleService()
    output = tmp_path / "out.xlsx"
    original = svc._resolve_output_path
    raced = False

    def resolve(path):
        nonlocal raced
        candidate = original(path)
        if not raced:
            raced = True
            candidate.write_bytes(b"other writer")
        return candidate

    monkeypatch.setattr(svc, "_resolve_output_path", resolve)
    actual = (
        svc.write_template(output)
        if template
        else svc.generate(input_book(tmp_path), output).output
    )
    assert output.read_bytes() == b"other writer"
    assert actual == tmp_path / "out_1.xlsx"
    wb = load_workbook(actual)
    wb.close()


def test_close_waits_for_real_worker(tmp_path, monkeypatch):
    pytest.importorskip("PySide6.QtWidgets")
    from PySide6.QtGui import QCloseEvent
    from PySide6.QtWidgets import QApplication, QMessageBox

    from file_toolbox.gui.dialogs.plan_schedule_tab import PlanScheduleTab

    app = QApplication.instance() or QApplication([])
    tab = PlanScheduleTab()
    entered, release = Event(), Event()

    class SlowService:
        def generate(self, *args, **kwargs):
            entered.set()
            assert release.wait(10)
            return ScheduleResult(error_message="synthetic")

    tab._svc = SlowService()
    tab.ui.edit_input.setText(str(input_book(tmp_path)))
    tab.ui.edit_outdir.setText(str(tmp_path))
    monkeypatch.setattr(QMessageBox, "warning", lambda *a: None)
    monkeypatch.setattr(QMessageBox, "critical", lambda *a: None)
    tab._generate()
    worker = tab._worker
    try:
        assert entered.wait(2)
        event = QCloseEvent()
        tab.closeEvent(event)
        assert not event.isAccepted()
        assert tab._worker is worker and worker.isRunning()
        assert tab.close_pending
    finally:
        release.set()
        assert worker.wait(2000)
        deadline = time.monotonic() + 2
        while tab._worker is not None and time.monotonic() < deadline:
            app.processEvents()
        app.processEvents()
    assert tab._worker is None
    assert not tab.close_pending


def test_partial_workbook_failure_removes_only_new_output(tmp_path, monkeypatch):
    src = input_book(tmp_path)
    output = tmp_path / "out.xlsx"
    output.write_bytes(b"existing")

    def fail_save(self, stream):
        stream.write(b"partial")
        raise OSError("disk full")

    monkeypatch.setattr(Workbook, "save", fail_save)
    result = PlanScheduleService().generate(src, output)
    assert not result.success and "disk full" in result.error_message
    assert output.read_bytes() == b"existing"
    assert not (tmp_path / "out_1.xlsx").exists()


def test_cli_and_gui_report_history_warning(tmp_path, monkeypatch):
    from typer.testing import CliRunner

    from file_toolbox.cli import plan_schedule_cmd
    from file_toolbox.cli.main import app
    from file_toolbox.gui.controllers.plan_schedule_controller import PlanScheduleController

    blocked = tmp_path / "history"
    blocked.write_text("occupied")
    monkeypatch.setattr(plan_schedule_cmd, "JsonHistoryStore", lambda: JsonHistoryStore(blocked))
    output = tmp_path / "out.xlsx"
    result = CliRunner().invoke(
        app, ["plan-schedule", str(input_book(tmp_path)), "--yes", "-o", str(output)]
    )
    assert result.exit_code == 0
    assert "历史未保存" in result.output
    assert "完成:" in result.output and output.exists()
    summary = PlanScheduleController.summarize(
        ScheduleResult(output=output, warning_message="历史未保存")
    )
    assert "已排布" in summary and "历史未保存" in summary


def test_whitespace_only_name_is_invalid(tmp_path):
    items, invalid = PlanScheduleService().parse(input_book(tmp_path, "  "))
    assert items == []
    assert len(invalid) == 1 and invalid[0].error == "缺少项点名称"
