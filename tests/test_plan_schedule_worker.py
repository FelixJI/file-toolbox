"""PlanScheduleWorker 测试:正常完成、异常、进度、真实生成集成。

worker.run() 同步调用(不走 QThread.start)验证逻辑;另含真实 start() 的
跨线程信号投递集成(与 test_excel_merge_worker 同范式)。基于虚构 xlsx,不触发 COM。
"""

import pytest

# 用 QtWidgets 子模块做 importorskip(而非顶层 PySide6):后者只校验包可 import,
# 不触发 libEGL/libGL 原生库加载;真实 import QtWidgets 才会,缺库时应跳过而非收集失败。
pytest.importorskip("PySide6.QtWidgets")

from datetime import date  # noqa: E402
from pathlib import Path  # noqa: E402

from openpyxl import load_workbook  # noqa: E402
from PySide6.QtWidgets import QApplication  # noqa: E402

from file_toolbox.core.plan_schedule import (  # noqa: E402
    SHEET_NAME,
    PlanScheduleService,
    ScheduleOptions,
    ScheduleResult,
)
from file_toolbox.gui.workers.plan_schedule_worker import PlanScheduleWorker  # noqa: E402


@pytest.fixture(scope="module")
def app():
    return QApplication.instance() or QApplication([])


class _FakeService:
    """假 PlanScheduleService:记录调用,可控返回/异常。"""

    def __init__(self, result=None, error=None):
        self._result = result
        self._error = error
        self.generate_calls: list[dict] = []

    def generate(self, input_path, output, options=None, progress_callback=None):
        self.generate_calls.append({"input": input_path, "output": output, "options": options})
        if self._error:
            raise self._error
        if progress_callback is not None:
            for i in range(2):
                progress_callback(i + 1, 2, f"生成 2026年{9 + i}月")
        return self._result


def test_worker_emits_finished_ok_on_success(app):
    """正常完成 → finished_ok 信号带 ScheduleResult,透传 input/output/options。"""
    expected = ScheduleResult(output=Path("o.xlsx"))
    svc = _FakeService(result=expected)
    options = ScheduleOptions(default_year=2026)
    worker = PlanScheduleWorker(svc, Path("in.xlsx"), Path("o.xlsx"), options)

    captured = {}
    worker.finished_ok.connect(lambda r: captured.setdefault("ok", r))
    worker.failed.connect(lambda m: captured.setdefault("fail", m))

    worker.run()  # 同步跑(不经 QThread.start)

    assert captured.get("ok") is expected
    assert "fail" not in captured
    assert len(svc.generate_calls) == 1
    assert svc.generate_calls[0]["options"] is options


def test_worker_emits_failed_on_exception(app):
    """service 抛异常 → failed 信号,不发 finished_ok。"""
    svc = _FakeService(error=RuntimeError("boom"))
    worker = PlanScheduleWorker(svc, Path("in.xlsx"), Path("o.xlsx"), ScheduleOptions())

    captured = {}
    worker.finished_ok.connect(lambda r: captured.setdefault("ok", r))
    worker.failed.connect(lambda m: captured.setdefault("fail", m))

    worker.run()

    assert "boom" in captured.get("fail", "")
    assert "ok" not in captured


def test_worker_emits_progress(app):
    """进度回调 → progress 信号,(current, total, msg) 形状正确。"""
    svc = _FakeService(result=ScheduleResult())
    worker = PlanScheduleWorker(svc, Path("in.xlsx"), Path("o.xlsx"), ScheduleOptions())

    progress_msgs = []
    worker.progress.connect(lambda c, t, m: progress_msgs.append((c, t, m)))

    worker.run()

    assert progress_msgs == [(1, 2, "生成 2026年9月"), (2, 2, "生成 2026年10月")]


def test_worker_start_generates_real_output_across_threads(app, make_xlsx, tmp_path):
    """集成:真实 service + worker.start() 后台线程 → 结果投递回主线程。"""
    src = make_xlsx(
        "清单.xlsx",
        {"S": [["项点名称", "起始日期", "终止日期"], ["A", date(2026, 9, 17), date(2026, 9, 21)]]},
    )
    out = tmp_path / "计划排布.xlsx"
    svc = PlanScheduleService()
    worker = PlanScheduleWorker(svc, src, out, ScheduleOptions())

    results: list[ScheduleResult] = []
    worker.finished_ok.connect(results.append)
    failed: list[str] = []
    worker.failed.connect(failed.append)

    worker.start()
    assert worker.wait(10000), "worker 未在 10s 内结束"
    app.processEvents()
    if not results:
        app.processEvents()

    assert failed == []
    assert len(results) == 1 and results[0].success
    assert load_workbook(out).sheetnames == [SHEET_NAME]
