"""Excel 源级失败与资源所有权:真实临时 xlsx,只在故障边界注入异常。"""

from unittest.mock import Mock

import pytest
from openpyxl import Workbook, load_workbook

from file_toolbox.common.operation_errors import OperationResultError
from file_toolbox.core.excel_merge import NAMING_KEEP, ExcelMergeService, MergeOptions


@pytest.mark.parametrize(
    "operation,fault",
    [
        ("plan", "enumerate"),
        ("plan", "close"),
        ("merge", "enumerate"),
        ("merge", "copy"),
        ("merge", "close"),
        ("merge", "copy-close"),
    ],
)
def test_failed_source_is_rolled_back_and_next_source_survives(
    operation, fault, make_xlsx, tmp_path, monkeypatch, caplog
):
    bad = make_xlsx("bad.xlsx", {"First": [["bad-first"]], "Second": [["bad-second"]]})
    good = make_xlsx("good.xlsx", {"First": [["good"]]})
    svc = ExcelMergeService()
    sources, closed = [], []
    real_load = svc._load_workbook
    real_copy = svc._copy_worksheet
    worksheets = Workbook.worksheets.fget

    def load(path, mode):
        src = real_load(path, mode)
        sources.append(src)
        close = src.close

        def finish():
            closed.append(path.name)
            close()
            if path == bad and "close" in fault:
                raise OSError("source close fault")

        monkeypatch.setattr(src, "close", finish)
        return src

    def enumerate_sheets(src):
        rows = worksheets(src)
        if sources and src is sources[0] and fault == "enumerate":

            def broken():
                yield rows[0]
                raise ValueError("source enumeration fault")

            return broken()
        return rows

    def copy_sheet(src, dest, title):
        result = real_copy(src, dest, title)
        if src.parent is sources[0] and src.title == "Second" and "copy" in fault:
            raise ValueError("source copy fault")
        return result

    monkeypatch.setattr(svc, "_load_workbook", load)
    monkeypatch.setattr(Workbook, "worksheets", property(enumerate_sheets))
    monkeypatch.setattr(svc, "_copy_worksheet", copy_sheet)
    options = MergeOptions(naming=NAMING_KEEP)
    output = tmp_path / "out.xlsx"
    if operation == "plan":
        sheets, failed = svc.plan_sheets([bad, good], options)
        assert not output.exists()
    else:
        result = svc.merge([bad, good], output, options)
        assert result.success
        sheets, failed = result.sheets, result.failed
        with_output = load_workbook(output)
        try:
            assert with_output.sheetnames == ["First"]
            assert with_output["First"]["A1"].value == "good"
        finally:
            with_output.close()
    assert [(s.file, s.target_name) for s in sheets] == [("good.xlsx", "First")]
    assert len(failed) == 1 and failed[0].file == "bad.xlsx"
    assert "fault" in failed[0].error
    assert closed == ["bad.xlsx", "good.xlsx"]
    if fault == "copy-close":
        assert "source copy fault" in failed[0].error
        assert "source close fault" in caplog.text


@pytest.mark.parametrize(
    "exit_path", ["success", "cancel", "all-failed", "save-failed", "callback", "memory"]
)
def test_destination_closes_on_each_exit(exit_path, make_xlsx, tmp_path, monkeypatch):
    source = make_xlsx("source.xlsx", {"Data": [[1]]})
    svc = ExcelMergeService()
    dest = svc._new_workbook()
    closed = Mock(wraps=dest.close)
    monkeypatch.setattr(dest, "close", closed)
    monkeypatch.setattr(svc, "_new_workbook", lambda: dest)
    output = tmp_path / "out.xlsx"

    def fail(*args):
        raise OSError("injected failure")

    if exit_path == "save-failed":
        monkeypatch.setattr(dest, "save", fail)
    if exit_path == "all-failed":
        source.write_bytes(b"invalid xlsx")

    def copy(*args):
        raise MemoryError("memory fault")

    if exit_path == "memory":
        monkeypatch.setattr(svc, "_copy_worksheet", copy)
    kwargs = {"cancel_check": lambda: exit_path == "cancel"}
    if exit_path == "callback":
        kwargs["progress_callback"] = fail
    if exit_path in {"callback", "memory"}:
        with pytest.raises((OSError, MemoryError), match="fault|failure"):
            svc.merge([source], output, **kwargs)
    else:
        result = svc.merge([source], output, **kwargs)
        assert result.success == (exit_path == "success")
    closed.assert_called_once()
    assert output.exists() == (exit_path == "success")


def test_destination_close_failure_preserves_written_result(make_xlsx, tmp_path, monkeypatch):
    source = make_xlsx("source.xlsx", {"Data": [["kept"]]})
    svc = ExcelMergeService()
    dest = svc._new_workbook()

    def close():
        raise OSError("destination close fault")

    monkeypatch.setattr(dest, "close", close)
    monkeypatch.setattr(svc, "_new_workbook", lambda: dest)
    with pytest.raises(OperationResultError, match="destination close fault") as caught:
        svc.merge([source], tmp_path / "out.xlsx")
    result = caught.value.result
    assert result.success and len(result.sheets) == 1
    wb = load_workbook(result.output)
    try:
        assert wb["source-Data"]["A1"].value == "kept"
    finally:
        wb.close()


def test_cleanup_does_not_replace_primary_error(make_xlsx, tmp_path, monkeypatch, caplog):
    source = make_xlsx("source.xlsx", {"Data": [[1]]})
    svc = ExcelMergeService()
    dest = svc._new_workbook()

    def close():
        raise OSError("destination close fault")

    def callback(*args):
        raise ValueError("primary callback fault")

    monkeypatch.setattr(dest, "close", close)
    monkeypatch.setattr(svc, "_new_workbook", lambda: dest)
    with pytest.raises(ValueError, match="primary callback fault"):
        svc.merge([source], tmp_path / "out.xlsx", progress_callback=callback)
    assert "destination close fault" in caplog.text
    assert not (tmp_path / "out.xlsx").exists()


def test_rollback_failure_is_fatal_and_preserves_copy_error(
    make_xlsx, tmp_path, monkeypatch, caplog
):
    source = make_xlsx("source.xlsx", {"Data": [[1]]})
    svc = ExcelMergeService()
    dest = svc._new_workbook()
    close = Mock(wraps=dest.close)
    monkeypatch.setattr(dest, "close", close)
    monkeypatch.setattr(svc, "_new_workbook", lambda: dest)

    def broken_copy(src, dest, title):
        dest.create_sheet(title)
        raise ValueError("primary copy fault")

    def broken_remove(sheet):
        raise OSError("rollback fault")

    monkeypatch.setattr(svc, "_copy_worksheet", broken_copy)
    monkeypatch.setattr(dest, "remove", broken_remove)
    with pytest.raises(ValueError, match="primary copy fault") as caught:
        svc.merge([source], tmp_path / "out.xlsx")
    assert isinstance(caught.value.__cause__, OSError)
    assert "rollback fault" in caplog.text
    close.assert_called_once()
    assert not (tmp_path / "out.xlsx").exists()


def test_copy_memory_error_closes_source_and_aborts(make_xlsx, tmp_path, monkeypatch):
    source = make_xlsx("source.xlsx", {"Data": [[1]]})
    svc = ExcelMergeService()
    src = svc._load_workbook(source, "values")
    close = Mock(wraps=src.close)
    monkeypatch.setattr(src, "close", close)
    monkeypatch.setattr(svc, "_load_workbook", lambda *args: src)

    def fail(*args):
        raise MemoryError("memory exhausted")

    monkeypatch.setattr(svc, "_copy_worksheet", fail)
    with pytest.raises(MemoryError, match="memory exhausted"):
        svc.merge([source], tmp_path / "out.xlsx")
    close.assert_called_once()
    assert not (tmp_path / "out.xlsx").exists()


def test_real_xlsx_has_no_open_file_after_source_copy_failure(make_xlsx, tmp_path, monkeypatch):
    import psutil

    bad = make_xlsx("bad.xlsx", {"Data": [[1]]})
    good = make_xlsx("good.xlsx", {"Data": [[2]]})
    svc = ExcelMergeService()
    original = svc._copy_worksheet

    def copy(src, dest, title):
        original(src, dest, title)
        if title.startswith("bad-"):
            raise ValueError("injected copy failure")

    monkeypatch.setattr(svc, "_copy_worksheet", copy)
    result = svc.merge([bad, good], tmp_path / "out.xlsx")
    assert result.success
    owned = {str(path.resolve()).casefold() for path in (bad, good, result.output)}
    assert not owned.intersection(f.path.casefold() for f in psutil.Process().open_files())
    for path in (bad, good, result.output):
        moved = path.with_suffix(".moved")
        path.rename(moved)
        moved.unlink()
