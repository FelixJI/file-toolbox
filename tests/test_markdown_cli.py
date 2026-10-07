"""markdown-convert CLI 测试:默认预览不落盘、--yes 执行、参数校验、错误退出码。

历史隔离:monkeypatch.chdir(tmp_path),默认 JsonHistoryStore() 落入临时目录。
docx 用例走真实内置 Pandoc。
"""

import json
from pathlib import Path

from typer.testing import CliRunner

from file_toolbox.cli.main import app

runner = CliRunner()

MD_TABLE = "# 标题\n\n| a | b |\n|---|---|\n| 1 | 2 |\n"


def _md(tmp_path: Path, name: str, content: str = MD_TABLE) -> Path:
    path = tmp_path / name
    path.write_text(content, encoding="utf-8")
    return path


def test_preview_shows_path_plan_without_writing(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    a = _md(tmp_path, "a.md")
    b = _md(tmp_path, "b.markdown")

    r = runner.invoke(app, ["markdown-convert", str(a), str(b)])

    assert r.exit_code == 0
    assert f"{a} -> {tmp_path / 'a.docx'}" in r.output
    assert f"{b} -> {tmp_path / 'b.docx'}" in r.output
    assert "共 2 个文件 -> docx" in r.output
    assert "预览模式" in r.output
    assert not any(tmp_path.glob("*.docx"))
    assert not (tmp_path / ".file_toolbox" / "history").exists()


def test_preview_xlsx_mode_notes_and_output_dir(tmp_path):
    a = _md(tmp_path, "a.md")
    out = tmp_path / "out"

    r = runner.invoke(app, ["markdown-convert", str(a), "--to", "xlsx", "--output-dir", str(out)])

    assert r.exit_code == 0
    assert f"{a} -> {out / 'a.xlsx'}" in r.output
    assert "(--excel-mode tables)" in r.output
    assert "没有表格的文件会在执行时跳过" in r.output
    assert not out.exists()


def test_execute_docx_writes_output_and_history(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    a = _md(tmp_path, "a.md", "# 标题\n\n正文 [l](https://example.com/x)。\n")

    r = runner.invoke(app, ["markdown-convert", str(a), "--yes"])

    assert r.exit_code == 0
    out = tmp_path / "a.docx"
    assert out.is_file() and out.read_bytes()[:2] == b"PK"
    assert "完成: 1 个输出" in r.output
    history = tmp_path / ".file_toolbox" / "history" / "markdown_convert.jsonl"
    assert history.is_file()
    record = json.loads(history.read_text(encoding="utf-8").splitlines()[0])
    assert record["data"]["success"] == 1
    assert record["data"]["target"] == "docx"
    assert "excel_mode" not in record["data"]


def test_execute_xlsx_tables_with_output_dir(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    from openpyxl import load_workbook

    a = _md(tmp_path, "a.md")
    out = tmp_path / "out"

    r = runner.invoke(
        app,
        [
            "markdown-convert",
            str(a),
            "--to",
            "xlsx",
            "--excel-mode",
            "tables",
            "--output-dir",
            str(out),
            "--yes",
        ],
    )

    assert r.exit_code == 0
    wb = load_workbook(out / "a.xlsx")
    assert wb.sheetnames == ["表1"]
    assert wb["表1"]["A1"].value == "a"
    wb.close()


def test_execute_xlsx_document_pure_text_still_outputs(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    from openpyxl import load_workbook

    a = _md(tmp_path, "a.md", "# 纯文本\n\n没有表格。\n")

    r = runner.invoke(
        app, ["markdown-convert", str(a), "--to", "xlsx", "--excel-mode", "document", "--yes"]
    )

    assert r.exit_code == 0
    wb = load_workbook(tmp_path / "a.xlsx")
    assert wb.sheetnames == ["正文"]
    wb.close()


def test_execute_tables_mode_reports_skip_with_nonzero(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    a = _md(tmp_path, "a.md", "# 无表格\n")

    r = runner.invoke(app, ["markdown-convert", str(a), "--to", "xlsx", "--yes"])

    assert r.exit_code == 1
    assert "跳过: a.md - 没有表格" in r.output
    assert "没有写出任何输出" in r.output


def test_execute_partial_failure_exits_nonzero(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    a = _md(tmp_path, "a.md")
    bad = tmp_path / "bad.md"
    bad.write_bytes("中文".encode("gbk"))

    r = runner.invoke(app, ["markdown-convert", str(a), str(bad), "--to", "xlsx", "--yes"])

    assert r.exit_code == 1
    assert (tmp_path / "a.xlsx").is_file()
    assert "失败: bad.md" in r.output
    assert "完成: 1 个输出" in r.output


def test_output_conflict_auto_numbered(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    a = _md(tmp_path, "a.md")
    out = tmp_path / "a.xlsx"
    out.write_text("precious", encoding="utf-8")

    r = runner.invoke(app, ["markdown-convert", str(a), "--to", "xlsx", "--yes"])

    assert r.exit_code == 0
    assert out.read_text(encoding="utf-8") == "precious"
    assert (tmp_path / "a_1.xlsx").is_file()


def test_invalid_to_and_excel_mode_rejected(tmp_path):
    a = _md(tmp_path, "a.md")
    r = runner.invoke(app, ["markdown-convert", str(a), "--to", "pdf", "--yes"])
    assert r.exit_code == 1
    assert "无效的 --to: pdf" in r.output
    r = runner.invoke(
        app, ["markdown-convert", str(a), "--to", "xlsx", "--excel-mode", "bad", "--yes"]
    )
    assert r.exit_code == 1
    assert "无效的 --excel-mode: bad" in r.output


def test_dir_batch_and_unsupported_ignored(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    _md(tmp_path, "a.md")
    _md(tmp_path, "b.markdown")
    (tmp_path / "note.txt").write_text("x", encoding="utf-8")

    r = runner.invoke(app, ["markdown-convert", "--dir", str(tmp_path), "--to", "xlsx", "--yes"])

    assert r.exit_code == 0
    assert "已忽略 1 个不支持的文件" in r.output
    assert (tmp_path / "a.xlsx").is_file()
    assert (tmp_path / "b.xlsx").is_file()
    assert not (tmp_path / "note.txt.xlsx").exists()


def test_recursive_and_no_supported_files_errors(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    (tmp_path / "only.txt").write_text("x", encoding="utf-8")
    r = runner.invoke(app, ["markdown-convert", "--dir", str(tmp_path), "--recursive"])
    assert r.exit_code == 1
    assert "未选择任何受支持的 Markdown 文件" in r.output

    f = _md(tmp_path, "f.md")
    r = runner.invoke(app, ["markdown-convert", "--dir", str(f)])
    assert r.exit_code == 1
    assert "不是目录" in r.output


def test_missing_pandoc_friendly_error_no_traceback(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    from file_toolbox.core import markdown_convert

    a = _md(tmp_path, "a.md")
    monkeypatch.setattr(markdown_convert, "_find_spec", lambda name: None)

    r = runner.invoke(app, ["markdown-convert", str(a), "--yes"])

    assert r.exit_code == 1
    assert "pypandoc" in r.output
    assert "Traceback" not in r.output
    assert not (tmp_path / "a.docx").exists()


def test_cancel_reports_completed_outputs(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    a = _md(tmp_path, "a.md")
    b = _md(tmp_path, "b.md")
    import file_toolbox.cli.markdown_cmd as cmd
    from file_toolbox.core.markdown_convert import ConversionItem, ConversionResult

    class FakeService:
        def __init__(self, history_store=None) -> None:
            self.history_store = history_store

        def convert(
            self,
            files,
            output_dir,
            target="docx",
            excel_mode="tables",
            progress_callback=None,
            cancel_check=None,
        ):
            return ConversionResult(
                [
                    ConversionItem(source=files[0], output=files[0].with_suffix(".xlsx")),
                    ConversionItem(source=files[1], skipped=True),
                ],
                cancelled=True,
            )

    monkeypatch.setattr(cmd, "MarkdownConvertService", FakeService)

    r = runner.invoke(app, ["markdown-convert", str(a), str(b), "--to", "xlsx", "--yes"])

    assert r.exit_code == 1
    assert "已取消:已完成 1 个输出" in r.output
