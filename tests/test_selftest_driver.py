"""--selftest 入口与驱动(Issue #143)契约:参数 fail closed、lazy import、
场景语义、EVIDENCE_MISSING 不假 PASS、源码级 pure/reopen 端到端。

边界:
- 正常启动不得加载 heavy selftest 模块(显式参数分支才 lazy import);
- 缺值/非法模式退出码 2,绝不静默进入正常 GUI(可能触发真实更新);
- full 代表路径在 Office 预筛不可用时返回 EVIDENCE_MISSING(退出码 3);
- 纯路径端到端在源码形态(离屏 Qt + 临时数据根)走真实页面/worker/事件循环。
"""

import json
import subprocess
import sys
import threading
from pathlib import Path
from types import SimpleNamespace

import pytest

pytest.importorskip("PySide6.QtWidgets")

from PySide6.QtWidgets import QApplication, QMessageBox  # noqa: E402

from file_toolbox.gui_entry import _parse_selftest_args, _source_selftest_root  # noqa: E402

# ---------------------------------------------------------------------------
# 入口参数:显式 selftest fail closed;正常参数不影响启动
# ---------------------------------------------------------------------------


def test_parse_selftest_args_normal_startup_ignores_unknown():
    assert _parse_selftest_args([]) == (None, None)
    assert _parse_selftest_args(["--foo", "bar"]) == (None, None)


def test_parse_selftest_args_explicit_mode_and_report():
    mode, report = _parse_selftest_args(["--selftest", "pure", "--selftest-report", "r.json"])
    assert mode == "pure" and report == Path("r.json")
    mode, report = _parse_selftest_args(["--selftest=full"])
    assert mode == "full" and report is None


def test_parse_selftest_args_missing_value_fails_closed(capsys):
    """旧实现缺陷回归:'--selftest' 缺值曾静默返回 None 进入正常 GUI。"""
    with pytest.raises(SystemExit) as exc_info:
        _parse_selftest_args(["--selftest"])
    assert exc_info.value.code == 2
    assert "--selftest" in capsys.readouterr().err


def test_parse_selftest_args_invalid_mode_fails_closed(capsys):
    with pytest.raises(SystemExit) as exc_info:
        _parse_selftest_args(["--selftest", "bogus"])
    assert exc_info.value.code == 2
    assert "bogus" in capsys.readouterr().err


def test_parse_selftest_args_missing_report_value_fails_closed():
    with pytest.raises(SystemExit) as exc_info:
        _parse_selftest_args(["--selftest", "pure", "--selftest-report"])
    assert exc_info.value.code == 2


def test_parse_selftest_args_orphan_report_fails_closed(capsys):
    """F5 回归:孤立 --selftest-report(无 --selftest)不得静默进入普通 GUI。"""
    with pytest.raises(SystemExit) as exc_info:
        _parse_selftest_args(["--selftest-report", "r.json"])
    assert exc_info.value.code == 2
    assert "--selftest" in capsys.readouterr().err


def test_parse_selftest_args_tolerates_velopack_arguments():
    """Velopack hook 参数等未知参数不影响正常启动(不误判、不吞掉)。"""
    assert _parse_selftest_args(["--veloapp-install", "--selftest", "pure"])[0] == "pure"
    assert _parse_selftest_args(["--veloapp-obsolete"]) == (None, None)


def test_source_selftest_root_env_override(monkeypatch, tmp_path):
    monkeypatch.setenv("FILE_TOOLBOX_SELFTEST_DATA_ROOT", str(tmp_path / "shared"))
    assert _source_selftest_root() == tmp_path / "shared"
    monkeypatch.delenv("FILE_TOOLBOX_SELFTEST_DATA_ROOT")
    default = _source_selftest_root()
    assert "build" in str(default) and "selftest-runtime" in str(default)
    assert default.name.startswith("run-")


def test_normal_startup_does_not_import_selftest_driver():
    """正常启动链(入口/主窗口导入)不得加载 selftest 驱动(heavy 模块)。"""
    code = (
        "import sys\n"
        "import file_toolbox.gui_entry\n"
        "import file_toolbox.gui.main_window\n"
        "leaked = [m for m in sys.modules if 'selftest_driver' in m]\n"
        "print(','.join(leaked))\n"
    )
    proc = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True, check=True)
    assert proc.stdout.strip() == "", f"正常启动加载了 selftest 驱动: {proc.stdout.strip()}"


# ---------------------------------------------------------------------------
# 数据根门禁(F2):拒绝已有真实数据根,范围标记记下精确记录 ID
# ---------------------------------------------------------------------------


def test_selftest_refuses_existing_real_data_root(monkeypatch, tmp_path):
    """F2 回归:已有真实历史的数据根拒绝自测,文件/历史保持不变。"""
    from file_toolbox.gui.selftest_driver import ensure_selftest_data_root

    monkeypatch.chdir(tmp_path)
    real_root = tmp_path / ".file_toolbox"
    history = real_root / "history"
    history.mkdir(parents=True)
    history_line = (
        '{"id": 7, "data": {"rename_map": {"真实甲.txt": "真实乙.txt"}}, "undone": false}\n'
    )
    (history / "rename.jsonl").write_text(history_line, encoding="utf-8")
    real_file = tmp_path / "真实乙.txt"
    real_file.write_text("用户真实文件\n", encoding="utf-8")

    allowed, reason = ensure_selftest_data_root()
    assert allowed is False
    assert "拒绝" in reason
    # 拒绝不改动任何内容:历史/文件/无标记写入
    assert (history / "rename.jsonl").read_text(encoding="utf-8") == history_line
    assert real_file.read_text(encoding="utf-8") == "用户真实文件\n"
    assert not (real_root / "selftest-scope.json").is_file()


def test_selftest_scope_marker_and_record_roundtrip(monkeypatch, tmp_path):
    from file_toolbox.gui.selftest_driver import (
        ensure_selftest_data_root,
        load_selftest_scope,
        record_selftest_rename_record,
    )

    monkeypatch.chdir(tmp_path)
    assert tmp_path.joinpath(".file_toolbox").exists() is False
    allowed, reason = ensure_selftest_data_root()
    assert allowed is True, reason
    scope = load_selftest_scope()
    assert scope is not None and scope["rename_record_id"] is None
    # 同根第二次(reopen):标记在 → 允许
    allowed, _ = ensure_selftest_data_root()
    assert allowed is True
    record_selftest_rename_record(42)
    scope = load_selftest_scope()
    assert scope is not None and scope["rename_record_id"] == 42


def test_driver_refuses_real_data_root_without_touching_it(app, monkeypatch, tmp_path):
    """driver 层门禁同样拒绝:返回非零,不向被拒根写任何文件。"""
    import json as json_module

    from file_toolbox.gui import selftest_driver as driver

    monkeypatch.chdir(tmp_path)
    real_root = tmp_path / ".file_toolbox"
    real_root.mkdir()
    (real_root / "settings.json").write_text("{}", encoding="utf-8")
    report = tmp_path / "outside-root" / "report.json"
    report.parent.mkdir()
    code = driver._execute_selftest(app, "pure", report)
    assert code == driver.EXIT_FAIL
    payload = json_module.loads(report.read_text(encoding="utf-8"))
    assert payload["scenarios"][0]["name"] == "scope"
    assert payload["scenarios"][0]["status"] == "fail"
    assert list(real_root.iterdir()) == [real_root / "settings.json"]  # 未写入


def test_reopen_requires_scope_record_id(app, monkeypatch, tmp_path):
    """F2 回归:范围标记缺记录 ID 时 reopen 失败,不碰任何历史。"""
    from file_toolbox.gui import selftest_driver as driver

    monkeypatch.chdir(tmp_path)
    assert driver.ensure_selftest_data_root()[0] is True  # 仅建标记,无首轮 rename
    outcome = driver._scenario_reopen_history_undo(_reopen_ctx(tmp_path))
    assert outcome.status == "fail"
    assert "记录 ID" in outcome.detail


def _reopen_ctx(tmp_path: Path):
    """reopen 场景的最小上下文(工作目录在临时根内)。"""
    ctx = SimpleNamespace(
        app=None,
        window=None,
        workdir=tmp_path / "selftest-work",
        coordinator=None,
        boxes=None,
        outcomes=[],
    )
    return ctx


# ---------------------------------------------------------------------------
# 驱动组件:模态框脚本、fake 协调器、聚合、场景表
# ---------------------------------------------------------------------------


def test_message_box_script_answers_apply_or_yes_and_restores():
    from file_toolbox.gui.selftest_driver import _MessageBoxScript

    with _MessageBoxScript() as script:
        # 打桩期间是纯 Python 函数(带 __code__);恢复后回到 PySide6 内置方法
        assert hasattr(QMessageBox.question, "__code__")
        assert QMessageBox.question(None, "t", "确认执行") == QMessageBox.StandardButton.Yes
        apply_buttons = QMessageBox.StandardButton.Apply | QMessageBox.StandardButton.Cancel
        assert (
            QMessageBox.question(None, "t", "确认更新", apply_buttons)
            == QMessageBox.StandardButton.Apply
        )
        assert QMessageBox.information(None, "t", "完成") == QMessageBox.StandardButton.Ok
        assert [record["kind"] for record in script.records] == [
            "question",
            "question",
            "information",
        ]
        assert any("确认更新" in record["text"] for record in script.records)
    assert not hasattr(QMessageBox.question, "__code__")
    assert not hasattr(QMessageBox.information, "__code__")


def test_fake_coordinator_download_honors_real_cancel_gate():
    from file_toolbox.gui.selftest_driver import SelftestUpdateCoordinator
    from file_toolbox.updater.coordinator import UpdateRequest
    from file_toolbox.updater.models import UpdateApplyStatus

    coordinator = SelftestUpdateCoordinator()
    request = UpdateRequest()
    result_box: list = []

    def run_download() -> None:
        result_box.append(coordinator.download_and_apply(request=request))

    thread = threading.Thread(target=run_download, daemon=True)
    thread.start()
    assert coordinator.release.wait(0) is False
    assert request.cancel() is True  # 真实取消门:更新页取消按钮同一路径
    thread.join(timeout=5)
    assert not thread.is_alive()
    assert result_box[0].status is UpdateApplyStatus.CANCELLED
    assert coordinator.download_calls == 1


def test_aggregate_exit_codes():
    from file_toolbox.gui.selftest_driver import (
        EXIT_EVIDENCE_MISSING,
        EXIT_FAIL,
        EXIT_PASS,
        ScenarioOutcome,
        _aggregate,
    )

    ok = ScenarioOutcome("a", "pass")
    assert _aggregate([ok]) == EXIT_PASS
    assert _aggregate([ok, ScenarioOutcome("b", "evidence_missing", "no office")]) == (
        EXIT_EVIDENCE_MISSING
    )
    assert _aggregate(
        [ok, ScenarioOutcome("b", "evidence_missing"), ScenarioOutcome("c", "fail")]
    ) == (EXIT_FAIL)


def test_scenarios_for_modes():
    from file_toolbox.gui.selftest_driver import _scenarios_for

    pure = [name for name, _ in _scenarios_for("pure")]
    assert pure == [
        "pages",
        "rename",
        "pdf_pure",
        "pdf_cancel",
        "markdown",
        "replace_pure",
        "update_mutex",
    ]
    full = [name for name, _ in _scenarios_for("full")]
    assert full[: len(pure)] == pure
    assert full[-3:] == ["pdf_office", "replace_office", "attendance"]
    assert [name for name, _ in _scenarios_for("reopen")] == ["reopen_history_undo"]
    with pytest.raises(ValueError, match="未知的 selftest 模式"):
        _scenarios_for("bogus")


def test_office_scenarios_report_evidence_missing_not_pass(monkeypatch):
    """Office 预筛不可用 → EVIDENCE_MISSING,不得 skip 成 PASS(AC 边界)。"""
    from file_toolbox.gui import selftest_driver as driver

    monkeypatch.setattr(driver, "_office_ready", lambda tool_id, kinds: (False, "Word: 未检测到"))
    ctx = SimpleNamespace(outcomes=[], workdir=Path("."), app=None)
    for name, scenario in (
        ("pdf_office", driver._scenario_pdf_office),
        ("replace_office", driver._scenario_replace_office),
        ("attendance", driver._scenario_attendance),
    ):
        outcome = scenario(ctx)
        assert outcome.status == "evidence_missing", name
        assert "未检测到" in outcome.detail


# ---------------------------------------------------------------------------
# 考勤虚构样本与方案(纯 openpyxl,不触发 COM)
# ---------------------------------------------------------------------------


def test_attendance_samples_follow_corrected_pitfalls(tmp_path):
    from openpyxl import load_workbook

    from file_toolbox.gui.selftest_driver import _build_attendance_samples

    samples = _build_attendance_samples(tmp_path)
    source = load_workbook(samples["source"], read_only=True)
    try:
        sheet = source["Sheet1"]
        departments = {sheet.cell(row, 3).value for row in range(2, 6) if sheet.cell(row, 1).value}
        assert departments == {"验收部门"}  # 坑 1:部门统一
    finally:
        source.close()
    roster_source = load_workbook(samples["source_roster"], read_only=True)
    try:
        names = [
            roster_source["Sheet1"].cell(row, 1).value
            for row in range(2, 5)
            if roster_source["Sheet1"].cell(row, 1).value
        ]
        assert len(set(names)) == len(names)  # 坑 3:名单源姓名唯一
    finally:
        roster_source.close()
    template = load_workbook(samples["template"], read_only=True)
    try:
        detail = template["出勤明细"]
        headers = [detail.cell(6, 3 + day).value for day in range(1, 31)]
        assert headers == [str(day) for day in range(1, 31)]  # 坑 2:字符串日期表头
        assert isinstance(headers[0], str)
        summary = template["考勤汇总表"]
        assert str(summary.cell(8, 5).value).startswith("=COUNTIF")
    finally:
        template.close()


def test_attendance_plan_roster_mode(tmp_path):
    from file_toolbox.gui.selftest_driver import _attendance_plan, _build_attendance_samples

    samples = _build_attendance_samples(tmp_path)
    plain = _attendance_plan(samples["template"], roster_mode=False)
    assert plain.roster is None
    assert [config.attendance_group for config in plain.group_sheet_configs] == ["甲组", "乙组"]
    roster = _attendance_plan(samples["template"], roster_mode=True)
    assert roster.roster is not None
    assert roster.roster.excluded_employee_ids == ()  # 排除经真实预览勾选(900003)
    assert roster.roster.workbook_path.name == "roster.xlsx"


# ---------------------------------------------------------------------------
# 源码级端到端:真实 MainWindow + 真实页面/worker/事件循环(离屏、临时数据根)
# ---------------------------------------------------------------------------


@pytest.fixture(scope="module")
def app():
    return QApplication.instance() or QApplication([])


def _run_mode(app, monkeypatch, tmp_path, mode: str) -> tuple[int, dict]:
    from file_toolbox.gui import selftest_driver as driver

    monkeypatch.chdir(tmp_path)  # CLI 数据根策略 → 临时 .file_toolbox,隔离真实配置/历史
    report_path = tmp_path / f"report-{mode}.json"
    code = driver._execute_selftest(app, mode, report_path)
    payload = json.loads(report_path.read_text(encoding="utf-8"))
    return code, payload


def test_execute_selftest_pure_then_reopen_e2e(app, monkeypatch, tmp_path):
    """pure:全场景真实执行并读回;reopen:同数据根历史 + 真实 HistoryDialog 撤销。"""
    code, payload = _run_mode(app, monkeypatch, tmp_path, "pure")
    assert code == 0, json.dumps(payload, ensure_ascii=False)
    statuses = {item["name"]: item["status"] for item in payload["scenarios"]}
    assert statuses == {
        "pages": "pass",
        "rename": "pass",
        "pdf_pure": "pass",
        "pdf_cancel": "pass",
        "markdown": "pass",
        "replace_pure": "pass",
        "update_mutex": "pass",
    }
    assert payload["evidence_missing"] is False
    # 数据根与工作目录都落在临时根内(不写真实配置/历史)
    assert str(payload["data_root"]).startswith(str(tmp_path))
    # 取消场景:产出数少于全部文件即真实取消,或全部完成(两种合法终态)且控件恢复
    cancel = next(item for item in payload["scenarios"] if item["name"] == "pdf_cancel")
    assert 1 <= len(cancel["artifacts"]["outputs"]) <= 3

    code, reopen_payload = _run_mode(app, monkeypatch, tmp_path, "reopen")
    assert code == 0, json.dumps(reopen_payload, ensure_ascii=False)
    reopen = reopen_payload["scenarios"][0]
    assert reopen["name"] == "reopen_history_undo" and reopen["status"] == "pass"
    for restored in reopen["artifacts"]["restored"]:
        assert Path(restored).is_file()
        assert (
            not Path(restored)
            .with_name(Path(restored).name.replace("selftest", "ftb-renamed"))
            .exists()
        )


def test_execute_selftest_full_without_office_is_evidence_missing(app, monkeypatch, tmp_path):
    """本机无 Office(预筛不可用)时 full → 退出码 3,不把缺失当 PASS。"""
    from file_toolbox.gui import selftest_driver as driver

    def fake_ready(tool_id: str, kinds: tuple[str, ...]) -> tuple[bool, str]:
        wanted = {"word": "Word", "excel": "Excel", "ppt": "PowerPoint"}
        problems = [f"{wanted.get(kind, kind)}: 未检测到" for kind in kinds]
        return (False, "; ".join(problems))

    monkeypatch.setattr(driver, "_office_ready", fake_ready)
    monkeypatch.chdir(tmp_path)
    code, payload = _run_mode(app, monkeypatch, tmp_path, "full")
    assert code == 3, json.dumps(payload, ensure_ascii=False)
    statuses = {item["name"]: item["status"] for item in payload["scenarios"]}
    assert statuses["pdf_office"] == "evidence_missing"
    assert statuses["replace_office"] == "evidence_missing"
    assert statuses["attendance"] == "evidence_missing"
    assert statuses["rename"] == "pass"  # 纯路径不受 Office 缺失影响
    assert payload["evidence_missing"] is True


# ---------------------------------------------------------------------------
# F1:关闭超时保留现场——期限先写失败报告,但绝不提前返回/退出
# ---------------------------------------------------------------------------


def test_close_timeout_writes_report_then_waits_real_finish(app, monkeypatch, tmp_path):
    """F1 回归:受控 worker 超出收尾期限后释放——期限内 driver 未返回
    (已先写失败报告),真实 finished/窗口关闭后才以非零返回。"""
    import time as time_module

    from PySide6.QtCore import QThread

    from file_toolbox.gui import selftest_driver as driver

    released_at: dict[str, float] = {}
    write_times: list[float] = []
    original_write = driver._write_report

    def spy_write(*args, **kwargs):
        write_times.append(time_module.monotonic())
        return original_write(*args, **kwargs)

    class _DelayedWorker(QThread):
        def run(self) -> None:
            time_module.sleep(0.8)
            released_at["t"] = time_module.monotonic()

    def controlled_scenario(ctx):
        tab = ctx.window._rename_tab
        worker = _DelayedWorker(tab)  # 挂在页面下:closeEvent 按运行中 worker 协作延迟
        worker.start()
        return driver._pass("controlled")

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        driver, "_scenarios_for", lambda mode: [("controlled", controlled_scenario)]
    )
    monkeypatch.setattr(driver, "_write_report", spy_write)
    started = time_module.monotonic()
    code = driver._execute_selftest(app, "pure", tmp_path / "close-report.json", close_wait_s=0.2)
    returned_at = time_module.monotonic()
    assert code == driver.EXIT_FAIL
    scenarios = {
        item["name"]: item["status"]
        for item in json.loads((tmp_path / "close-report.json").read_text(encoding="utf-8"))[
            "scenarios"
        ]
    }
    assert scenarios == {"controlled": "pass", "close": "fail"}
    # 关键契约:worker 真实结束(释放时刻)之后才返回,未在 0.2s 期限处截断
    assert returned_at >= released_at["t"]
    assert returned_at - started >= 0.8
    # 期限内先写失败报告,收尾后再写终版
    assert len(write_times) >= 2
    assert write_times[0] < released_at["t"]


# ---------------------------------------------------------------------------
# F6:结果输出成功但严格清理失败(cleanup_warning)→ 场景/总结果非零,产物保留
# ---------------------------------------------------------------------------


def test_track_observer_subscribes_strictly_before_start(app):
    """F6 竞态根源:TaskLifecycle.track 观察缝在 start 前订阅。"""
    from PySide6.QtCore import QThread, Signal
    from PySide6.QtWidgets import QWidget

    from file_toolbox.gui.task_lifecycle import TaskLifecycle

    class _ImmediateFailWorker(QThread):
        failed = Signal(str)

        def run(self) -> None:
            self.failed.emit("立即失败")  # 线程一启动即发射,不留连接窗口

    lifecycle = TaskLifecycle(QWidget())
    seen: list[str] = []
    lifecycle.on_worker_tracked = lambda worker: worker.failed.connect(seen.append)
    worker = _ImmediateFailWorker()
    lifecycle.track(worker)
    assert lifecycle.worker is worker
    worker.start()
    assert worker.wait(5000)
    app.processEvents()
    assert seen == ["立即失败"]  # 订阅先于 start,立即发射也不丢


def test_immediate_worker_failure_captured_via_track_seam(app, monkeypatch, tmp_path):
    """F6 竞态回归:batch_generate 立即抛错(worker 可能在主线程连接前发射 failed)。

    此前 71 pass 的回归只覆盖"慢任务结束后清理失败"的发射时序,未覆盖 start
    前竞态;track 观察缝订阅严格先于 start → 场景 detail 必含 failed 信号证据,
    而非仅靠缺输出推断。
    """
    from file_toolbox.core.batch_pdf.service import PDFGeneratorService
    from file_toolbox.gui import selftest_driver as driver

    def boom(self: object, *args: object, **kwargs: object) -> object:
        raise RuntimeError("立即失败")

    monkeypatch.setattr(PDFGeneratorService, "batch_generate", boom)
    monkeypatch.chdir(tmp_path)
    code, payload = _run_mode(app, monkeypatch, tmp_path, "pure")
    assert code == driver.EXIT_FAIL
    pdf_pure = next(item for item in payload["scenarios"] if item["name"] == "pdf_pure")
    assert pdf_pure["status"] == "fail"
    assert "failed: 立即失败" in pdf_pure["detail"]  # 信号捕获,非仅缺输出


def test_pdf_cleanup_warning_fails_scenario_and_keeps_outputs(app, monkeypatch, tmp_path):
    """F6 回归:PDF 输出已写出但 svc.close(strict) 抛错 → cleanup_warning
    信号使场景失败,输出文件保留,总退出码非零(旧实现仅记录模态框仍 PASS)。"""
    from file_toolbox.core.batch_pdf.service import PDFGeneratorService
    from file_toolbox.gui import selftest_driver as driver

    original_close = PDFGeneratorService.close

    def broken_close(self, _from_del=False, *, strict=False):
        if strict:
            raise RuntimeError("模拟 COM 严格清理失败")
        return original_close(self, _from_del, strict=strict)

    monkeypatch.setattr(PDFGeneratorService, "close", broken_close)
    monkeypatch.chdir(tmp_path)
    code, payload = _run_mode(app, monkeypatch, tmp_path, "pure")
    assert code == driver.EXIT_FAIL
    pdf_pure = next(item for item in payload["scenarios"] if item["name"] == "pdf_pure")
    assert pdf_pure["status"] == "fail"
    assert "cleanup_warning" in pdf_pure["detail"]
    # 产物保留:输出仍存在且可读页数(不以文件存在放行,但也不删除)
    workdir = Path(payload["workdir"])
    outputs = sorted(workdir.glob("pdf-pure/pure-image-*.pdf"))
    assert outputs
    from pypdf import PdfReader

    assert all(len(PdfReader(str(path)).pages) >= 1 for path in outputs)
